# -*- coding: utf-8 -*-
"""레시피 템플릿(recipes/Recipe_Template.csv·.xlsx) — 두 파일의 열이 같고, 엑셀 입력 제한·설명 시트가 열과 맞는지,
머리글 메모 상자가 줄지 않았는지. 열을 하나 추가할 때 한쪽만 고치는 실수를 막는다.
열은 머리글 이름으로, 데이터 범위는 # 열 입력 제한의 마지막 행으로 찾는다(예시 행을 사용자가 지우거나 늘려도 그대로 동작)."""
import asyncio
import os
import sys
import csv
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest                                                     # noqa: E402
from openpyxl import load_workbook                                # noqa: E402

CSV_PATH = os.path.join(_ROOT, "recipes", "Recipe_Template.csv")
XLSX_PATH = os.path.join(_ROOT, "recipes", "Recipe_Template.xlsx")


def _csv_rows():
    with open(CSV_PATH, encoding="utf-8-sig", newline="") as fp:
        return list(csv.reader(fp))


@pytest.fixture(scope="module")
def wb():
    return load_workbook(XLSX_PATH)


def _dvs(ws):
    return {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}


def _cols(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}


def _last_row(ws):
    return max(int(str(dv.sqref).split(":")[1][1:]) for dv in ws.data_validations.dataValidation
               if str(dv.sqref).startswith("A2:"))


def _dv(ws, name):
    L = ws.cell(1, _cols(ws)[name]).column_letter
    return _dvs(ws)[f"{L}2:{L}{_last_row(ws)}"]


def test_csv_and_xlsx_columns_match(wb):
    ws = wb["Recipe"]
    xlsx = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = _csv_rows()
    assert xlsx == rows[0] and all(len(r) == len(rows[0]) for r in rows)
    assert xlsx[-4:] == ["use_heater", "heater_temp", "heater_ramp", "heater_soak"]
    i = xlsx.index("process_time")
    assert xlsx[i + 1:i + 6] == ["dep_rate", "thickness", "tooling", "density", "z_factor"]
    j = xlsx.index("rf_pulse_duty_cycle")
    assert xlsx[j + 1:j + 3] == ["power_change_time", "change_power_value"]


def test_every_column_has_comment_and_validation(wb):
    ws = wb["Recipe"]
    dvs = _dvs(ws)
    last = _last_row(ws)
    assert last >= 60
    for c in range(1, ws.max_column + 1):
        col = ws.cell(1, c).column_letter
        assert ws.cell(1, c).comment is not None, ws.cell(1, c).value
        assert f"{col}2:{col}{last}" in dvs, ws.cell(1, c).value


def test_heater_and_special_validations(wb):
    ws = wb["Recipe"]
    assert (_dv(ws, "use_heater").type, _dv(ws, "use_heater").formula1) == ("list", '"T,F"')
    t = _dv(ws, "heater_temp")
    assert (t.type, t.formula1, t.formula2) == ("decimal", "1", "900")
    r = _dv(ws, "heater_ramp")
    assert (r.type, r.formula1, r.formula2) == ("decimal", "0.1", "60")
    a = _dv(ws, "#")
    assert a.type == "custom" and 'LOWER(A2)="test"' in a.formula1                              # TEST 행 허용
    assert _dv(ws, "use_dc_power2").formula1 == '"F"'                                          # DC2 폐지
    # 열을 끼워 넣어도 수식 입력 제한·조건부 서식이 자기 열을 가리키는지(Off Time·RF duty)
    cols = _cols(ws)
    off, freq = ws.cell(1, cols["dc_pulse_off_time"]).column_letter, ws.cell(1, cols["dc_pulse_freq"]).column_letter
    assert f"UPPER({off}2)" in _dv(ws, "dc_pulse_off_time").formula1 and f"400/{freq}2" in _dv(ws, "dc_pulse_off_time").formula1
    cf = {str(c.sqref): [f for rr in c.rules for f in rr.formula] for c in ws.conditional_formatting}
    duty, rfq = ws.cell(1, cols["rf_pulse_duty_cycle"]).column_letter, ws.cell(1, cols["rf_pulse_freq"]).column_letter
    assert any(f"1.6*{rfq}2" in f and f"ISNUMBER({duty}2)" in f for f in cf.get(f"{duty}2:{duty}{_last_row(ws)}", []))
    assert any(f"UPPER({off}2)" in f for f in cf.get(f"{off}2:{off}{_last_row(ws)}", []))


def test_thickness_and_power_change_columns(wb):
    ws, desc = wb["Recipe"], wb["설명"]
    cols = _cols(ws)
    for name, unit in (("dep_rate", "nm/s"), ("thickness", "nm")):
        d = _dv(ws, name)
        assert (d.type, d.formula1) == ("decimal", "0") and d.promptTitle == f"{name} [{unit}]"
        assert ws.cell(1, cols[name]).comment.text.startswith(f"단위: {unit}\n")
    assert "Thickness 모드" in _dv(ws, "thickness").prompt and "Factor" in _dv(ws, "thickness").prompt
    p = _dv(ws, "power_change_time")
    assert p.type == "textLength" and "E844" in p.prompt and "'초'" in p.prompt
    v = _dv(ws, "change_power_value")
    assert (v.type, v.operator, v.formula1) == ("decimal", "greaterThan", "0")
    last = _last_row(ws)
    for name in ("dep_rate", "thickness", "power_change_time", "change_power_value"):
        assert all(ws.cell(r, cols[name]).value is None for r in range(2, last + 1)), name   # 예시 행은 비움
    drow = {desc.cell(r, 1).value: r for r in range(5, desc.max_row + 1)}
    assert "Thickness 모드" in desc.cell(drow["thickness"], 3).value
    assert "E844" in desc.cell(drow["power_change_time"], 3).value + str(desc.cell(drow["power_change_time"], 4).value)
    # CSV 예시 행도 새 칸은 비어 있다(기존 예시 공정이 Thickness 모드로 바뀌지 않게)
    rows = _csv_rows()
    for name in ("dep_rate", "thickness", "power_change_time", "change_power_value"):
        k = rows[0].index(name)
        assert all(r[k] == "" for r in rows[1:]), name


def test_recipe_factor_columns(wb):
    """tooling·density·z_factor — thickness 바로 뒤, 범위 입력 제한, 메모·설명(Thickness 모드에서만·레시피 우선·비우면 Factor), 예시 행 비움."""
    ws, desc = wb["Recipe"], wb["설명"]
    cols = _cols(ws)
    last = _last_row(ws)
    drow = {desc.cell(r, 1).value: r for r in range(5, desc.max_row + 1)}
    rows = _csv_rows()
    for name, lo, hi in (("tooling", "10", "399"), ("density", "0.5", "99.99"), ("z_factor", "0.1", "9.999")):
        d = _dv(ws, name)
        assert (d.type, d.operator, d.formula1, d.formula2) == ("decimal", "between", lo, hi), name
        txt = ws.cell(1, cols[name]).comment.text
        assert "Thickness 모드에서만" in txt and "레시피" in txt and "Factor" in txt and f"{lo}" in d.prompt.replace("0.500", "0.5")
        assert all(ws.cell(r, cols[name]).value is None for r in range(2, last + 1)), name
        assert "Thickness 모드에서만" in desc.cell(drow[name], 3).value and "Factor" in desc.cell(drow[name], 4).value
        k = rows[0].index(name)
        assert all(r[k] == "" for r in rows[1:]), name


def test_validate_rejects_out_of_range_factor_values():
    from runtime.chamber_runtime import ChamberRuntime
    rt = ChamberRuntime.__new__(ChamberRuntime)
    ok = {"Process_name": "t", "thickness": "100", "tooling": "10", "density": "99.99", "z_factor": "0.1"}
    assert rt._validate_raw_rows([ok]) == []
    assert rt._validate_raw_rows([dict(ok, tooling="", density="", z_factor="")]) == []        # 비우면 Factor 값
    errs = rt._validate_raw_rows([dict(ok, tooling="5", density="abc", z_factor="10")])
    assert len(errs) == 3 and "tooling='5' → 범위 10~399 % 밖" in errs[0]
    assert "density='abc' → 숫자가 아닙니다" in errs[1] and "z_factor='10' → 범위 0.1~9.999 밖" in errs[2]


def test_description_sheet_lists_every_column(wb):
    ws, desc = wb["Recipe"], wb["설명"]
    cols = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    listed = [desc.cell(r, 1).value for r in range(5, desc.max_row + 1)]
    assert listed == cols


def test_freeze_panes_keep_horizontal_scroll(wb):
    """Recipe 시트 틀 고정 = C2(1행 + #·Process_name). 고정 열이 많으면(예: AE2 = 30열) 화면보다 넓어 좌우 스크롤이 안 된다."""
    ws = wb["Recipe"]
    assert ws.freeze_panes == "C2"
    pane = ws.sheet_view.pane
    assert pane.state == "frozen" and (pane.xSplit or 0) <= 2 and pane.ySplit == 1
    panes = [s.pane for s in ws.sheet_view.selection]
    assert len(panes) == len(set(panes)), panes                      # 같은 pane 의 selection 이 겹쳐 쌓이지 않게
    assert ws.max_column == 48 and all(ws.cell(1, c).comment is not None for c in range(1, 49))


def test_comment_boxes_not_shrunk():
    """머리글 메모 상자가 openpyxl 기본 크기(144×79)로 줄지 않았는지 — 줄면 설명 글이 잘린다.
    openpyxl 은 기존 메모 상자 크기를 읽지 못해 다시 저장하면 기본값으로 바꾸므로, 템플릿을 고치는 스크립트가 크기를 다시 넣어야 한다."""
    with zipfile.ZipFile(XLSX_PATH) as z:
        vml = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".vml"))
    assert vml and "width:144px;height:79px" not in vml


def test_loader_reads_new_columns():
    """템플릿 머리글 그대로 쓴 행 — 로더가 dep_rate·thickness·power_change_time·change_power_value 를 읽는다(이미 지원하는 이름)."""
    from runtime.chamber_runtime import ChamberRuntime

    async def go():
        rt = ChamberRuntime.__new__(ChamberRuntime)
        rt.ch = 2
        rt.append_log = lambda *a: None
        rt._get_text = lambda name: ""
        rt.supports_rf_pulse = rt.supports_dc_pulse = rt.supports_dc_cont = rt.supports_rf_cont = True
        rt.supports_dc_cont2 = False
        head = _csv_rows()[0]
        row = {h: "" for h in head}
        row.update({"#": "1", "Process_name": "t", "base_pressure": "7e-6", "working_pressure": "5",
                    "shutter_delay": "5", "dep_rate": "0.05", "thickness": "100", "gun1": "T", "Ar": "T",
                    "Ar_flow": "20", "use_dc_pulse": "T", "dc_pulse_power": "80",
                    "power_change_time": "2m", "change_power_value": "100"})
        assert rt._validate_raw_rows([row]) == []
        n = rt._normalize_params_for_process(row)
        assert n["dep_rate"] == 0.05 and n["thickness"] == 100.0
        assert n["process_time"] == pytest.approx(100 / 0.05 / 60)        # process_time 공란 → 역산(예상시간)
        assert n["power_change_time"] == "2m" and n["change_power_value"] == "100"
        assert (n["tooling"], n["density"], n["z_factor"]) == (None, None, None)          # 비우면 None(Factor 팝업 값)
        n = rt._normalize_params_for_process(dict(row, tooling="120", density="2.73", z_factor="1.08"))
        assert (n["tooling"], n["density"], n["z_factor"]) == (120.0, 2.73, 1.08)
        bad = dict(row, change_power_value="")
        assert rt._validate_raw_rows([bad])                                 # 두 칸은 함께 입력
    asyncio.run(go())


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
