# -*- coding: utf-8 -*-
"""레시피 heater_soak(분) — 공정 행이 히터 RAMP → 도달 확인 → SOAK N분(히터 유지) → 베이스 압력 … 순서로 돈다.

도달 확인은 예전 그대로(목표 ±HEATER_SOAK_TOLERANCE 에 HEATER_SOAK_TIME_SEC 연속, 시간 초과 E831).
SOAK 은 그 뒤 시간만 센다 — 온도는 래더 PID 가 계속 맞추고, 흔들려도 다시 세지 않고 로그·카드에 기록만.
히터 꺼짐·이상·출력 무응답/끊김은 SOAK 중에도 공정을 멈춘다. 공란·열 없음이면 SOAK 없음(예전과 같음).
템플릿(CSV·엑셀)에 heater_soak 열, 히터 설명은 현재 설정값(목표 상한 900·램프 기본 6°C/min).
실제 HeaterService + 래더·열 에뮬레이터(tests/heater_ladder_sim.py, 가상 시계). 장치·네트워크 접근 없음.
"""
import os
import sys
import csv
import asyncio
import contextlib

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _ROOT)
sys.path.insert(0, _HERE)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
with contextlib.suppress(Exception):
    sys.stdout.reconfigure(encoding="utf-8")

import pytest                                                          # noqa: E402
from openpyxl import load_workbook                                     # noqa: E402
import heater_ladder_sim as sim                                        # noqa: E402
from lib import config_common as cfgc                                  # noqa: E402
import runtime.chamber_runtime as cr                                   # noqa: E402
from controller.process_controller import ProcessController, ActionType             # noqa: E402
from controller.heater_process import HeaterSoakWaiter                 # noqa: E402

CSV_PATH = os.path.join(_ROOT, "recipes", "Recipe_Template.csv")
XLSX_PATH = os.path.join(_ROOT, "recipes", "Recipe_Template.xlsx")
# 장비 settings.json ch2 와 같은 히터 값(에뮬레이터 기본값 대신)
REAL = dict(mv_limit=1200, mv_abs_max=1300, ramp_rate=6.0, max_temp=900.0, sv_limit_c=900.0, ot_limit_c=950.0,
            hold_mode="off")


class _PCRec:
    """공정 컨트롤러 대신 on_heater_* 호출만 기록."""

    def __init__(self):
        self.calls = []

    def __getattr__(self, name):
        if name.startswith("on_heater"):
            return lambda *a, **k: self.calls.append((name, a, k))
        raise AttributeError(name)

    def names(self):
        return [c[0] for c in self.calls]


class _FieldChat(sim.FakeChat):
    """에뮬레이터 가짜 챗 + 카드 필드까지 남긴다."""

    def __init__(self, clk):
        super().__init__(clk)
        self.full = []

    def _post_card(self, title, subtitle="", status="INFO", fields=None, urgent=False):
        super()._post_card(title, subtitle, status, fields, urgent)
        self.full.append((title, status, dict(fields or {})))


def _rt():
    """런타임 최소 인스턴스 + 실제 _bind_process_controller 콜백. (rt, 콜백을 가진 진짜 PC, 기록기)"""
    rt = cr.ChamberRuntime.__new__(cr.ChamberRuntime)
    rt.ch = 2
    rt.logs = []
    rt.append_log = lambda s, m: rt.logs.append((None, s, m))
    rt._get_text = lambda name: ""
    rt.supports_rf_pulse = rt.supports_dc_pulse = rt.supports_dc_cont = rt.supports_rf_cont = True
    rt.supports_dc_cont2 = False
    rt._loop = asyncio.get_running_loop()
    rt._ensure_task_alive = lambda *a, **k: None
    rt._heater = None
    rt._heater_gen = 0
    rt._heater_job = None
    rt._heater_watch = False
    rt._heater_rate = None
    rt._runner_queue_mode = True
    rt.process_queue = []
    rt.current_process_index = -1
    rt._bind_process_controller()
    pc = rt.process_controller
    rec = _PCRec()
    rt.process_controller = rec
    return rt, pc, rec


def _logs(rt, word):
    return [(t, m) for t, s, m in rt.logs if word in m]


async def _tick(n=5):
    for _ in range(n):
        await asyncio.sleep(0)


async def _heat(target, soak_s, **cfg):
    """HEATER_SET → 승온 대기(soak_s=None 이면 인자 1개 — 예전 호출) 시작까지. (rt, rec, svc, plc, clk, chat, t0)"""
    rt, pc, rec = _rt()
    svc, plc, clk, _logs_, _chat = sim.make(**cfg)
    chat = _FieldChat(clk)
    svc._chat = chat
    rt.append_log = lambda s, m: rt.logs.append((clk.t, s, m))
    rt.attach_heater(svc)
    await sim.run(svc, plc, clk, 5)
    pc._heater_set(float(target), None)
    await _tick()
    await sim.run(svc, plc, clk, 2)
    assert rec.names() == ["on_heater_set_ok"], rec.calls
    t0 = clk.t
    if soak_s is None:
        pc._heater_wait(float(target))
    else:
        pc._heater_wait(float(target), float(soak_s))
    return rt, rec, svc, plc, clk, chat, t0


async def _run_until(svc, plc, clk, cond, limit_s, on_step=None):
    t0 = clk.t
    while not cond():
        await sim.run(svc, plc, clk, 1, on_step=on_step)
        await _tick()
        assert clk.t - t0 < limit_s, "시간 초과"


def _done(rec):
    return lambda: any(n in ("on_heater_ready", "on_heater_failed") for n in rec.names())


# ───────────────────────── 레시피 정규화·검증 ─────────────────────────
def test_normalize_validate_heater_soak():
    async def go():
        rt, _pc, _rec = _rt()
        svc, plc, clk, logs, chat = sim.make(**REAL)
        raw = {"Ar": "T", "Ar_flow": "10", "gun1": "T", "use_heater": "T", "heater_temp": "400"}
        assert rt._normalize_params_for_process(dict(raw, heater_soak="60"))["heater_soak"] == 60.0
        assert rt._normalize_params_for_process(dict(raw, heater_soak=""))["heater_soak"] is None
        assert rt._normalize_params_for_process(dict(raw))["heater_soak"] is None          # 열이 없는 예전 레시피
        n = rt._normalize_params_for_process(dict(raw, use_heater="F", heater_soak="60"))
        assert n["heater_soak"] is None and _logs(rt, "heater_soak=60.0 은 무시")

        rt._heater = svc
        base = dict(use_ar=True, ar_flow=10, use_g1=True, use_dc_pulse=True, dc_pulse_power=50,
                    use_heater=True, heater_temp=400.0)
        for ok in (None, 0.0, 0.5, 60.0, cr.HEATER_SOAK_MAX_MIN):
            assert rt._validate_norm_params(dict(base, heater_soak=ok)) == [], ok
        for bad in (-1.0, cr.HEATER_SOAK_MAX_MIN + 1, float("nan")):
            errs = rt._validate_norm_params(dict(base, heater_soak=bad))
            assert len(errs) == 1 and "heater_soak" in errs[0] and "480" in errs[0], (bad, errs)
        errs = rt._validate_raw_rows([{"use_heater": "T", "heater_temp": "400", "heater_soak": "1h"}])
        assert len(errs) == 1 and "heater_soak" in errs[0] and "숫자" in errs[0], errs
        assert rt._validate_raw_rows([{"use_heater": "T", "heater_temp": "400", "heater_soak": "60"}]) == []
    asyncio.run(go())


# ───────────────────────── 공정 단계 ─────────────────────────
def _bare_pc(**cb):
    def nop(*_a, **_k):
        return None
    kw = dict(send_plc=nop, send_mfc=nop, send_dc_power=nop, stop_dc_power=nop, send_rf_power=nop,
              stop_rf_power=nop, send_dc_power2=nop, stop_dc_power2=nop, start_dc_pulse=nop, stop_dc_pulse=nop,
              start_rf_pulse=nop, stop_rf_pulse=nop, ig_wait=nop, cancel_ig=nop, rga_scan=nop, oes_run=nop,
              ch=2, supports_dc_cont=True, supports_rf_cont=True, supports_dc_pulse=True, supports_rf_pulse=True)
    kw.update(cb)
    return ProcessController(**kw)


def test_wait_step_carries_soak_and_blank_keeps_one_arg_call():
    async def go():
        calls = []
        pc = _bare_pc(heater_set=lambda t, r: None, heater_wait=lambda *a: calls.append(a))
        p = dict(base_pressure=7e-6, working_pressure=60, process_time=1, shutter_delay=1, use_ar=True, ar_flow=50,
                 use_dc_pulse=True, dc_pulse_power=80, use_ms=True, use_g1=True, use_heater=True, heater_temp=400.0)
        seq = pc._create_process_sequence(dict(p, heater_soak=60.0))
        assert [s.action for s in seq[:3]] == [ActionType.HEATER_SET, ActionType.HEATER_WAIT, ActionType.IG_CMD]
        assert seq[1].params == (3600.0,) and seq[1].message == "히터 400°C 도달·안정 대기 (SOAK 60분)"
        old = pc._create_process_sequence(dict(p))                       # 공란 → 예전 그대로
        assert old[1].params is None and old[1].message == "히터 400°C 도달·안정 대기"
        assert [(s.action, s.message, s.value) for s in old[2:]] == [(s.action, s.message, s.value) for s in seq[2:]]

        assert pc._send_and_collect_tokens(seq[1])[0].kind == "HEATER_READY"
        assert pc._send_and_collect_tokens(old[1])[0].kind == "HEATER_READY"
        assert calls == [(400.0, 3600.0), (400.0,)]                       # 공란이면 인자 1개(예전 호출 그대로)
        zero = pc._create_process_sequence(dict(p, heater_soak=0.0))[1]   # 0분 = 도달 확인 뒤 바로(공란과 같은 동작)
        assert zero.params == (0.0,)
    asyncio.run(go())


# ───────────────────────── 400°C · SOAK 60분 (에뮬레이터) ─────────────────────────
def test_400c_soak_60min_counts_time_after_arrival_with_heater_on():
    async def go():
        rt, rec, svc, plc, clk, chat, t0 = await _heat(400.0, 3600.0, **REAL)
        runs = []
        await _run_until(svc, plc, clk, _done(rec), 4 * 3600,
                         on_step=lambda s, p: runs.append(bool(s.status().get("run"))))
        assert rec.names() == ["on_heater_set_ok", "on_heater_ready"], rec.calls
        assert all(runs)                                                    # 히터는 대기 내내 켜짐
        st = svc.status()
        assert st["run"] and svc.owner == "process" and abs(st["pv"] - 400.0) <= 2.0

        (t_band, _m), = _logs(rt, "히터 목표 도달 —")                        # 도달 확인(설정 60초)은 예전 그대로
        (t_hold, _m2), = _logs(rt, "SOAK 60분 시작")
        (t_ok, m_ok), = _logs(rt, "히터 준비 완료")
        assert 60 <= t_hold - t_band <= 61 and 3600 <= t_ok - t_hold <= 3601   # 도달 확인 뒤 60분
        assert 50 * 60 <= t_band - t0 <= 80 * 60                            # 25→400°C RAMP(6°C/min, 감속 접근)
        assert "SOAK 60분 끝" in m_ok and "SOAK 중 " in m_ok
        (_t, m_start), = _logs(rt, "그 뒤 SOAK 60분(레시피 heater_soak")
        assert "도달 한도" in m_start
        prog = _logs(rt, "히터 SOAK — ")
        assert len(prog) >= 55 and all(t_hold <= t <= t_ok for t, _m3 in prog) and "남은" in prog[0][1]

        reach = [c for c in chat.full if c[0] == "[CH2] 히터 도달"]
        assert len(reach) == 1 and reach[0][1] == "SUCCESS"
        f = reach[0][2]
        assert list(f) == ["목표", "도달 PV", "승온 소요", "SOAK", "SOAK 중 PV"]
        assert f["SOAK"] == "60분 0초" and f["SOAK 중 PV"].endswith("°C")
        ramp = int(t_hold - t0)
        assert f["승온 소요"] == f"{ramp // 60}분 {ramp % 60}초"             # SOAK 을 뺀 시간
    asyncio.run(go())


def test_blank_soak_uses_setting_and_card_unchanged():
    async def go():
        rt, rec, svc, plc, clk, chat, t0 = await _heat(200.0, None, mv_limit=700)
        await _run_until(svc, plc, clk, _done(rec), 3 * 3600)
        assert rec.names() == ["on_heater_set_ok", "on_heater_ready"], rec.calls
        assert svc.cfg.soak_time_s == 60
        (t_in, _m), = _logs(rt, "히터 목표 도달 —")
        (t_ok, _m2), = _logs(rt, "히터 준비 완료")
        assert 60 <= t_ok - t_in <= 61
        assert _logs(rt, "±2°C 를 60초 유지하면 진행")
        reach = [c for c in chat.full if c[0] == "[CH2] 히터 도달"]
        assert list(reach[0][2]) == ["목표", "도달 PV", "승온 소요"]           # 예전 카드 그대로
    asyncio.run(go())


async def _into_soak(soak_s):
    rt, rec, svc, plc, clk, chat, t0 = await _heat(200.0, soak_s, mv_limit=700)
    word = f"SOAK {soak_s / 60:g}분 시작"
    await _run_until(svc, plc, clk, lambda: bool(_logs(rt, word)), 3 * 3600)
    (t_hold, _m), = _logs(rt, word)
    return rt, rec, svc, plc, clk, t_hold


def test_wobble_during_soak_is_recorded_not_restarted():
    async def go():
        rt, rec, svc, plc, clk, t_hold = await _into_soak(600.0)
        await sim.run(svc, plc, clk, 300)
        await _tick()
        plc.tc_offset = 5.0                                                 # TC 가 5°C 높게 읽힘 → ±2°C 밖
        await sim.run(svc, plc, clk, 10)
        await _tick()
        plc.tc_offset = 0.0
        await _run_until(svc, plc, clk, _done(rec), 3600)
        assert rec.names() == ["on_heater_set_ok", "on_heater_ready"], rec.calls
        (t_ok, m_ok), = _logs(rt, "히터 준비 완료")
        assert 600 <= t_ok - t_hold <= 601                                  # 흔들려도 처음부터 다시 세지 않음
        (_t, m_out), = _logs(rt, "목표 ±2°C 밖(SOAK 시간은 계속 셈")
        (_t2, m_back), = _logs(rt, "히터 PV 범위 복귀")
        assert "벗어났었음" in m_back and "±2°C 밖" in m_ok
    asyncio.run(go())


@pytest.mark.parametrize("fault,code", [("interlock", "E830"), ("output_loss", "E832")])
def test_heater_trouble_during_soak_still_stops(fault, code):
    """SOAK 은 시간만 세지만 안전장치는 그대로 — 인터락 해제(RUN 꺼짐) E830, 출력 끊김(DAC 상한에서 식음) E832."""
    async def go():
        rt, rec, svc, plc, clk, t_hold = await _into_soak(3600.0)
        await sim.run(svc, plc, clk, 60)
        await _tick()
        if fault == "interlock":
            plc.vac_ok = False
        else:
            plc.gain = 0.0                                                  # 히터 출력이 열로 안 감(차단기·단선 흉내)
        await _run_until(svc, plc, clk, _done(rec), 1800)
        fails = [c for c in rec.calls if c[0] == "on_heater_failed"]
        assert len(fails) == 1 and fails[0][2].get("code") == code, rec.calls
        assert not _logs(rt, "히터 준비 완료") and clk.t - t_hold < 1800
    asyncio.run(go())


# ───────────────────────── 판정기 (순수) ─────────────────────────
class _Clk:
    def __init__(self):
        self.t = 1000.0                     # 실제 시계(monotonic)처럼 0 이 아닌 값에서 시작

    def __call__(self):
        return self.t


def test_waiter_hold_counts_time_and_ignores_arrival_timeout():
    clk, logs = _Clk(), []
    w = HeaterSoakWaiter(200.0, tol=2.0, soak_sec=60, timeout_s=300, hold_sec=600, clock=clk, log=logs.append)
    st = {"run": True, "fault": False, "stale": False, "pv": 200.3, "sv": 200.0, "sv_ramp": 200.0}
    assert w.tick(st, False) is None and logs[-1].startswith("히터 목표 도달 — 60초 안정 확인 중")
    clk.t = 1060
    assert w.tick(st, False) is None and "SOAK 10분 시작" in logs[-1] and w.hold_t0 == 1060
    clk.t = 1200
    assert w.tick(st, False) is None and logs[-1].startswith("히터 SOAK — 2.3분 / 10분, 남은 7.7분")
    clk.t = 1201
    assert w.tick({**st, "pv": 203.0}, False) is None and "목표 ±2°C 밖(SOAK 시간은 계속 셈" in logs[-1]
    clk.t = 1210
    assert w.tick(st, False) is None and logs[-1] == "히터 PV 범위 복귀 — 200.3°C (9초 벗어났었음)"
    clk.t = 1659                                                     # 도달 한도(300초)는 지났지만 SOAK 중이라 E831 아님
    assert w.tick(st, False) is None
    assert w.tick({**st, "run": False}, False)[2] == "E830"          # 히터 꺼짐은 SOAK 중에도 실패
    clk.t = 1660
    ok, msg, code = w.tick(st, False)
    assert ok is True and code is None and "SOAK 10분 끝" in msg and "200.3~203.0°C, ±2°C 밖 9초" in msg
    w2 = HeaterSoakWaiter(200.0, tol=2.0, soak_sec=60, timeout_s=300, hold_sec=600, clock=clk, log=logs.append)
    clk.t += 301                                                     # 도달 못 하면 예전처럼 E831
    assert w2.tick({**st, "pv": 150.0}, False)[2] == "E831"


# ───────────────────────── 템플릿 ─────────────────────────
def _tpl_cols(ws):
    """머리글 이름 → 열 번호, 데이터 마지막 행(# 열 입력 제한 범위)."""
    cols = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    last = max(int(str(dv.sqref).split(":")[1][1:]) for dv in ws.data_validations.dataValidation
               if str(dv.sqref).startswith("A2:"))
    return cols, last


def test_template_has_soak_column_and_current_heater_values():
    with open(CSV_PATH, encoding="utf-8-sig", newline="") as fp:
        rows = list(csv.reader(fp))
    assert rows[0][-1] == "heater_soak" and all(len(r) == len(rows[0]) for r in rows)

    wb = load_workbook(XLSX_PATH)
    ws, desc = wb["Recipe"], wb["설명"]
    cols, last = _tpl_cols(ws)
    c_soak, c_ramp = cols["heater_soak"], cols["heater_ramp"]
    assert c_soak == ws.max_column and ws.cell(1, c_soak).comment.text.startswith("단위: min\n")
    for r in (1, 2, last):                                                      # 서식은 옆 열(heater_ramp)과 같게
        assert tuple(ws.cell(r, c_soak)._style) == tuple(ws.cell(r, c_ramp)._style)
    dvs = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}

    def _dv(name):
        L = ws.cell(1, cols[name]).column_letter
        return dvs[f"{L}2:{L}{last}"]
    dv = _dv("heater_soak")
    assert (dv.type, dv.formula1, float(dv.formula2)) == ("decimal", "0", cr.HEATER_SOAK_MAX_MIN)
    assert dv.promptTitle == "heater_soak [min]"
    drow = {desc.cell(r, 1).value: r for r in range(5, desc.max_row + 1)}
    assert [desc.cell(drow["heater_soak"], c).value for c in (1, 2)] == ["heater_soak", "min"]

    # 설명 문구가 코드·설정 기본값과 맞는지(예전: 목표 1~500, 램프 기본 12, '60초 머물면 준비 완료')
    lim = min(cfgc.HEATER_MAX_TEMP, cfgc.HEATER_SV_LIMIT_C)
    assert float(_dv("heater_temp").formula2) == lim and f"1~{lim:g}" in _dv("heater_temp").prompt
    assert f"HEATER_RAMP_RATE_C_PER_MIN, {cfgc.HEATER_RAMP_RATE_C_PER_MIN:g})" in _dv("heater_ramp").prompt
    assert (f"±{cfgc.HEATER_SOAK_TOLERANCE:g}°C 안에 {cfgc.HEATER_SOAK_TIME_SEC:g}초 머물면 도달"
            in desc.cell(drow["heater_temp"], 4).value)
    assert "다시 세지 않음" in dv.prompt
    texts = [ws.cell(1, c).comment.text for c in range(1, ws.max_column + 1)]
    texts += [str(d.prompt) for d in dvs.values()] + [str(d.error) for d in dvs.values()]
    texts += [str(desc.cell(r, c).value) for r in range(1, desc.max_row + 1) for c in range(1, 5)]
    for old in ("1~500", "RATE_C_PER_MIN, 12)", "준비 완료로 보고", "'예열 대기'"):
        assert not [t for t in texts if old in t], old


def test_template_rows_load_and_normalize():
    """엑셀 템플릿을 로더처럼 읽어(첫 시트, 1행 머리글) 전수 검증·정규화 — 히터 열 값이 정규화 결과와 같다(예시 행은 사용자가 고칠 수 있음)."""
    async def go():
        rt, _pc, _rec = _rt()
        wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
        try:
            it = wb.worksheets[0].iter_rows(values_only=True)
            head = [str(h).strip() if h is not None else "" for h in next(it)]
            rows = []
            for vals in it:
                row = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, vals) if h}
                if any(row.values()):
                    rows.append(row)
        finally:
            wb.close()
        assert "heater_soak" in head and rows
        assert rt._validate_raw_rows(rows) == []
        for row in rows:
            if cr.eta_delay_seconds(row) > 0:
                continue
            n = rt._normalize_params_for_process(row)
            on = row.get("use_heater", "").upper() in ("T", "TRUE", "1")
            assert n["use_heater"] is on
            want = float(row["heater_soak"]) if (on and row.get("heater_soak")) else None
            assert n["heater_soak"] == want
    asyncio.run(go())


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
