# -*- coding: utf-8 -*-
"""CH2 히터를 공정 기록에 넣기 — Calib CSV(Ch2_log.csv)·공정 엑셀(CH2.xlsx) 히터 열, 공정 로그(.txt) 상태 줄·요약.
표본은 메인 공정 구간(폴링 중 + 셔터 열림)만. CH1 파일은 그대로. 기존 파일은 머리글만 늘린다(데이터 행 보존)."""
import os
import sys
import csv
import asyncio
import contextlib

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _ROOT)
sys.path.insert(0, _HERE)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
with contextlib.suppress(Exception):
    sys.stdout.reconfigure(encoding="utf-8")

import re                                                         # noqa: E402
import pytest                                                     # noqa: E402
from openpyxl import load_workbook                                # noqa: E402
from controller import data_logger as DLM                         # noqa: E402
from controller.data_logger import DataLogger, HEATER_COLUMNS     # noqa: E402
from util import gdrive_logger as G                               # noqa: E402
from runtime.chamber_runtime import ChamberRuntime                # noqa: E402
import heater_ladder_sim as sim                                   # noqa: E402


def _dl(ch, tmp_path, logs=None):
    return DataLogger(ch=ch, csv_dir=tmp_path, log_func=(logs.append if logs is not None else None))


def _rows(path):
    with open(path, encoding="utf-8-sig", newline="") as fp:
        return list(csv.reader(fp))


def _feed(dl):
    """셔터 닫힘(버림) → 열림 3표본(TC 이상 1개 포함) → 닫힘(버림)."""
    dl.start_new_log_session({"use_ms": True, "process_name": "T1"})
    dl.log_heater(100.0, 250.0, 600, True)
    dl.notify_shutter_open()
    dl.log_heater(249.0, 250.0, 468, True)
    dl.log_heater(251.0, 250.0, 472, True)
    dl.log_heater(None, 250.0, 470, True)
    dl.notify_shutter_close()
    dl.log_heater(300.0, 250.0, 900, True)


# ───────────────────────── Calib CSV ─────────────────────────
def test_header_ch2_only(tmp_path):
    h1, h2 = _dl(1, tmp_path).header, _dl(2, tmp_path).header
    assert len(h1) == 36 and not any(c.startswith("Heater") for c in h1)
    assert h2[:36] == h1 and h2[36:] == HEATER_COLUMNS


def test_existing_csv_header_upgraded_rows_kept(tmp_path):
    old = _dl(1, tmp_path).header
    path = tmp_path / "Ch2_log.csv"
    with open(path, "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.DictWriter(fp, fieldnames=old)
        w.writeheader()
        w.writerow({h: "" for h in old} | {"Timestamp": "2026-10-13 14:34:19", "Process Name": "A"})
    dl = _dl(2, tmp_path)
    rows = _rows(path)
    assert rows[0] == dl.header and dl._fields == dl.header
    assert rows[1][:2] == ["2026-10-13 14:34:19", "A"] and rows[1][36:] == [""] * 5


def test_upgrade_failure_keeps_old_column_order(tmp_path, monkeypatch):
    """머리글 갱신(임시 파일→교체)이 실패하면 기존 열 순서로 쓴다 — 열이 어긋나지 않게(히터 값은 버림)."""
    old = _dl(1, tmp_path).header
    path = tmp_path / "Ch2_log.csv"
    with open(path, "w", encoding="utf-8-sig", newline="") as fp:
        csv.DictWriter(fp, fieldnames=old).writeheader()
    monkeypatch.setattr(DLM.os, "replace", lambda *a, **k: (_ for _ in ()).throw(PermissionError("locked")))
    logs = []
    dl = _dl(2, tmp_path, logs)
    assert dl._fields == old and any("머리글 갱신 실패" in m for m in logs)
    _feed(dl)
    dl.finalize_and_write_log(True)
    dl._nas_app.close()
    rows = _rows(path)
    assert rows[0] == old and len(rows[1]) == 36 and rows[1][1] == "T1"


def test_ch2_row_heater_values(tmp_path):
    dl = _dl(2, tmp_path)
    _feed(dl)
    s = dl.heater_stats()
    assert (s["sv"], s["pv_avg"], s["pv_min"], s["pv_max"], s["mv_avg"], s["n"]) == (250.0, 250.0, 249.0, 251.0, 470.0, 2)
    assert dl.heater_summary_text() == ("히터 요약(메인 공정 구간) — 목표 250.0°C · PV 평균 250.0°C "
                                        "(최소 249.0 · 최대 251.0) · MV 평균 470 · 표본 2개")
    dl.finalize_and_write_log(True)
    dl._nas_app.close()
    rows = _rows(tmp_path / "Ch2_log.csv")
    row = dict(zip(rows[0], rows[1]))
    assert [row[c] for c in HEATER_COLUMNS] == ["250.0", "250.0", "249.0", "251.0", "470"]


def test_heater_off_only_pv(tmp_path):
    """히터는 꺼져 있고 TC 만 읽힌 공정 — 목표·MV 는 공란, PV 는 기록."""
    dl = _dl(2, tmp_path)
    dl.start_new_log_session({"use_ms": False, "process_name": "T2"})
    for pv in (45.0, 47.0):
        dl.log_heater(pv, 0.0, 0, False)
    assert dl.heater_summary_text().startswith("히터 요약(메인 공정 구간) — 히터 꺼짐 · PV 평균 46.0°C")
    dl.finalize_and_write_log(True)
    dl._nas_app.close()
    rows = _rows(tmp_path / "Ch2_log.csv")
    row = dict(zip(rows[0], rows[1]))
    assert [row[c] for c in HEATER_COLUMNS] == ["", "46.0", "45.0", "47.0", ""]


def test_no_samples_blank_and_ch1_untouched(tmp_path):
    dl2 = _dl(2, tmp_path)
    dl2.start_new_log_session({"use_ms": True})
    assert dl2.heater_stats() is None and dl2.heater_summary_text() == ""
    dl2.finalize_and_write_log(True)
    dl2._nas_app.close()
    row = dict(zip(*_rows(tmp_path / "Ch2_log.csv")[:2]))
    assert [row[c] for c in HEATER_COLUMNS] == [""] * 5
    dl1 = _dl(1, tmp_path)
    dl1.start_new_log_session({"use_ms": True})
    dl1.finalize_and_write_log(True)
    dl1._nas_app.close()
    assert len(_rows(tmp_path / "Ch1_log.csv")[0]) == 36


# ───────────────────────── 공정 엑셀 ─────────────────────────
def _save(ch, dl, tmp_path):
    rows = G._build_data(ch, dl, "", "", "", None)
    G._save_sync(ch, rows, tmp_path, 5, 20.0, False, "", tmp_path / "fb", "Main Process", *G._main_cols(ch))


def test_excel_new_ch2_file_has_heater_group(tmp_path):
    dl = _dl(2, tmp_path)
    _feed(dl)
    _save(2, dl, tmp_path)
    ws = load_workbook(tmp_path / "CH2.xlsx")["Main Process"]
    names = [ws.cell(2, c).value for c in range(1, 44)]
    assert names[:38] == [c[0] for c in G._COLS_MAIN] and names[38:] == ["SP Temp", "Avg Temp", "Min Temp", "Max Temp", "Avg DAC"]
    assert ws["AM1"].value == "Heater" and "AM1:AQ1" in {str(r) for r in ws.merged_cells.ranges}
    assert [ws.cell(4, c).value for c in range(39, 44)] == [250.0, 250.0, 249.0, 251.0, 470.0]


def test_excel_existing_ch2_file_gets_heater_headers(tmp_path):
    """옛 배치(38열)로 만든 CH2.xlsx — 기존 머리글·행은 그대로, 뒤에 히터 머리글을 붙이고 새 행에 값."""
    path = tmp_path / "CH2.xlsx"
    wb, ws = G._get_or_create_wb(path, "Main Process", G._COLS_MAIN, G._GRP_COLS_MAIN, G._GRP_STARTS_MAIN)
    old_row = {"process_name": "OLD", "power_source": "DC", "soft_arc": 3, "hard_arc": 4}
    G._write_data_row(ws, 4, old_row, 5, 20.0, G._COLS_MAIN)
    wb.save(path)
    dl = _dl(2, tmp_path)
    _feed(dl)
    _save(2, dl, tmp_path)
    ws = load_workbook(path)["Main Process"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A1:M1", "N1:AL1", "AM1:AQ1"} <= merged and ws["AM1"].value == "Heater"
    assert [ws.cell(2, c).value for c in (37, 38, 39, 43)] == ["Soft Arc", "Hard Arc", "SP Temp", "Avg DAC"]
    assert [ws.cell(3, c).value for c in (39, 43)] == ["°C", "count"]
    assert ws["C4"].value == "OLD" and ws["AK4"].value == 3 and ws["AL4"].value == 4
    assert [ws.cell(4, c).value for c in range(39, 44)] == [None] * 5
    assert ws["C5"].value == "T1" and [ws.cell(5, c).value for c in range(39, 44)] == [250.0, 250.0, 249.0, 251.0, 470.0]


def test_excel_ch1_unchanged_and_mismatched_header_left_alone(tmp_path):
    dl = _dl(1, tmp_path)
    dl.start_new_log_session({"use_ms": True, "process_name": "C1"})
    _save(1, dl, tmp_path)
    ws = load_workbook(tmp_path / "CH1.xlsx")["Main Process"]
    assert ws.max_column == 38 and ws.cell(2, 39).value is None
    wb = load_workbook(tmp_path / "CH1.xlsx")
    ws = wb["Main Process"]
    ws.cell(2, 5).value = "바뀐 머리글"                      # 앞부분이 다른 파일 — 새 열을 만들지 않는다
    assert G._extend_headers(ws, *G._main_cols(2)) == 0 and ws.cell(2, 39).value is None


# ───────────────────────── 공정 로그(.txt) ─────────────────────────
def _rt():
    """런타임 최소 인스턴스 + 실제 공정 컨트롤러. 실행 중인 루프 안에서 부른다."""
    rt = ChamberRuntime.__new__(ChamberRuntime)
    rt.ch = 2
    rt.logs = []
    rt.append_log = lambda s, m: rt.logs.append((s, m))
    rt._get_text = lambda name: ""
    rt.supports_rf_pulse = rt.supports_dc_pulse = rt.supports_dc_cont = rt.supports_rf_cont = True
    rt.supports_dc_cont2 = False
    rt._loop = asyncio.get_running_loop()
    rt._ensure_task_alive = lambda *a, **k: None
    rt._heater = None
    rt._heater_gen = 0
    rt._heater_job = None
    rt._heater_watch = False
    rt._heater_rate = None
    rt._runner_queue_mode = False
    rt.process_queue = []
    rt.current_process_index = -1
    rt._bind_process_controller()
    return rt, rt.process_controller


def _status_lines(rt):
    return [m for s, m in rt.logs if s == "HEATER" and m.startswith("히터 상태 —")]


def test_runtime_status_lines_and_samples(tmp_path):
    async def go():
        rt, pc = _rt()
        svc, plc, clk, logs, chat = sim.make()
        rt.attach_heater(svc)
        dl = _dl(2, tmp_path)
        rt.data_logger = dl
        dl.start_new_log_session({"use_ms": True})
        dl.notify_shutter_open()
        await sim.run(svc, plc, clk, 5)
        ok, why = await svc.turn_on(150.0)
        assert ok, why
        rt._log_file_path = "run.txt"
        await sim.run(svc, plc, clk, 30)
        assert _status_lines(rt) == [] and dl.heater_stats() is None          # 공정이 안 돌면 기록 없음
        pc.is_running = True
        await sim.run(svc, plc, clk, 60)
        lines = _status_lines(rt)
        assert 5 <= len(lines) <= 7 and dl.heater_stats() is None             # 10초마다 한 줄, 폴링 전이라 표본 없음
        assert re.fullmatch(r"히터 상태 — PV \d+\.\d°C · SV \d+\.\d°C · 목표 150\.0°C · MV \d+/520 · "
                            r"(램프|제어|유지\(DAC 상한 \d+\)) · 수동", lines[-1]), lines[-1]
        rt._proc_polling_active = True
        await sim.run(svc, plc, clk, 30)
        n = int(dl.heater_stats()["n"])
        assert 28 <= n <= 31 and dl.heater_stats()["sv"] == 150.0
        rt._log_file_path = None                                              # 공정 로그가 없으면 줄을 남기지 않는다
        before = len(_status_lines(rt))
        await sim.run(svc, plc, clk, 30)
        assert len(_status_lines(rt)) == before
        plc.connected = False                                                 # 응답 없음(stale) — 표본 멈춤
        await sim.run(svc, plc, clk, 15)
        n1 = int(dl.heater_stats()["n"])
        await sim.run(svc, plc, clk, 10)
        assert int(dl.heater_stats()["n"]) == n1
        rt._heater_log_summary()
        assert [m for s, m in rt.logs if m.startswith("히터 요약(메인 공정 구간) — 목표 150.0°C · PV 평균")]
    asyncio.run(go())


def test_runtime_status_lines_off_when_period_zero(tmp_path):
    async def go():
        rt, pc = _rt()
        svc, plc, clk, logs, chat = sim.make(run_log_period_s=0.0)
        rt.attach_heater(svc)
        rt.data_logger = _dl(2, tmp_path)
        await sim.run(svc, plc, clk, 5)
        ok, why = await svc.turn_on(150.0)
        assert ok, why
        rt._log_file_path = "run.txt"
        pc.is_running = True
        await sim.run(svc, plc, clk, 60)
        assert _status_lines(rt) == []
    asyncio.run(go())


def test_summary_needs_attached_heater(tmp_path):
    async def go():
        rt, pc = _rt()
        dl = _dl(2, tmp_path)
        rt.data_logger = dl
        dl.start_new_log_session({"use_ms": False})
        dl.log_heater(30.0, None, None, False)
        rt._heater_log_summary()                                              # 히터 기능이 꺼져 있으면 아무것도 안 남김
        assert rt.logs == []
    asyncio.run(go())


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
