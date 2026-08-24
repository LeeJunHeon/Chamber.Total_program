# controller/chamber_runtime.py
# -*- coding: utf-8 -*-
from __future__ import annotations

import csv, asyncio, contextlib, inspect, re, traceback, os, time
import json
import uuid
from datetime import datetime
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from typing import Any, Callable, Coroutine, Deque, Literal, Mapping, Optional, Sequence, TypedDict, cast, Union
from collections import deque

from PySide6.QtWidgets import QMessageBox, QFileDialog, QPlainTextEdit, QDialog, QApplication
from PySide6.QtGui import QTextCursor
from PySide6.QtCore import Qt, QTimer  # ← 추가: 모달리티/속성 지정용

# ✅ Qt(C++) 객체가 이미 delete된 상태인지 체크 (Access Violation 방어)
try:
    from shiboken6 import isValid as _qt_is_valid
except Exception:
    def _qt_is_valid(obj: Any) -> bool:
        return obj is not None

# 팝업 자동 닫기(5초) 유틸
from util.timed_popup import attach_autoclose

# 장비
from device.ig import AsyncIG
from device.mfc import AsyncMFC
from device.oes import OESAsync
from device.rga import RGAWorkerClient
from device.dc_power import DCPowerAsync
from device.rf_power import RFPowerAsync
from device.rf_pulse import RFPulseAsync
from device.dc_pulse import AsyncDCPulse

# Google Drive 로그 (선택적 import — 없어도 동작)
try:
    from util.gdrive_logger import save_process_log as _gdrive_save
    _GDRIVE_OK = True
except ImportError:
    _GDRIVE_OK = False

# 그래프/로거/알림
from controller.graph_controller import GraphController
from controller.data_logger import DataLogger
from controller.chat_notifier import ChatNotifier
from controller.process_monitor import ProcessMonitor
from util.log_hub import SessionTextAppender

# ⬇️ 추가: 전역 런타임 상태 레지스트리
from controller.runtime_state import runtime_state

# 공정 컨트롤러(기존 CH2) + CH1은 별도 모듈이 있으면 사용, 없으면 CH2를 공용으로
from controller.process_controller import ProcessController

# ---- 타입 (main.py의 정의를 최소 필요만 가져와 복제) -------------------------
RawParams = TypedDict('RawParams', {
    'Process_name': str,
    'process_note': str,
    'base_pressure': float | str,
    'working_pressure': float | str,
    'process_time': float | str,
    'shutter_delay': float | str,
    'dep_rate': float | str | None,
    'thickness': float | str | None,
    'integration_time': int | str,
    'Ar': Literal['T','F'] | bool,
    'O2': Literal['T','F'] | bool,
    'N2': Literal['T','F'] | bool,
    'Ar_flow': float | str,
    'O2_flow': float | str,
    'N2_flow': float | str,
    'use_dc_power': Literal['T','F'] | bool,
    'use_dc_power2': Literal['T','F'] | bool,
    'use_rf_power': Literal['T','F'] | bool,
    'dc_power': float | str,
    'dc_power2': float | str,
    'rf_power': float | str,

    # 🔥 펄스 완전 분리(레거시 키 전부 제거)
    'use_dc_pulse': Literal['T','F'] | bool,
    'dc_pulse_power': float | str,
    'dc_pulse_freq': int | str | None,
    'dc_pulse_duty_cycle': int | str | None,

    'use_rf_pulse': Literal['T','F'] | bool,
    'rf_pulse_power': float | str,
    'rf_pulse_freq': int | str | None,
    'rf_pulse_duty_cycle': int | str | None,

    'gun1': Literal['T','F'] | bool,
    'gun2': Literal['T','F'] | bool,
    'gun3': Literal['T','F'] | bool,
    'main_shutter': Literal['T','F'] | bool,
    'G1 Target': str,
    'G2 Target': str,
    'G3 Target': str,
    'power_select': Literal['T','F'] | bool,
}, total=False)

NormParams = TypedDict('NormParams', {
    'base_pressure': float,
    'working_pressure': float,
    'process_time': float,
    'shutter_delay': float,
    'dep_rate': float | None,
    'thickness': float | None,
    'integration_time': int,
    'use_ar': bool, 'use_o2': bool, 'use_n2': bool,
    'ar_flow': float, 'o2_flow': float, 'n2_flow': float,
    'use_dc_power': bool, 'dc_power': float,
    'use_dc_power2': bool, 'dc_power2': float,
    'use_rf_power': bool, 'rf_power': float,

    'use_dc_pulse': bool, 'dc_pulse_power': float,
    'dc_pulse_freq': int | None, 'dc_pulse_duty': int | None,

    'use_rf_pulse': bool, 'rf_pulse_power': float,
    'rf_pulse_freq': int | None, 'rf_pulse_duty': int | None,

    'use_g1': bool, 'use_g2': bool, 'use_g3': bool, 'use_ms': bool,
    'process_note': str,
    'G1_target_name': str, 'G2_target_name': str, 'G3_target_name': str,
    'G1 Target': str, 'G2 Target': str, 'G3 Target': str,
    'use_power_select': bool,
    'chuck_position': str, # ★ CSV의 up/mid/down (또는 "")
}, total=False)

# 폴링 타깃도 명확히 분리
TargetsMap = Mapping[Literal["mfc", "dc", "dc2", "rf", "dc_pulse", "rf_pulse"], bool]

@dataclass(frozen=True)
class _RunnerCmd:
    kind: Literal["START", "START_QUEUE", "STOP", "PC_FINISHED"]
    params: NormParams | None = None
    ok: bool | None = None
    detail: dict[str, Any] | None = None
    user_initiated: bool = False

# -----------------------------------------------------------------------------


@dataclass
class _CfgAdapter:
    """config_ch1 / config_ch2 모듈을 추상화해서 접근(필수 키만)."""
    mod: Any
    ch: int  # ← 채널 번호 저장(기본 파일명 등에 사용)

    def _get(self, name: str, default=None):
        """채널별 모듈 값 우선, 없으면 config_common으로 폴백."""
        v = getattr(self.mod, name, None)
        if v is not None:
            return v
        try:
            from lib import config_common as _cc
            return getattr(_cc, name, default)
        except Exception:
            return default

    @property
    def IG_POLLING_INTERVAL_MS(self) -> int:
        return int(self._get("IG_POLLING_INTERVAL_MS", 500))

    @property
    def RGA_CSV_PATH(self) -> Path:
        p = self._get("RGA_CSV_PATH", None)

        # ✅ dict 형태면 채널 키로 선택
        if isinstance(p, dict):
            p = p.get(f"ch{self.ch}") or p.get(str(self.ch)) or p.get(self.ch)

        # ✅ 혹시 RGA_CSV_PATH가 없고, 레거시로 RGA_XLSX_PATH만 있는 경우도 처리
        if not p:
            legacy = self._get("RGA_XLSX_PATH", None)
            if isinstance(legacy, dict):
                p = legacy.get(f"ch{self.ch}")
            elif isinstance(legacy, str):
                p = legacy

        if p:
            return Path(p)
        return Path.cwd() / f"RGA_CH{self.ch}.csv"

    @property
    def RGA_NET(self) -> Mapping[str, Any]:
        # 기존 코드 호환(필요시 여전히 접근 가능)
        return self._get("RGA_NET", {}) or {}

    def rga_creds(self) -> tuple[str, str, str]:
        """
        RGA 연결 정보 반환.
        - 단일 dict: {"ip","user","password"}
        - 채널별 dict: {"ch1":{...},"ch2":{...}}
        """
        rnet = self._get("RGA_NET", {}) or {}
        if isinstance(rnet, dict) and "ip" in rnet:
            # 단일 dict 형태
            return (
                rnet.get("ip", ""),
                rnet.get("user", "admin"),
                rnet.get("password", "admin"),
            )
        # 채널별 dict 형태
        block = rnet.get(f"ch{self.ch}", {}) if isinstance(rnet, dict) else {}
        return (
            block.get("ip", ""),
            block.get("user", "admin"),
            block.get("password", "admin"),
        )
    
    @property
    def IG_TCP(self) -> tuple[str, int]:
        return (
            str(self._get("IG_TCP_HOST", "192.168.1.50")),
            int(self._get("IG_TCP_PORT", 4001 if self.ch == 1 else 4002)),
        )

    @property
    def MFC_TCP(self) -> tuple[str, int]:
        return (
            str(self._get("MFC_TCP_HOST", "192.168.1.50")),
            int(self._get("MFC_TCP_PORT", 4003 if self.ch == 1 else 4006)),
        )
    
    @property
    def DCPULSE_TCP(self) -> tuple[str, int]:
        return (
            str(self._get("DCPULSE_TCP_HOST", "192.168.1.50")),
            int(self._get("DCPULSE_TCP_PORT", 4007)),
        )

    @property
    def RFPULSE_TCP(self) -> Optional[tuple[str, int]]:
        """
        (우선) RFPULSE_TCP_HOST/RFPULSE_TCP_PORT
        (폴백) RFPULSE_PORT="ip:port" 형태
        없으면 None
        """
        host = self._get("RFPULSE_TCP_HOST", None)
        port = self._get("RFPULSE_TCP_PORT", None)
        if host is not None and port is not None:
            try:
                return (str(host), int(port))
            except Exception:
                return None

        s = self._get("RFPULSE_PORT", None)
        if not s:
            return None

        try:
            # tuple/list (host, port)도 허용
            if isinstance(s, (tuple, list)) and len(s) == 2:
                return (str(s[0]), int(s[1]))

            s = str(s).strip()
            if ":" in s:
                h, p = s.split(":", 1)
                return (h.strip(), int(p.strip()))
        except Exception:
            return None

        return None
    
class ChamberRuntime:
    """
    한 챔버 실행 단위(장치/이벤트펌프/그래프/로그/버튼 바인딩).
    - PLC는 외부에서 공유 주입
    - CH1은 건셔터 없음: PLC 콜백에서 MS/G1~G3는 무시(즉시 confirmed)
    - 파워 구성:
        * CH1: DC-Pulse
        * CH2: DC(연속) + RF-Pulse (필요 시 RF 연속도 옵션)
    """

    def __init__(
        self,
        ui: Any,
        chamber_no: int,
        prefix: str,
        loop: asyncio.AbstractEventLoop,
        plc: Any,
        chat: Optional[ChatNotifier],
        cfg: Any,
        log_dir: Path,
        *,
        mfc: Optional[AsyncMFC] = None,
        ig: Optional[AsyncIG] = None,
        supports_dc_cont: Optional[bool] = None,   # DC 연속
        supports_dc_cont2: Optional[bool] = None,  # DC 연속 2호기 (CH2 전용)
        supports_rf_cont: Optional[bool] = None,   # RF 연속
        supports_dc_pulse: Optional[bool] = None,  # DC-Pulse
        supports_rf_pulse: Optional[bool] = None,  # RF-Pulse
        owns_plc: Optional[bool] = None,   # ← 추가: PLC 로그 소유자
        on_plc_owner: Optional[Callable[[Optional[int]], None]] = None,   # ★ 추가
    ) -> None:
        self.ui = ui
        self.ch = int(chamber_no)
        self.prefix = str(prefix)
        self._loop = loop
        self.plc = plc
        self.chat = chat

        # --- 공정 모니터링 (gas/pressure 편차 → 별도 웹훅) ---
        try:
            from lib import config_local as _cfgl
            _monitor_url = getattr(_cfgl, "CHAT_WEBHOOK_MONITOR_URL", "").strip()
        except Exception:
            _monitor_url = ""
        from lib import config_common as _cfgc
        self._process_monitor = ProcessMonitor(
            ch=self.ch,
            webhook_url=_monitor_url,
            flow_tolerance=getattr(_cfgc, "PROCESS_MONITOR_TOLERANCE_FLOW_SCCM", 0.3),
            pressure_tolerance=getattr(_cfgc, "PROCESS_MONITOR_TOLERANCE_PRESSURE_MTORR", 0.3),
            log_func=self.append_log,
        )

        self.cfg = _CfgAdapter(cfg, self.ch)
        self._bg_tasks: list[asyncio.Task[Any]] = []
        self._mfc_seq_lock = asyncio.Lock()
        self._starter_threads: dict[str, asyncio.Task] = {}
        self._bg_started = False
        self._pending_device_cleanup = False
        # ✅ cleanup 중 timeout 발생 여부(정상정리 실패 → 강제 복구 승격 기준)
        self._cleanup_timed_out = False

        # ✅ 강제 복구(Force Recovery) 반복 가드
        # - cleanup timeout이 연속 발생할 때 무한 복구 루프를 막기 위해 카운터를 둔다.
        # - 성공적으로 정리되면 0으로 리셋된다.
        self._force_recover_count = 0

        self._last_polling_targets: TargetsMap | None = None
        self._last_state_text: str | None = None
        # ✅ Runner 구조에서는 delay/cooldown을 stage 코루틴에서 처리하므로 레거시 delay task를 사용하지 않는다.
        self._dc_failed_flag: bool = False     # ★ 추가
        self._dc2_failed_flag: bool = False    # ★ DC2용 (dc_power2)
        # ★ DC 실측 표시 버퍼: {유닛: (P, V, I) | None}. None이면 '-' 로 표시
        self._dc_disp: dict[int, tuple[float, float, float] | None] = {1: None, 2: None}
        self._auto_connect_enabled = True  # ← 실패시 False로 내려 자동 재연결 차단
        self._run_select: dict[str, bool] | None = None  # ← 이번 런에서 펄스 선택 상태
        self._owns_plc = bool(owns_plc if owns_plc is not None else (int(chamber_no) == 1))  # 기본 CH1
        self._notify_plc_owner = on_plc_owner 
        self._main_done_callback: Optional[Callable] = None
        self._last_running_state: Optional[bool] = None

        self._cmd_q: asyncio.Queue[_RunnerCmd] = asyncio.Queue(maxsize=200)
        self._runner_task: asyncio.Task | None = None
        self._runner_state: Literal["IDLE","PREFLIGHT","RUNNING","COOLDOWN","DELAY","CLEANUP","STOPPING"] = "IDLE"
        self._runner_stage_task: asyncio.Task | None = None
        self._runner_stage_kind: str | None = None
        self._runner_queue_mode: bool = False

        # ✅ Runner cmd 디바운스(중복/잔여 STOP이 다음 START를 끊는 레이스 차단)
        self._runner_cmd_start_enqueued: bool = False
        self._runner_cmd_stop_enqueued: bool = False

        # ✅ stage 취소(교체) 의도 플래그: Runner가 stage를 교체하기 위해 cancel 하는 경우를 구분
        self._expected_stage_cancel_task: asyncio.Task | None = None
        self._expected_stage_cancel_kind: str | None = None
    
        # ✅ 런(시작) 세대 번호: 프리플라이트/정리 레이스 방지 + watchdog 식별
        self._run_gen: int = 0
        self._active_run_gen: int = 0
    
        # ✅ Host 응답용 Future (프리플라이트가 끝나면 결과를 세팅)
        self._host_start_future: Optional[asyncio.Future] = None

        # QMessageBox 참조 저장소(비모달 유지용)
        self._msg_boxes: list[QMessageBox] = []  # ← 추가

        # ✅ 종료/정리 중에는 UI 콜백 예약을 중단하기 위한 플래그
        self._shutting_down: bool = False

        # ✅ 기본 전략: config의 SUPPORTS_*를 최우선으로, 없으면 기존 CH 기본값으로 폴백
        def _cfg_bool(*names: str):
            for n in names:
                v = self.cfg._get(n, None)
                if v is not None:
                    return bool(v)
            return None

        if supports_dc_cont is None:
            v = _cfg_bool("SUPPORTS_DC_CONT", "SUPPORTS_DC")  # 레거시 키도 허용
            supports_dc_cont = v if v is not None else (self.ch == 2)

        if supports_dc_cont2 is None:
            v = _cfg_bool("SUPPORTS_DC_CONT2", "SUPPORTS_DC2")
            supports_dc_cont2 = v if v is not None else False   # 기본 OFF

        if supports_rf_cont is None:
            v = _cfg_bool("SUPPORTS_RF_CONT")
            supports_rf_cont = v if v is not None else False

        if supports_dc_pulse is None:
            v = _cfg_bool("SUPPORTS_DC_PULSE", "SUPPORTS_DCPULSE")
            supports_dc_pulse = v if v is not None else (self.ch == 1)

        if supports_rf_pulse is None:
            v = _cfg_bool("SUPPORTS_RFPULSE", "SUPPORTS_RF_PULSE")
            supports_rf_pulse = v if v is not None else (self.ch == 2)

        self.supports_dc_cont  = bool(supports_dc_cont)
        self.supports_dc_cont2 = bool(supports_dc_cont2) and (self.ch == 2)   # DC2는 CH2 전용
        self.supports_rf_cont  = bool(supports_rf_cont)
        self.supports_dc_pulse = bool(supports_dc_pulse)
        self.supports_rf_pulse = bool(supports_rf_pulse)

        # UI 포인터
        self._w_log: QPlainTextEdit | None = self._u("logMessage_edit")
        self._w_state: QPlainTextEdit | None = self._u("processState_edit")

        # 그래프 컨트롤러
        self.graph = GraphController(self._u("rgaGraph_widget"), self._u("oesGraph_widget"))
        try:
            self.graph.reset()
        except Exception:
            self.append_log("Graph", "reset skipped (headless)")

        self._log_root = Path(log_dir)

        # ✅ config_common.py 경로 사용
        self._local_log_dir = Path(
            self.cfg._get(
                f"LOCAL_FALLBACK_CH{self.ch}_DIR",
                Path.cwd() / "Logs_LocalFallback" / f"CH{self.ch}",
            )
        )

        # ✅ CH 로그는 NAS 경로를 우선 사용
        self._log_dir = self._ensure_log_dir(self._log_root / f"CH{self.ch}")
        self._log_file_path: Path | None = None
        self._prestart_buf: Deque[str] = deque(maxlen=1000)

        # ✅ 실제 write 실패 시에만 fallback_dir 쪽에 생성/전환
        self._run_log_appender = SessionTextAppender(
            fallback_dir=self._local_log_dir,
            encoding="utf-8",
        )

        self._log_q: asyncio.Queue[str | None] = asyncio.Queue(maxsize=4096)
        self._log_writer_task: asyncio.Task | None = None

        # ✅ 로그 writer shutdown 중복 호출 방지(동시 shutdown 레이스 방지)
        self._log_shutdown_lock = asyncio.Lock()

        # ✅ 로그 파일 I/O는 이벤트루프 밖(전용 1-thread)에서만 수행
        self._log_io_exec = ThreadPoolExecutor(
            max_workers=1,
            thread_name_prefix=f"LogIO.CH{self.ch}"
        )

        # ✅ UI 로그 무한 누적 방지(프리징 완화)
        self._ui_log_buf = deque(maxlen=5000)   # UI에 쌓을 임시 버퍼(메모리 보호)
        self._ui_log_timer = None

        if self._w_log:
            # 1) UI 문서 줄 수 제한(이미 있다면 유지/조정)
            self._w_log.setMaximumBlockCount(2000)   # 또는 5000

            # 2) Undo/Redo 끄면 QPlainTextEdit 비용이 줄어듦
            with contextlib.suppress(Exception):
                self._w_log.setUndoRedoEnabled(False)

            # 3) 배치로 찍기 위한 타이머
            self._ui_log_timer = QTimer(self._w_log)
            self._ui_log_timer.setInterval(100)      # 100ms마다 한번만 UI 업데이트
            self._ui_log_timer.timeout.connect(self._flush_ui_log_to_ui)
            self._ui_log_timer.start()

        # 데이터 로거 (Sputter Calib CSV) - CH 로그로 로그를 흘려보내도록 콜백 전달
        self.data_logger = DataLogger(
            ch=self.ch,
            csv_dir=Path(str(self.cfg._get(
                "SPUTTER_CALIB_DB_DIR",
                r"G:\공유 드라이브\VanaM_Sputter\Sputter\Calib\Database",
            ))),
            log_func=lambda msg: self.append_log("CSV", msg),
        )

        # 장치 인스턴스(각 챔버 독립)
        mfc_host, mfc_port = self.cfg.MFC_TCP
        ig_host,  ig_port  = self.cfg.IG_TCP

        # 설정에서 채널별 스케일 정보를 불러와 주입
        scale_map = self.cfg._get("MFC_SCALE_FACTORS", {1: 1.0, 2: 1.0, 3: 1.0})
        if not isinstance(scale_map, dict):
            scale_map = {1: 1.0, 2: 1.0, 3: 1.0}

        # MFC/IG를 외부에서 주입하면 그대로 사용하고, 없으면 기존 방식대로 생성
        self.mfc = mfc or AsyncMFC(
            host=mfc_host, port=mfc_port, enable_verify=False, enable_stabilization=True,
            # ★ 챔버별 스케일을 드라이버에 주입
            scale_factors=scale_map,  # ✅ CH별 MFC 스케일 전달
            cfg=self.cfg.mod,  # ✅ 채널 config 주입(=UI Apply 반영/CH별 오버라이드 반영)
        )
        self.ig  = ig or AsyncIG(host=ig_host, port=ig_port, cfg=self.cfg.mod)

        # OES 인스턴스 생성 시 현재 챔버 번호에 따라 USB 채널을 명시적으로 매핑한다.
        # CH1 → USB0, CH2 → USB1. OESAsync 내부 기본 동작도 동일하지만 명확성을 위해 전달한다.
        if self.ch == 1:
            _usb_index = int(self.cfg._get("CHAMBER_OES_USB_INDEX_CH1", 0))
        else:
            _usb_index = int(self.cfg._get("CHAMBER_OES_USB_INDEX_CH2", 1))

        self.oes = OESAsync(chamber=self.ch, usb_index=_usb_index)
        self._oes_initialized = False  # ✅ 워커 init(장치 스캔) 1회만 수행하도록 캐시

        # RGA: worker client (메인에서 srsinst import 안함)
        self.rga = None
        try:
            # logger는 선택사항인데 ChamberRuntime에는 self.logger가 없으니 전달하지 않는다.
            timeout_s = float(self.cfg._get("RGA_WORKER_TIMEOUT_S", 60.0))
            self.rga = RGAWorkerClient(ch=self.ch, logger=None, default_timeout_s=timeout_s)
        except Exception as e:
            # 기존 로그 시스템(append_log)로만 남긴다.
            self.append_log(f"RGA{self.ch}", f"RGAWorkerClient init failed: {e!r} (RGA disabled)")
            self.rga = None

        # 펄스 파워(완전 분리)
        # - on_telemetry를 DataLogger로 직결(있으면 log_dcpulse_power, 없으면 log_dc_power 폴백)
        # - 생성 시점에 host/port도 지정
        if self.supports_dc_pulse:
            _cb = getattr(self.data_logger, "log_dcpulse_power", None)
            if not callable(_cb):
                def _cb(p, v, i):
                    try:
                        self.data_logger.log_dc_power(float(p), float(v), float(i))
                    except Exception:
                        pass
            # host/port는 드라이버가 cfg에서 스스로 읽게 둠(override 고정 방지)
            self.dc_pulse = AsyncDCPulse(on_telemetry=_cb, cfg=self.cfg.mod)
        else:
            self.dc_pulse = None

        # RF Pulse (공유 장비)
        if self.supports_rf_pulse:
            # ✅ cfg 주입: UI Apply로 바뀐 값(타임아웃/간격/백오프 등)을 드라이버가 읽을 수 있게
            self.rf_pulse = RFPulseAsync(cfg=self.cfg.mod)

            # ✅ config 값 런타임 캐시 재반영(드라이버가 __init__에 캐시하는 값이 있으면 특히 중요)
            with contextlib.suppress(Exception):
                if hasattr(self.rf_pulse, "reload_runtime_cfg"):
                    self.rf_pulse.reload_runtime_cfg()

            # ✅ 프리플라이트 전에 endpoint를 먼저 세팅
            # - (우선) RFPULSE_TCP_HOST/PORT
            # - (폴백) RFPULSE_PORT="ip:port"
            rf_tcp = getattr(self.cfg, "RFPULSE_TCP", None)

            # ✅ 레거시 호환: RFPULSE_TCP_HOST/PORT가 없고, RFPULSE_PORT="ip:port"만 있을 때만 override를 건다
            host_defined = self.cfg._get("RFPULSE_TCP_HOST", None)
            port_defined = self.cfg._get("RFPULSE_TCP_PORT", None)

            if (host_defined is None or port_defined is None) and rf_tcp and hasattr(self.rf_pulse, "set_endpoint"):
                host, port = rf_tcp
                with contextlib.suppress(Exception):
                    self.rf_pulse.set_endpoint(host, port)

        else:
            self.rf_pulse = None

        # 연속 파워
        self.dc_power = None
        if self.supports_dc_cont:
            async def _dc_send(power: float):
                # 연속 제어 루프에서는 SET을 건드리지 않는다 → WRITE만 수행
                await self.plc.power_write(power, family="DCV", write_idx=0)

            async def _dc_send_unverified(power: float):
                # no-reply: WRITE만
                await self.plc.power_write(power, family="DCV", write_idx=0)

            async def _dc_read():
                try:
                    P, V, I = await self.plc.power_read(family="DCV", v_idx=0, i_idx=1)
                    return (P, V, I)
                except Exception as e:
                    self.append_log("DCpower", f"read failed: {e!r}")

            # ⬇️ 추가: SET 코일 ON/OFF 콜백
            async def _dc_toggle_enable(on: bool):
                await self.plc.power_enable(on, family="DCV", set_idx=0)

            self.dc_power = DCPowerAsync(
                send_dc_power=_dc_send,
                send_dc_power_unverified=_dc_send_unverified,
                request_status_read=_dc_read,
                toggle_enable=_dc_toggle_enable,   # ← 추가
                name="DC1",
            )

        # 연속 파워 2호기 (CH2 전용)
        #  - SET: DCV_SET_3(M00053), WRITE: DCV_WRITE_3(D00007)
        #  - READ: DCV_READ_6/7(D00010/D00011) = V/I
        #  - 스케일은 1호기와 동일 기종이므로 cfg.dc_* 기본값 그대로 사용
        self.dc_power2 = None
        if self.supports_dc_cont2 and self.plc:
            async def _dc2_send(power: float):
                await self.plc.power_write(power, family="DCV", write_idx=3)

            async def _dc2_send_unverified(power: float):
                await self.plc.power_write(power, family="DCV", write_idx=3)

            async def _dc2_read():
                try:
                    P, V, I = await self.plc.power_read(family="DCV", v_idx=6, i_idx=7)
                    return (P, V, I)
                except Exception as e:
                    self.append_log("DCpower2", f"read failed: {e!r}")

            async def _dc2_toggle_enable(on: bool):
                await self.plc.power_enable(on, family="DCV", set_idx=3)

            self.dc_power2 = DCPowerAsync(
                send_dc_power=_dc2_send,
                send_dc_power_unverified=_dc2_send_unverified,
                request_status_read=_dc2_read,
                toggle_enable=_dc2_toggle_enable,
                name="DC2",
            )

        self.rf_power = None
        if self.supports_rf_cont and self.plc:
            # (CH2 전용) RF 연속 제어 — RF channel 2 사용
            # - SET: DCV_SET_2, WRITE: DCV ch2, READ: DCV_READ_4/5 (FWD/REF)
            async def _rf_send(power: float):
                # SET 래치는 RFPowerAsync의 toggle_enable(True)에서 한 번만 걸어도 됨
                # 여기서는 중복 SET 방지하려면 ensure_set=False로 호출
                await self.plc.rf_apply(float(power), ensure_set=False, rf_ch=2)

            async def _rf_send_unverified(power: float):
                await self.plc.rf_write_w(float(power), rf_ch=2)

            async def _rf_request_read():
                try:
                    rf_zeroing = bool(self.cfg._get("CHAMBER_RF_CONT_ZEROING", False))
                    return await self.plc.rf_read_fwd_ref(rf_ch=2, zeroing=rf_zeroing)
                except Exception as e:
                    self.append_log("RF", f"read failed: {e!r}")
                    return None

            async def _rf_toggle_enable(on: bool):
                await self.plc.rf_enable(bool(on), rf_ch=2)

            self.rf_power = RFPowerAsync(
                send_rf_power=_rf_send,
                send_rf_power_unverified=_rf_send_unverified,
                request_status_read=_rf_request_read,
                toggle_enable=_rf_toggle_enable,
                poll_interval_ms=int(self.cfg._get("CHAMBER_RF_CONT_POLL_INTERVAL_MS", 1000)),
                rampdown_interval_ms=int(self.cfg._get("CHAMBER_RF_CONT_RAMPDOWN_INTERVAL_MS", 50)),
                direct_mode=bool(self.cfg._get("CHAMBER_RF_CONT_DIRECT_MODE", False)),
                write_inv_a=float(self.cfg._get("CHAMBER_RF_CONT_WRITE_INV_A", 1.6546)),
                write_inv_b=float(self.cfg._get("CHAMBER_RF_CONT_WRITE_INV_B", 2.6323)),
            )

        # Gun Target 자동 불러오기 (재고 DB 연동)
        asyncio.ensure_future(self._load_gun_targets())

        # === ProcessController 바인딩 ===
        self._bind_process_controller()

        # === UI 버튼 바인딩 (자기 챔버 것만) ===
        self._connect_my_buttons()

        # === 백그라운드 워치독/이벤트펌프 준비는 최초 Start 때 올림 ===
        self._on_process_status_changed(False)

    # ------------------------------------------------------------------
    # 공정 컨트롤러 바인딩
    # ------------------------------------------------------------------
    def reload_runtime_cfg(self) -> None:
        """
        원칙
        - 기존 공정 로직(Runner / preflight / cleanup / polling 흐름)은 건드리지 않는다.
        - 실행 중/정리 중에는 위험한 값(지원 플래그/장비 생성 의존 값)은 바꾸지 않는다.
        - config 값이 비어 있으면 기존처럼 config_ch* -> config_common -> default 로 폴백한다.
        """
        try:
            cfg_mod = getattr(getattr(self, "cfg", None), "mod", None)
            if cfg_mod is None:
                return

            # ✅ 항상 최신 모듈 값을 보도록 adapter만 다시 감싼다.
            self.cfg = _CfgAdapter(cfg_mod, self.ch)

            # ------------------------------------------------------------------
            # 1) 항상 안전한 경로/캐시만 먼저 갱신
            # ------------------------------------------------------------------
            self._local_log_dir = Path(
                self.cfg._get(
                    f"LOCAL_FALLBACK_CH{self.ch}_DIR",
                    Path.cwd() / "Logs_LocalFallback" / f"CH{self.ch}",
                )
            )

            # SessionTextAppender fallback_dir 갱신 (가능한 API만 안전 호출)
            with contextlib.suppress(Exception):
                if hasattr(self._run_log_appender, "fallback_dir"):
                    self._run_log_appender.fallback_dir = self._local_log_dir
                elif hasattr(self._run_log_appender, "set_fallback_dir"):
                    self._run_log_appender.set_fallback_dir(self._local_log_dir)

            # DataLogger 저장 경로 갱신
            with contextlib.suppress(Exception):
                if getattr(self, "data_logger", None) is not None:
                    self.data_logger.csv_dir = Path(
                        str(
                            self.cfg._get(
                                "SPUTTER_CALIB_DB_DIR",
                                r"G:\공유 드라이브\VanaM_Sputter\Sputter\Calib\Database",
                            )
                        )
                    )

            # ------------------------------------------------------------------
            # 2) 실행 중/정리 중이면 여기서 중단
            #    - 기존 plasma cleaning runtime 때처럼 공정 흐름을 흔들지 않기 위함
            # ------------------------------------------------------------------
            busy = bool(getattr(self.process_controller, "is_running", False)) or (
                str(getattr(self, "_runner_state", "IDLE")).upper() != "IDLE"
            )

            if busy:
                self.append_log(
                    "Config",
                    "Apply(Runtime): 공정/정리 진행 중이므로 ChamberRuntime 고정 캐시는 다음 공정부터 반영됩니다."
                )
                return

            # ------------------------------------------------------------------
            # 3) 비실행 중일 때만 support 플래그 재평가
            #    - 단, 이미 생성된 장비 객체가 없는 항목은 즉시 ON 반영하지 않는다.
            #      (예: rf_pulse 객체가 없는데 SUPPORTS_RFPULSE만 True로 바꾸는 경우)
            # ------------------------------------------------------------------
            def _cfg_bool(*names: str):
                for n in names:
                    v = self.cfg._get(n, None)
                    if v is not None:
                        return bool(v)
                return None

            desired_dc_cont = self.supports_dc_cont
            v = _cfg_bool("SUPPORTS_DC_CONT", "SUPPORTS_DC")
            if v is not None:
                desired_dc_cont = bool(v)

            desired_dc_cont2 = self.supports_dc_cont2
            v = _cfg_bool("SUPPORTS_DC_CONT2", "SUPPORTS_DC2")
            if v is not None:
                desired_dc_cont2 = bool(v)

            desired_rf_cont = self.supports_rf_cont
            v = _cfg_bool("SUPPORTS_RF_CONT")
            if v is not None:
                desired_rf_cont = bool(v)

            desired_dc_pulse = self.supports_dc_pulse
            v = _cfg_bool("SUPPORTS_DC_PULSE", "SUPPORTS_DCPULSE")
            if v is not None:
                desired_dc_pulse = bool(v)

            desired_rf_pulse = self.supports_rf_pulse
            v = _cfg_bool("SUPPORTS_RFPULSE", "SUPPORTS_RF_PULSE")
            if v is not None:
                desired_rf_pulse = bool(v)

            # ✅ 실제 즉시 반영 가능한 값만 적용
            #    - 객체가 이미 존재하는 장치만 True 반영 가능
            #    - 객체가 없는 장치를 True로 바꾸는 건 런타임 재생성/프로그램 재시작이 필요
            effective_dc_cont = bool(desired_dc_cont and (self.dc_power is not None))
            effective_dc_cont2 = bool(desired_dc_cont2 and (self.dc_power2 is not None))
            effective_rf_cont = bool(desired_rf_cont and (self.rf_power is not None))
            effective_dc_pulse = bool(desired_dc_pulse and (self.dc_pulse is not None))
            effective_rf_pulse = bool(desired_rf_pulse and (self.rf_pulse is not None))

            if desired_dc_cont and self.dc_power is None:
                self.append_log("Config", "DC 연속파 지원 ON 요청은 현재 runtime에 dc_power 객체가 없어 즉시 반영하지 않습니다. 프로그램 재시작 후 반영됩니다.")
            if desired_dc_cont2 and self.dc_power2 is None:
                self.append_log("Config", "DC 연속파 2호기 지원 ON 요청은 현재 runtime에 dc_power2 객체가 없어 즉시 반영하지 않습니다. 프로그램 재시작 후 반영됩니다.")
            if desired_rf_cont and self.rf_power is None:
                self.append_log("Config", "RF 연속파 지원 ON 요청은 현재 runtime에 rf_power 객체가 없어 즉시 반영하지 않습니다. 프로그램 재시작 후 반영됩니다.")
            if desired_dc_pulse and self.dc_pulse is None:
                self.append_log("Config", "DC-Pulse 지원 ON 요청은 현재 runtime에 dc_pulse 객체가 없어 즉시 반영하지 않습니다. 프로그램 재시작 후 반영됩니다.")
            if desired_rf_pulse and self.rf_pulse is None:
                self.append_log("Config", "RF-Pulse 지원 ON 요청은 현재 runtime에 rf_pulse 객체가 없어 즉시 반영하지 않습니다. 프로그램 재시작 후 반영됩니다.")

            self.supports_dc_cont = effective_dc_cont
            self.supports_dc_cont2 = effective_dc_cont2
            self.supports_rf_cont = effective_rf_cont
            self.supports_dc_pulse = effective_dc_pulse
            self.supports_rf_pulse = effective_rf_pulse

            # ProcessController도 동일하게 맞춘다.
            with contextlib.suppress(Exception):
                self.process_controller.supports_dc_cont = self.supports_dc_cont
                self.process_controller.supports_dc_cont2 = self.supports_dc_cont2
                self.process_controller.supports_rf_cont = self.supports_rf_cont
                self.process_controller.supports_dc_pulse = self.supports_dc_pulse
                self.process_controller.supports_rf_pulse = self.supports_rf_pulse

            # ------------------------------------------------------------------
            # 4) 참고 로그
            #    - OES USB index 같은 값은 OES 객체를 새로 만들지 않는 이상 hot reload 대상이 아님
            # ------------------------------------------------------------------
            self.append_log(
                "Config",
                "ChamberRuntime 설정 캐시 갱신 완료 (실행 중 위험한 값은 건드리지 않음, 일부 장치 생성 의존 값은 재시작 필요)"
            )

        except Exception as e:
            self.append_log("Config", f"reload_runtime_cfg 실패: {e!r}")

    def _bind_process_controller(self) -> None:
        # === 콜백 정의(PLC/MFC/파워/OES/RGA/IG) ===

        def cb_plc(cmd: str, on: Any, ch: int | None = None) -> None:
            async def run():
                raw = str(cmd)
                nname = raw.upper()
                onb = bool(on)

                # ⬇️ 추가: 요청 로그 + 총 소요시간 계측 시작
                t0 = 0.0
                try:
                    t0 = asyncio.get_running_loop().time()
                except RuntimeError:
                    pass
                self.append_log("PLC", f"[CH{self.ch}] 요청: {nname} -> {onb} (raw='{raw}', ch={self.ch})")

                try:
                    # CH1: 셔터 + N2 가스 무시
                    if self.ch == 1 and nname in ("G1", "G2", "G3", "N2"):
                        reason = "건 셔터 없음" if nname in ("G1", "G2", "G3") else "N2 라인 없음"
                        self.append_log("PLC", f"[CH1] '{nname}' 명령은 무시({reason}).")
                        self.process_controller.on_plc_confirmed(nname)
                        return

                    if nname == "MV":
                        await self.plc.write_switch(f"MAIN_{int(self.ch)}_GAS_SW", onb)
                    elif nname in ("AR", "O2", "N2", "MAIN"):
                        await self.plc.gas(int(self.ch), nname, on=onb)

                    elif nname == "MS":
                        await self.plc.main_shutter(int(self.ch), open=onb)

                        # ✅ process_monitor shutter 연동
                        with contextlib.suppress(Exception):
                            if onb:
                                self._process_monitor.notify_shutter_open()
                            else:
                                self._process_monitor.notify_shutter_close()

                        # ✅ data_logger shutter 연동
                        with contextlib.suppress(Exception):
                            if onb:
                                self.data_logger.notify_shutter_open()
                            else:
                                self.data_logger.notify_shutter_close()

                        # ✅ Main Shutter 기준 카메라 녹화 시작/정지
                        # RF 파워 사용 공정에서만 카메라 활성화
                        #   CH1: RF Pulse 사용 시
                        #   CH2: RF Power 또는 RF Pulse 사용 시
                        #   CLEANING: RF Power 사용 시
                        try:
                            recorder = getattr(self, "camera_recorder", None)
                            if recorder:
                                if onb:
                                    pc = getattr(self, "process_controller", None)
                                    params = getattr(pc, "current_params", {}) or {}
                                    use_rf       = bool(params.get("use_rf_power", False))
                                    use_rf_pulse = bool(params.get("use_rf_pulse", False))

                                    if use_rf or use_rf_pulse:
                                        recorder.set_log_callback(self._cam_log)  # ✅ 카메라 로그 → 공정 로그+화면
                                        recorder.start(f"CH{self.ch}")
                                        self.append_log("CAM", f"[CH{self.ch}] 카메라 녹화 시작 (RF={'RF' if use_rf else ''}{'Pulse' if use_rf_pulse else ''})")
                                    else:
                                        self.append_log("CAM", f"[CH{self.ch}] RF 미사용 공정 → 카메라 건너뜀")
                                else:
                                    recorder.stop()
                                    self.append_log("CAM", f"[CH{self.ch}] 카메라 녹화 정지")
                            else:
                                self.append_log("CAM", "camera_recorder 없음 (None)")
                        except Exception as e:
                            self.append_log("CAM", f"카메라 오류: {e!r}")

                    elif nname in ("G1", "G2", "G3"):
                        idx = int(nname[1])
                        await self.plc.write_switch(f"SHUTTER_{idx}_SW", onb)

                    elif nname == "SW_RF_SELECT":
                        # 일반 write 동작은 동일
                        await self.plc.write_switch("SW_RF_SELECT", onb)
                        # ✅ Start 시퀀스의 ON write 직후, 실측값을 로그용
                        #    process_params["use_power_select"] 에 기록.
                        #    - 종료 시퀀스의 OFF write 에서는 덮어쓰지 않도록
                        #      `if onb` 로 가드 (Start 시점 상태만 보존).
                        #    - PLC read 실패 시 레시피 값(start_new_log_session
                        #      에서 미리 채워진 값)을 그대로 유지하고 로그만 남김.
                        if onb:
                            try:
                                actual = bool(await self.plc.read_bit("SW_RF_SELECT"))
                                pp = getattr(self.data_logger, "process_params", None)
                                if isinstance(pp, dict):
                                    pp["use_power_select"] = actual
                                self.append_log(
                                    "PLC",
                                    f"[CH{self.ch}] SW_RF_SELECT 실측 기록(use_power_select={actual})"
                                )
                            except Exception as e:
                                # PLC 읽기 실패 → process_params 변경하지 않음
                                # (start_new_log_session 에서 채워진 레시피 값이 유지됨)
                                pp = getattr(self.data_logger, "process_params", None)
                                recipe_val = pp.get("use_power_select") if isinstance(pp, dict) else None
                                self.append_log(
                                    "PLC",
                                    f"[CH{self.ch}] SW_RF_SELECT 읽기 실패({e!r}) → "
                                    f"레시피 값 유지(use_power_select={recipe_val})"
                                )

                    else:
                        await self.plc.write_switch(raw, onb)

                    self.process_controller.on_plc_confirmed(nname)

                    # ⬇️ 추가: 완료 로그(+ 소요시간)
                    dt_ms = 0.0
                    try:
                        if t0:
                            dt_ms = (asyncio.get_running_loop().time() - t0) * 1000.0
                    except RuntimeError:
                        pass
                    self.append_log("PLC", f"[CH{self.ch}] 완료: {nname} -> {onb} ({dt_ms:.0f} ms)")

                except Exception as e:
                    # ⬇️ 추가: 실패 로그(+ 소요시간)
                    dt_ms = 0.0
                    try:
                        if t0:
                            dt_ms = (asyncio.get_running_loop().time() - t0) * 1000.0
                    except RuntimeError:
                        pass

                    self.process_controller.on_plc_failed(
                        nname,
                        e,
                        code=getattr(e, "code", None) or getattr(e, "error_code", None),
                        meta={
                            "cmd": nname,
                            "raw_cmd": raw,
                            "ch": self.ch,
                            "requested_on": onb,
                            "op": getattr(e, "op", None),
                            "addr": getattr(e, "addr", None),
                        },
                    )
                    self.append_log("PLC", f"명령 실패: {raw} -> {onb}: {e!r}")
            self._spawn_detached(run())

        def cb_mfc(cmd: str, args: Mapping[str, Any]) -> None:
            # 🔒 CH1에선 N2 가스를 완전히 무시
            gas = str(args.get("gas", "")).upper() if isinstance(args, Mapping) else ""
            if self.ch == 1 and gas == "N2":
                self.append_log("MFC", "[CH1] N2 요청 무시 (라인 없음)")
                # 프로세스 진행이 끊기지 않도록 '확인' 신호만 넘겨줌
                self.process_controller.on_mfc_confirmed(cmd)
                return
            
            self._spawn_detached(self.mfc.handle_command(cmd, args))

        def cb_dc_power(value: float):
            if not self.dc_power:
                self.append_log("DCpower", "이 챔버는 DC 연속 파워를 지원하지 않습니다.")
                return
            self._spawn_detached(self.dc_power.start_process(float(value)))

        def cb_dc_stop():
            if self.dc_power:
                self._spawn_detached(self.dc_power.cleanup())

        def cb_dc_power2(value: float):
            if not self.dc_power2:
                self.append_log("DCpower2", "이 챔버는 DC 연속 파워 2호기를 지원하지 않습니다.")
                return
            self._spawn_detached(self.dc_power2.start_process(float(value)))

        def cb_dc2_stop():
            if self.dc_power2:
                self._spawn_detached(self.dc_power2.cleanup())

        def cb_rf_power(value: float):
            if not self.rf_power:
                self.append_log("RFpower", "이 챔버는 RF 연속 파워를 지원하지 않습니다.")
                return
            self._spawn_detached(self.rf_power.start_process(float(value)))

        def cb_rf_stop():
            if self.rf_power:
                self._spawn_detached(self.rf_power.cleanup())

        def cb_dc_pulse_start(
            power: float,
            freq: Union[int, float, str, None],
            duty: Union[int, float, str, None],
        ) -> None:
            async def run():
                if not self.dc_pulse:
                    self.append_log("DCPulse", "DC-Pulse 미지원 챔버입니다."); return
                try:
                    self._ensure_background_started()
                    # (선행 단계에서 이미 연결/워치독이 올라와 있으므로 start()는 생략해도 무방)
                    ok = await self.dc_pulse.prepare_and_start(power_w=float(power), freq=freq, duty=duty)
                    if ok:
                        with contextlib.suppress(Exception):
                            pr = await self.dc_pulse.read_actual_pulse_params()
                            if pr:
                                pp = getattr(self.data_logger, "process_params", None)
                                if isinstance(pp, dict):
                                    pp["dc_pulse_freq"]       = pr["freq_khz"]
                                    pp["dc_pulse_duty_cycle"] = pr["duty_pct"]
                                    pp["dc_pulse_off_time_us"]= pr["off_time_us"]

                    if not ok:
                        self.process_controller.on_dc_pulse_failed(
                            "prepare_and_start failed",
                            meta={
                                "stage": "prepare_and_start",
                                "ch": self.ch,
                                "power": float(power),
                                "freq": freq,
                                "duty": duty,
                                "returned_ok": False,
                            },
                        )
                        return
                except Exception as e:
                    why = f"DC-Pulse start failed: {e!r}"
                    self.append_log("DCPulse", why)
                    self.process_controller.on_dc_pulse_failed(
                        e,
                        code=getattr(e, "code", None) or getattr(e, "error_code", None),
                        meta={
                            "stage": "prepare_and_start",
                            "ch": self.ch,
                            "power": float(power),
                            "freq": freq,
                            "duty": duty,
                        },
                    )
            self._spawn_detached(run())

        # ✅ Output ON 상태에서 Power setpoint만 변경
        #    - 변경 중에는 DCP polling을 잠시 멈춰 queue 혼선을 줄인다.
        #    - 실패 판정은 dc_pulse.events() -> _pump_dcpulse_events 에서 일원화한다.
        def cb_dc_pulse_set_power(power: float) -> None:
            async def run():
                if not self.dc_pulse:
                    self.append_log("DCPulse", "DC-Pulse 미지원 챔버입니다.")
                    return
                try:
                    self._ensure_background_started()

                    ok = await self.dc_pulse.set_reference_power(
                        float(power),
                        pause_polling=True,
                    )
                    if not ok:
                        self.append_log(
                            "DCPulse",
                            "set_reference_power returned False (will be handled by event pump)"
                        )

                except Exception as e:
                    why = f"DC-Pulse set_reference_power failed: {e!r}"
                    self.append_log("DCPulse", why)
                    self.process_controller.on_dc_pulse_failed(
                        e,
                        code=getattr(e, "code", None) or getattr(e, "error_code", None),
                        meta={
                            "stage": "set_reference_power",
                            "ch": self.ch,
                            "power": float(power),
                        },
                    )

            self._spawn_detached(run())

        def cb_dc_pulse_stop():
            async def run():
                if self.dc_pulse:
                    try:
                        await self.dc_pulse.output_off()
                    except Exception as e:
                        self.append_log("DCPulse", f"output_off failed: {e!r}")
                        self.process_controller.on_dc_pulse_failed(
                            e,
                            code=getattr(e, "code", None) or getattr(e, "error_code", None),
                            meta={
                                "stage": "output_off",
                                "ch": self.ch,
                            },
                        )
            self._spawn_detached(run())

        def cb_rf_pulse_start(power: float, freq: int | None, duty: int | None) -> None:
            async def run():
                if not self.rf_pulse:
                    self.append_log("RFPulse", "RF-Pulse 미지원 챔버입니다.")
                    return
                try:
                    self._ensure_background_started()
                    await self.rf_pulse.start_pulse_process(float(power), freq, duty)
                    with contextlib.suppress(Exception):
                        pr = await self.rf_pulse.read_actual_pulse_params()
                        if pr:
                            pp = getattr(self.data_logger, "process_params", None)
                            if isinstance(pp, dict):
                                pp["rf_pulse_freq"]        = pr["freq_khz"]
                                pp["rf_pulse_duty_cycle"]  = pr["duty_pct"]
                                pp["rf_pulse_off_time_us"] = pr["off_time_us"]

                except Exception as e:
                    why = f"RF-Pulse start failed: {e!r}"
                    self.append_log("RFPulse", why)
                    self.process_controller.on_rf_pulse_failed(
                        e,
                        code=getattr(e, "code", None) or getattr(e, "error_code", None),
                        meta={
                            "stage": "start_pulse_process",
                            "ch": self.ch,
                            "power": float(power),
                            "freq": freq,
                            "duty": duty,
                        },
                    )
            self._spawn_detached(run())

        # ✅ 추가: Output ON 상태에서 RF-Pulse Power setpoint만 변경
        def cb_rf_pulse_set_power(power: float) -> None:
            async def run():
                if not self.rf_pulse:
                    self.append_log("RFPulse", "RF-Pulse 미지원 챔버입니다.")
                    return
                try:
                    self._ensure_background_started()

                    ok = await self.rf_pulse.set_reference_power(float(power), pause_polling=True)
                    if not ok:
                        # rf_pulse가 command_failed 이벤트를 이미 올리므로,
                        # 여기서는 로그만 남기고 실패 판정은 _pump_rfpulse_events로 일원화
                        self.append_log("RFPulse", "set_reference_power returned False (will be handled by event pump)")

                except Exception as e:
                    why = f"RF-Pulse set_reference_power failed: {e!r}"
                    self.append_log("RFPulse", why)
                    self.process_controller.on_rf_pulse_failed(
                        e,
                        code=getattr(e, "code", None) or getattr(e, "error_code", None),
                        meta={
                            "stage": "set_reference_power",
                            "ch": self.ch,
                            "power": float(power),
                        },
                    )

            self._spawn_detached(run())

        def cb_rf_pulse_stop():
            async def run():
                if not self.rf_pulse:
                    return
                try:
                    self.rf_pulse.stop_process()
                except Exception as e:
                    self.append_log("RFPulse", f"stop_process failed: {e!r}")
                    self.process_controller.on_rf_pulse_failed(
                        e,
                        code=getattr(e, "code", None) or getattr(e, "error_code", None),
                        meta={
                            "stage": "stop_process",
                            "ch": self.ch,
                        },
                    )
            self._spawn_detached(run())

        def cb_ig_wait(base_pressure: float) -> None:
            async def _run():
                self._ensure_background_started()
                ok = await self.ig.wait_for_base_pressure(
                    float(base_pressure),
                    interval_ms=self.cfg.IG_POLLING_INTERVAL_MS
                )
                self.append_log("IG", f"wait_for_base_pressure returned: {ok}")
            self._spawn_detached(_run())

        def cb_ig_cancel():
            self._spawn_detached(self.ig.cancel_wait())

        def cb_oes_run(duration_sec: float, integration_ms: int):
            async def run():
                try:
                    self._ensure_task_alive(f"Pump.OES.{self.ch}", self._pump_oes_events)
                    self._ensure_background_started()

                    if hasattr(self.oes, "drain_events"):
                        with contextlib.suppress(Exception):
                            await self.oes.drain_events()

                    # ✅ 실제 공정 경로에서도 동일 timeout 정책을 강제
                    oes_init_timeout_s = float(self.cfg._get("CHAMBER_OES_INIT_TIMEOUT_S", 30.0))

                    if not getattr(self, "_oes_initialized", False):
                        ok = await self.oes.initialize_device(timeout_s=oes_init_timeout_s, force=False)
                        if not ok:
                            detail = getattr(self.oes, "_init_error", None) or "unknown"
                            raise RuntimeError(f"OES 초기화 실패: {detail}")
                        self._oes_initialized = True

                    self._soon(self._safe_clear_oes_plot)
                    self._oes_active = True
                    await self.oes.run_measurement(duration_sec, integration_ms)

                except Exception as e:
                    self._oes_active = False
                    self._oes_initialized = False  # 다음 런에서 재초기화 시도
                    self.append_log("OES", f"OES 실패: {e!r}")

                    if self.chat:
                        with contextlib.suppress(Exception):
                            self.chat.notify_text(f"[OES] 실패: {e!r}")
                            if hasattr(self.chat, "flush"):
                                self.chat.flush()

                    self.process_controller.on_oes_failed(
                        "OES",
                        e,
                        code=getattr(e, "code", None) or getattr(e, "error_code", None),
                        meta={
                            "stage": "run_measurement",
                            "ch": self.ch,
                            "duration_sec": float(duration_sec),
                            "integration_ms": int(integration_ms),
                        },
                    )

            self._spawn_detached(run())

        def cb_rga_scan():
            async def _run():
                timeout_s = float(self.cfg._get("RGA_WORKER_TIMEOUT_S", 60.0))
                self._soon(self._graph_clear_rga_plot_safe)

                try:
                    # ✅ auto_connect 차단 상태여도 Pump.RGA는 올려야 finished/data를 소비함
                    if self.rga:
                        self._ensure_task_alive(f"Pump.RGA.{self.ch}", self._pump_rga_events)

                    if self.rga:
                        await self.rga.scan_histogram_to_csv(timeout_s=timeout_s)
                    else:
                        raise RuntimeError("RGA 어댑터 없음")

                except Exception as e:
                    self.append_log("RGA", f"예외로 RGA 스캔 실패: {e!r} → 다음 단계")

                finally:
                    # ✅ 핵심: 펌프/이벤트 누락이 있어도 공정이 여기서 영원히 멈추지 않게 한다
                    try:
                        self.process_controller.on_rga_finished()
                    except Exception:
                        pass

            self._spawn_detached(_run())

        # 컨트롤러 생성
        self.process_controller = ProcessController(
            send_plc=cb_plc,
            send_mfc=cb_mfc,

            # 연속 파워
            send_dc_power=cb_dc_power, 
            stop_dc_power=cb_dc_stop,
            send_dc_power2=cb_dc_power2,
            stop_dc_power2=cb_dc2_stop,
            send_rf_power=cb_rf_power, 
            stop_rf_power=cb_rf_stop,

            # 펄스 파워(완전 분리)
            start_dc_pulse=cb_dc_pulse_start, 
            stop_dc_pulse=cb_dc_pulse_stop,
            set_dc_pulse_power=cb_dc_pulse_set_power,   # ✅ 추가

            start_rf_pulse=cb_rf_pulse_start, 
            stop_rf_pulse=cb_rf_pulse_stop,
            set_rf_pulse_power=cb_rf_pulse_set_power,   # ✅ 추가

            ig_wait=cb_ig_wait, 
            cancel_ig=cb_ig_cancel,
            rga_scan=cb_rga_scan, 
            oes_run=cb_oes_run,

            ch=self.ch,
            supports_dc_cont=self.supports_dc_cont,
            supports_dc_cont2=self.supports_dc_cont2,
            supports_rf_cont=self.supports_rf_cont,
            supports_dc_pulse=self.supports_dc_pulse,
            supports_rf_pulse=self.supports_rf_pulse,
        )

        # 이벤트 펌프 루프(컨트롤러 → UI/로거/다음공정)
        self._ensure_task_alive("Pump.PC", self._pump_pc_events)

    # ------------------------------------------------------------------
    # 이벤트 펌프들
    async def _pump_pc_events(self) -> None:
        q = self.process_controller.event_q
        while True:
            ev = await q.get()
            kind = ev.kind
            payload = ev.payload or {}
            try:
                if kind == "log":
                    self.append_log(payload.get("src", f"PC{self.ch}"), payload.get("msg", ""))

                elif kind == "state":
                    self._apply_process_state_message(payload.get("text", ""))

                elif kind == "status":
                    self._on_process_status_changed(bool(payload.get("running", False)))

                elif kind == "started":
                    params = payload.get("params", {}) or {}

                    # ✅ 공정명 키 통일: process_name을 정식 키로 사용
                    #    (레거시 키 Process_name / process_note는 동일값으로 유지)
                    pname = str(params.get("process_name") or params.get("Process_name") or params.get("process_note") or "").strip()
                    if not pname:
                        pname = f"Run CH{self.ch}"
                    params["process_name"] = pname
                    params["Process_name"] = pname
                    params["process_note"] = pname

                    # ✅ 시작 카드 전송(성공 시 로그 X, 실패만 로그)
                    # AFTER: 시작 카드 전송 후 즉시 flush
                    if self.chat:
                        p = dict(params)
                        p.setdefault("ch", self.ch)
                        p["prefix"] = f"CH{self.ch} Sputter"

                        # ➋ 리스트 공정이면 공정명에 " (i/n)"을 덧붙이고 인덱스도 함께 넘김
                        try:
                            total = len(getattr(self, "process_queue", []) or [])
                            cur   = int(getattr(self, "current_process_index", -1)) + 1
                            if total > 0 and cur > 0:
                                name_key = "Process_name" if "Process_name" in p else ("process_name" if "process_name" in p else None)
                                if name_key:
                                    base = (str(p.get(name_key, "")) or f"Run CH{self.ch}").strip()
                                    p[name_key] = f"{base} ({cur}/{total})"
                                p["process_index"] = cur
                                p["process_total"] = total
                        except Exception:
                            pass

                        p = self._format_card_payload_for_chat(p)
                        try:
                            ret = self.chat.notify_process_started(p)
                            if inspect.iscoroutine(ret):
                                await ret
                            # ★ 추가: 버퍼링 드롭 방지(즉시 밀어내기)
                            if hasattr(self.chat, "flush"):
                                self.chat.flush()
                        except Exception as e:
                            self.append_log("CHAT", f"구글챗 시작 카드 전송 실패: {e!r}")

                    # ✅ 시작시각 확정: 버튼-누른-시각 우선, 없으면 지금 시각 (둘 다 tz 없음)
                    params = dict(params)
                    t0 = params.get("t0_pressed_wall") or datetime.now().isoformat(timespec="seconds")
                    params["t0_wall"]   = t0
                    params["started_at"] = t0  # 하위호환 키 동일값

                    self._oes_active = False  # OES는 별도 cb에서 True로 바꿈

                    # Plasma Cleaning 스타일 헤더 포함한 오픈 (중복 방지)
                    if not getattr(self, "_log_file_path", None):
                        self._open_run_log(params)
                    else:
                        self.append_log("Logger", f"이미 열린 로그 파일 사용: {self._log_file_path.name}")

                    self.data_logger.start_new_log_session(params)
                    self._runtime_arc_notified = False  # ✅ 추가: 런타임 레벨 arc 알림 중복 방지 플래그

                    try:
                        # ✅ Arc 카운터 초기화
                        if self.dc_pulse is not None:
                            with contextlib.suppress(Exception):
                                self.dc_pulse.reset_arc_counts()
                        self._gdrive_arc_sent = False
                        
                        # 성공 시에도 명시적으로 남겨 두면 나중에 추적이 쉬움
                        self.append_log("CSV", "Sputter Calib 로그 세션 시작")
                    except Exception as e:
                        # 시작 자체가 실패한 경우도 CH 로그에 남김
                        self.append_log("CSV", f"Sputter Calib 로그 세션 시작 실패: {e!r}")

                    self._soon(self._graph_reset_safe)

                    # ✅ 텍스트 알림은 기존 그대로 유지
                    name = (params.get("process_name")
                            or params.get("Process_name")
                            or params.get("process_note")
                            or f"Run CH{self.ch}")
                    t = params.get("process_time", 0) or 0
                    line = f"▶️ CH{self.ch} '{name}' 시작 (t={float(t):.1f}s)"
                    self.append_log("MAIN", line)

                    # ★ 공정 모니터 활성화 (gas/pressure 편차 감시)
                    with contextlib.suppress(Exception):
                        self._process_monitor.activate(params)

                    # 폴링 타깃 초기화
                    self._last_polling_targets = None

                elif kind == "finished":
                    ok = False
                    detail = {}
                    try:
                        ok = bool(payload.get("ok", False))
                        detail = payload.get("detail", {}) or {}

                        # ✅ test/stop 판별 (기존 로직 영향 없음)
                        is_test = bool(detail.get("test_mode", False))
                        is_stopped = bool(detail.get("stopped", False))
                        is_test_cancel = is_test and is_stopped

                        # ============================
                        # ✅ CSV용 Result/실제 시간 반영
                        # ============================

                        # 1) Result 매핑: (process_controller는 SUCCESS/FAIL/STOP을 보냄)
                        raw_result = str(detail.get("result") or "").strip().upper()
                        if raw_result in ("SUCCESS", "OK", "TRUE", "1"):
                            result = "성공"
                        elif raw_result in ("STOP", "STOPPED", "CANCEL", "CANCELED", "CANCELLED"):
                            result = "stop"
                        elif raw_result in ("FAIL", "FAILED", "ERROR", "FALSE", "0"):
                            result = "실패"
                        else:
                            # detail에 result가 없을 때만 최소 폴백(추정 X: ok/stopped로만 결정)
                            result = "stop" if is_stopped else ("성공" if ok else "실패")

                        # 2) DataLogger의 process_params에 덮어쓰기(여기가 핵심)
                        try:
                            pp = getattr(self.data_logger, "process_params", None)
                            if isinstance(pp, dict):
                                # 공정명도 확실히 통일
                                pname = str(detail.get("process_name") or detail.get("Process_name") or detail.get("process_note") or "").strip()
                                if pname:
                                    pp["process_name"] = pname
                                    pp["Process_name"] = pname
                                    pp["process_note"] = pname

                                # Result 컬럼용
                                pp["result"] = result
                                pp["stopped"] = bool(is_stopped)

                                # 실제 진행 시간(분) → 기존 컬럼(shutter_delay / process_time)에 그대로 넣기
                                v = detail.get("actual_shutter_delay_min", None)
                                if v is not None:
                                    pp["shutter_delay"] = float(v)

                                v = detail.get("actual_process_time_min", None)
                                if v is not None:
                                    pp["process_time"] = float(v)

                        except Exception as e:
                            self.append_log("CSV", f"Sputter Calib CSV 메타 갱신 실패(무시): {e!r}")

                        # 3) 이제 기록(성공/실패/stop 모두 기록)
                        self.append_log("CSV", f"Sputter Calib CSV 기록 요청 (ok={ok}, result={result})")
                        
                        # - data_logger가 'ok=False면 return'인 현재 버전이면, 아래 인자를 ok로 두면 실패/stop이 기록되지 않습니다.
                        # - data_logger를 실패/stop도 기록하도록 수정한 뒤에는 ok를 그대로 넘겨도 됩니다.
                        self.data_logger.finalize_and_write_log(ok)

                        # ✅ Google Drive 엑셀 로그 (메인 공정과 완전 독립)
                        if _GDRIVE_OK and ok:
                            try:
                                from lib import config_common as _cfgc
                                if self.dc_pulse is not None:
                                    s, h = self.dc_pulse.arc_counts
                                    self.data_logger.process_params["soft_arc_count"] = s
                                    self.data_logger.process_params["hard_arc_count"] = h

                                _pname = str(
                                    self.data_logger.process_params.get("process_name")
                                    or self.data_logger.process_params.get("Process_name")
                                    or ""
                                ).strip()
                                _cb = getattr(self, "_main_done_callback", None)
                                _pc_params = _cb(self.ch, _pname, self.data_logger) if callable(_cb) else None

                                # ✅ PENDING은 더 이상 반환되지 않으므로 None or dict만 옴
                                _t = asyncio.create_task(
                                    _gdrive_save(
                                        ch=self.ch,
                                        data_logger=self.data_logger,
                                        operator=self.data_logger.process_params.get("operator", ""),
                                        substrate=self.data_logger.process_params.get("substrate", ""),
                                        note=self.data_logger.process_params.get("note", ""),
                                        pc_params=_pc_params if isinstance(_pc_params, dict) else None,
                                        log_dir=None,
                                        arc_thresh=getattr(_cfgc, "GDRIVE_ARC_ALERT_THRESH", 5),
                                        refp_warn=getattr(_cfgc, "GDRIVE_REF_P_WARN_W", 20.0),
                                        webhook_url=getattr(_cfgc, "CHAT_WEBHOOK_MONITOR_URL", ""),
                                        arc_alert_sent=getattr(self, "_gdrive_arc_sent", False),
                                    )
                                )
                                def _on_gdrive_done(t):
                                    try:
                                        if t.result():
                                            self._gdrive_arc_sent = True
                                        self.append_log("CSV", "GDrive 저장 완료")
                                    except Exception as e:
                                        self.append_log("CSV", f"GDrive 저장 실패: {e!r}")
                                _t.add_done_callback(_on_gdrive_done)

                            except Exception as e:
                                self.append_log("CSV", f"GDrive 저장 준비 실패: {e!r}")

                        # ➊ 카드 헤더용 prefix: "CHx Sputter"
                        detail.setdefault("ch", self.ch)
                        detail.setdefault("prefix", f"CH{self.ch} Sputter")

                        # ➋ 리스트 공정 표기 (i/n) 동일하게 적용
                        try:
                            total = len(getattr(self, "process_queue", []) or [])
                            cur   = int(getattr(self, "current_process_index", -1)) + 1
                            if total > 0 and cur > 0:
                                name_key = "Process_name" if "Process_name" in detail else ("process_name" if "process_name" in detail else None)
                                if name_key:
                                    base = (str(detail.get(name_key, "")) or f"Run CH{self.ch}").strip()
                                    detail[name_key] = f"{base} ({cur}/{total})"
                                detail["process_index"] = cur
                                detail["process_total"] = total
                        except Exception:
                            pass

                        # ✅ 종료 카드 전송(성공 시 로그 X, 실패만 로그)
                        if self.chat:
                            # 라우팅/표시용 힌트 보강: CH2 누락으로 전송이 드롭/오경로 되는 문제 방지
                            payload = dict(detail)
                            payload.setdefault("ch", self.ch)           # ← 필수(라우팅)
                            payload.setdefault("prefix", self.prefix)   # ← 표시/구분용
                            # 시작 카드와 키를 맞춰 카드 템플릿이 동일하게 먹히도록 보정
                            if "process_note" not in payload and "process_name" in payload:
                                payload["process_note"] = payload["process_name"]

                            try:
                                # ✅ chuck 경고/목표 위치를 종료 카드로 전달
                                pos = str(getattr(self, "_run_chuck_position", "") or "").strip()
                                if pos:
                                    payload.setdefault("chuck_position", pos)

                                warns = list(getattr(self, "_run_warnings", []) or [])
                                if warns:
                                    payload.setdefault("warnings", warns)

                                ret = self.chat.notify_process_finished_detail(ok, payload)
                                if inspect.iscoroutine(ret):
                                    await ret
                                # Plasma cleaning과 동일하게 즉시 밀어내기(버퍼링 드롭 방지)
                                if hasattr(self.chat, "flush"):
                                    self.chat.flush()
                            except Exception as e:
                                self.append_log("CHAT", f"구글챗 종료 카드 전송 실패: {e!r}")

                            # 👇 추가: 카드가 잘려 보일 때를 대비해 '실패 이유'만 텍스트로 별도 전송
                            # ✅ 실패 이유 텍스트는 "진짜 실패"에만
                            if (not ok) and (not detail.get("stopped", False)):
                                reason = (str(detail.get("reason") or "")).strip()
                                if not reason:
                                    errs = detail.get("errors", [])
                                    if isinstance(errs, (list, tuple)) and errs:
                                        reason = str(errs[0])
                                    elif isinstance(errs, str):
                                        reason = errs
                                if reason:
                                    try:
                                        r = self.chat.notify_text(f"❌ CH{self.ch} 공정 실패 이유: {reason}")
                                        if inspect.iscoroutine(r):
                                            await r
                                        # ★ 추가: 실패 텍스트도 카드 직후에 바로 나가도록 즉시 flush
                                        if hasattr(self.chat, "flush"):
                                            self.chat.flush()
                                    except Exception as _e:
                                        self.append_log("CHAT", f"실패 이유 텍스트 알림 실패: {_e!r}")

                        try:
                            if self._skip_mfc_finalize_due_to_pc():
                                self.append_log("MFC", "PC 실행 중 → mfc on_process_finished 생략(공유 자원 보호)")
                            else:
                                self.mfc.on_process_finished(ok)
                        except Exception:
                            pass

                        # ★ 공정 모니터 비활성화
                        with contextlib.suppress(Exception):
                            self._process_monitor.deactivate()

                        # 0) 재연결 선차단 + 폴링 완전 OFF
                        self._auto_connect_enabled = False
                        self._run_select = None
                        self._last_polling_targets = None
                        # 남아 있을 수 있는 폴링 스위치를 즉시 모두 내림(장치 내부 워치독 종료 유도)
                        self._apply_polling_targets({"mfc": False, "dc_pulse": False, "rf_pulse": False, "dc": False, "dc2": False, "rf": False})

                        # ✅ finished 이벤트에서는 "결과 기록/카드/알림"까지만 수행한다.
                        #   - 장치 정리(_stop_device_watchdogs)
                        #   - 다음 공정 진행(큐 advance)
                        #   은 Runner가 PC_FINISHED 명령을 받아 "순차적으로" 처리한다.

                        self._last_polling_targets = None
                    except Exception as e:
                        self.append_log("MAIN", f"예외 발생 (finished 처리): {e}")
                        # ✅ 여기서 Runner/큐/UI를 리셋하지 않는다.
                        # (stale finished/후행 이벤트가 다음 공정을 끊는 레이스 방지)
                        pass

                    finally:
                        try:
                            stopped = bool(detail.get("stopped"))

                            if ok or stopped:
                                # ✅ 정상 종료(ok=True) 또는 사용자 STOP(stopped=True)은 idle로 표시
                                runtime_state.clear_error("chamber", self.ch)
                            else:
                                # ✅ 실패만 error
                                _reason = (str(detail.get("reason") or "")).strip()
                                if not _reason:
                                    _errs = detail.get("errors", None)
                                    if isinstance(_errs, (list, tuple)) and _errs:
                                        _reason = str(_errs[0])
                                    elif isinstance(_errs, str):
                                        _reason = _errs
                                if not _reason:
                                    _reason = "process failed"

                                runtime_state.set_error("chamber", self.ch, _reason)

                                # ✅ 실패 알림창: 사용자가 OK를 누르면 status 표시가 idle로 돌아가도록
                                try:
                                    _pname = (str(detail.get("process_name") or "").strip() or
                                            str(detail.get("Process_name") or "").strip() or
                                            "(process)")
                                    self._post_critical(
                                        f"CH{self.ch} 공정 실패",
                                        f"{_pname}\n\n사유: {_reason}\n\n확인을 누르면 상태 표시가 Idle로 변경됩니다.",
                                        clear_status_to_idle=True,
                                    )
                                except Exception:
                                    pass

                            runtime_state.mark_finished("chamber", self.ch)

                            # ✅ Runner에 finished 통지 (정리/다음 공정 진행은 Runner가 담당)
                            self._runner_put(_RunnerCmd(kind="PC_FINISHED", ok=bool(ok), detail=dict(detail)))

                        except Exception:
                            pass

                elif kind == "aborted":
                    try:
                        if self.chat:
                            try:
                                ret = self.chat.notify_text(f"🛑 CH{self.ch} 공정 중단")
                                if inspect.iscoroutine(ret):
                                    await ret
                            except Exception as e:
                                self.append_log("CHAT", f"구글챗 중단 알림 전송 실패: {e!r}")

                        # ✅ runtime_state는 '중단'으로만 마킹(선택)
                        try:
                            if not runtime_state.has_error("chamber", self.ch):
                                runtime_state.set_error("chamber", self.ch, "aborted")
                            runtime_state.mark_finished("chamber", self.ch)
                        except Exception:
                            pass

                        # ✅ 핵심: aborted는 Runner에만 통지하고,
                        #    UI/큐/딜레이/스테이지는 절대 여기서 건드리지 않는다.
                        self._runner_put(_RunnerCmd(
                            kind="PC_FINISHED",
                            ok=False,
                            detail={"stopped": True, "reason": "aborted"}
                        ))

                        self.append_log("MAIN", f"[PC] aborted event received → forwarded to Runner (ignored if stale)")

                    except Exception as e:
                        self.append_log("MAIN", f"예외 발생 (aborted 처리): {e}")
                        # ❌ 여기서도 _clear_queue_and_reset_ui() 같은 리셋 금지
                        pass

                elif kind == "polling_targets":
                    targets = dict(payload.get("targets") or {})
                    self._last_polling_targets = targets
                    self._apply_polling_targets(targets)

                elif kind == "polling":
                    active = bool(payload.get("active", False))

                    # ✅ 공정이 실제 실행 중일 때만 자동 기동
                    if active and self._auto_connect_enabled and self.process_controller.is_running:
                        self._ensure_background_started()

                    # ✅ 제거: "안전망" _apply_polling_targets 호출 삭제
                    #    → polling_targets 이벤트에서 이미 처리됨. 여기서 또 호출하면 모든 장비에
                    #      set_process_status(False)가 중복 실행되어 로그 노이즈 + 불필요한 큐 정리 발생
                    # if not active:
                    #     self._apply_polling_targets({"mfc": False, ...})

                    params = getattr(self.process_controller, "current_params", {}) or {}
                    use_dc_pulse = bool(params.get("use_dc_pulse", False))
                    use_rf_pulse = bool(params.get("use_rf_pulse", False))
                    use_dc_cont  = bool(params.get("use_dc_power", False))
                    use_dc_cont2 = bool(params.get("use_dc_power2", False))
                    use_rf_cont  = bool(params.get("use_rf_power", False))

                    # ✅ 동시 사용 허용 정책: 같은 계열(펄스↔연속) 상호 배타 제거
                    base_targets = {
                        "mfc":      active,
                        "dc_pulse": active and self.supports_dc_pulse and use_dc_pulse,
                        "rf_pulse": active and self.supports_rf_pulse and use_rf_pulse,
                        "dc":       active and self.supports_dc_cont  and use_dc_cont,
                        "dc2":      active and self.supports_dc_cont2 and use_dc_cont2,
                        "rf":       active and self.supports_rf_cont  and use_rf_cont,
                    }

                    if self._last_polling_targets:
                        lt = self._last_polling_targets
                        targets = {
                            "mfc":      base_targets["mfc"]      and bool(lt.get("mfc", False)),
                            "dc_pulse": base_targets["dc_pulse"] and bool(lt.get("dc_pulse", False)),
                            "rf_pulse": base_targets["rf_pulse"] and bool(lt.get("rf_pulse", False)),
                            "dc":       base_targets["dc"]       and bool(lt.get("dc", False)),
                            "dc2":      base_targets["dc2"]      and bool(lt.get("dc2", False)),
                            "rf":       base_targets["rf"]       and bool(lt.get("rf", False)),
                        }
                    else:
                        targets = base_targets

                    self._apply_polling_targets(targets)

                else:
                    self.append_log(f"MAIN{self.ch}", f"알 수 없는 PC 이벤트: {kind} {payload}")

            except Exception as e:
                self.append_log(f"MAIN{self.ch}", f"PC 이벤트 처리 예외: {e!r} (kind={kind})")
            finally:
                await asyncio.sleep(0)

    async def _pump_mfc_events(self) -> None:
        async for ev in self.mfc.events():
            k = ev.kind
            if k == "status":
                self.append_log(f"MFC{self.ch}", ev.message or "")
            elif k == "command_confirmed":
                self.process_controller.on_mfc_confirmed(ev.cmd or "")
            elif k == "command_failed":
                self.process_controller.on_mfc_failed(
                    ev.cmd or "",
                    ev.reason or "unknown",
                    code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                    meta={
                        "cmd": ev.cmd,
                        "gas": getattr(ev, "gas", None),
                        "kind": k,
                        "ch": self.ch,
                    },
                )
                # 중복 방지: 런타임에서 MFC 장비오류 카드는 전송하지 않음
            elif k == "flow":
                gas = ev.gas or ""
                flow = float(ev.value or 0.0)
                with contextlib.suppress(Exception):
                    self._dl_fire_and_forget(self.data_logger.log_mfc_flow, gas, flow)
                self.append_log(f"MFC{self.ch}", f"[poll] {gas}: {flow:.2f} sccm")
                
                # ★ 공정 모니터: 메인 공정 폴링 구간에서만 체크
                targets = getattr(self, "_last_polling_targets", None) or {}
                if targets.get("mfc"):
                    with contextlib.suppress(Exception):
                        self._process_monitor.check_flow(gas, flow)

            elif k == "pressure":
                txt = ev.text or (f"{ev.value:.3g}" if ev.value is not None else "")

                # ✅ Working Pressure는 메인 공정(process time) 폴링 구간에서만 수집
                #    - process_controller 에서 polling=True 인 DELAY(step) 동안만
                #      _last_polling_targets["mfc"] 가 True 가 됨
                targets = getattr(self, "_last_polling_targets", None) or {}
                if targets.get("mfc"):
                    with contextlib.suppress(Exception):
                        self.data_logger.log_mfc_pressure(txt)

                    # ★ 공정 모니터: pressure 체크
                    if ev.value is not None:
                        with contextlib.suppress(Exception):
                            self._process_monitor.check_pressure(float(ev.value))

                # UI / 로그에는 기존처럼 항상 표시
                self.append_log(f"MFC{self.ch}", f"[poll] ChamberP: {txt}")

    async def _pump_ig_events(self) -> None:
        async for ev in self.ig.events():
            k = ev.kind
            if k == "status":
                self.append_log(f"IG{self.ch}", ev.message or "")
            elif k == "pressure":
                try:
                    if ev.pressure is not None:
                        self._dl_fire_and_forget(self.data_logger.log_ig_pressure, float(ev.pressure))
                    elif ev.message:
                        self.data_logger.log_ig_pressure(ev.message)
                except Exception:
                    pass
            elif k == "base_reached":
                self.process_controller.on_ig_ok()
            elif k == "base_failed":
                self.process_controller.on_ig_failed(
                    "IG",
                    ev.message or "unknown",
                    code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                    meta={
                        "kind": k,
                        "pressure": getattr(ev, "pressure", None),
                        "ch": self.ch,
                    },
                )
                # 중복 방지: 런타임에서 IG 오류 카드는 전송하지 않음

    async def _pump_rga_events(self) -> None:
        adapter = self.rga
        if not adapter:
            return

        tag = f"RGA{self.ch}"
        finished_called = False  # ✅ 중복 방지(혹시 failed/finished 둘 다 들어오거나 예외 발생 시)

        try:
            async for ev in adapter.events():
                if ev.kind == "status":
                    self.append_log(tag, ev.message or "")

                elif ev.kind == "data":
                    self._graph_update_rga_safe(ev.mass_axis, ev.pressures)
                    # finish는 finished 이벤트에서만 처리(중복 방지)

                elif ev.kind == "finished":
                    if not finished_called:
                        finished_called = True
                        try:
                            self.process_controller.on_rga_finished()
                        except Exception as e:
                            self.append_log(tag, f"on_rga_finished() error: {e!r}")

                elif ev.kind == "failed":
                    why = ev.message or "RGA failed"
                    self.append_log(tag, f"측정 실패: {why} → 다음 단계")

                    # ✅ 워커가 준 stdout/stderr 있으면 같이 남겨서 원인 추적 가능하게
                    with contextlib.suppress(Exception):
                        payload = getattr(ev, "payload", None) or {}
                        stderr = (payload.get("stderr") or "").strip()
                        stdout = (payload.get("stdout") or "").strip()
                        if stderr:
                            self.append_log(tag, f"stderr: {stderr[-800:]}")
                        if stdout:
                            self.append_log(tag, f"stdout: {stdout[-800:]}")

                    if self.chat:
                        with contextlib.suppress(Exception):
                            self.chat.notify_text(f"[{tag}] 측정 실패: {why} → 건너뜀")
                            if hasattr(self.chat, "flush"):
                                self.chat.flush()

                    if not finished_called:
                        finished_called = True
                        self.process_controller.on_rga_finished()

        except Exception as e:
            # ✅ adapter.events() 자체가 예외로 끊겨도 공정은 계속 진행되게
            self.append_log(tag, f"RGA event pump crashed: {e!r} → 다음 단계")
            if self.chat:
                with contextlib.suppress(Exception):
                    self.chat.notify_text(f"[{tag}] 이벤트 루프 예외: {e!r} → 건너뜀")
                    if hasattr(self.chat, "flush"):
                        self.chat.flush()

            if not finished_called:
                self.process_controller.on_rga_finished()

    async def _pump_dc_events(self) -> None:
        if not self.dc_power:
            return
        async for ev in self.dc_power.events():
            k = ev.kind
            if k == "status":
                self.append_log("DC1", ev.message or "")
            elif k == "display":
                with contextlib.suppress(Exception):
                    self.data_logger.log_dc_power(
                        float(ev.power  or 0.0),
                        float(ev.voltage or 0.0),
                        float(ev.current or 0.0),
                    )
                self._display_dc(ev.power, ev.voltage, ev.current)
                self.append_log("DC1", f"측정: {float(ev.power or 0.0):.1f} W, {float(ev.voltage or 0.0):.1f} V, {float(ev.current or 0.0):.3f} A")
            elif k == "target_reached":
                self.process_controller.on_dc_target_reached()
            elif k == "target_failed":
                self._dc_failed_flag = True
                self.process_controller.on_dc_target_failed(
                    ev.message or "low-power",
                    code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                    meta={
                        "kind": k,
                        "power": getattr(ev, "power", None),
                        "voltage": getattr(ev, "voltage", None),
                        "current": getattr(ev, "current", None),
                        "ch": self.ch,
                    },
                )
            elif k == "power_off_finished":
                if not self._dc_failed_flag:                # ★ 추가: 실패 시에는 OK 토큰(다음 스텝 진행) 차단
                    self.process_controller.on_dc_off_finished()
                else:
                    self._dc_failed_flag = False            #    1회성 플래그 해제

    async def _pump_dc2_events(self) -> None:
        if not self.dc_power2:
            return
        async for ev in self.dc_power2.events():
            k = ev.kind
            if k == "status":
                self.append_log("DC2", ev.message or "")
            elif k == "display":
                with contextlib.suppress(Exception):
                    self.data_logger.log_dc2_power(
                        float(ev.power  or 0.0),
                        float(ev.voltage or 0.0),
                        float(ev.current or 0.0),
                    )
                self._display_dc(ev.power, ev.voltage, ev.current, unit=2)
                self.append_log("DC2", f"측정: {float(ev.power or 0.0):.1f} W, {float(ev.voltage or 0.0):.1f} V, {float(ev.current or 0.0):.3f} A")
            elif k == "target_reached":
                self.process_controller.on_dc2_target_reached()
            elif k == "target_failed":
                self._dc2_failed_flag = True
                self.process_controller.on_dc2_target_failed(
                    ev.message or "low-power",
                    code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                    meta={
                        "kind": k,
                        "power": getattr(ev, "power", None),
                        "voltage": getattr(ev, "voltage", None),
                        "current": getattr(ev, "current", None),
                        "ch": self.ch,
                        "unit": 2,
                    },
                )
            elif k == "power_off_finished":
                if not self._dc2_failed_flag:
                    self.process_controller.on_dc2_off_finished()
                else:
                    self._dc2_failed_flag = False

    async def _pump_rf_events(self) -> None:
        if not self.rf_power:
            return
        async for ev in self.rf_power.events():
            k = ev.kind
            if k == "status":
                self.append_log(f"RF{self.ch}", ev.message or "")
            elif k == "display":
                fwd = float(ev.forward or 0.0)
                ref = float(ev.reflected or 0.0)
                # 데이터 로거 저장 + UI 갱신 + 텍스트 로그
                with contextlib.suppress(Exception):
                    self._dl_fire_and_forget(self.data_logger.log_rf_power, fwd, ref)
                self._display_rf(fwd, ref)
                self.append_log(f"RF{self.ch}", f"[poll] fwd={fwd:.1f}W, ref={ref:.1f}W")
            elif k == "target_reached":
                self.process_controller.on_rf_target_reached()
            elif k == "target_failed":
                self.process_controller.on_rf_target_failed(
                    ev.message or "unknown",
                    code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                    meta={
                        "kind": k,
                        "forward": getattr(ev, "forward", None),
                        "reflected": getattr(ev, "reflected", None),
                        "ch": self.ch,
                    },
                )
            elif k == "power_off_finished":
                self.process_controller.on_rf_off_finished()

    async def _pump_rfpulse_events(self) -> None:
        if not self.rf_pulse:
            return
        async for ev in self.rf_pulse.events():
            k = ev.kind
            if k == "status":
                self.append_log(f"RFPulse{self.ch}", ev.message or "")
                
                # ★ 변경: REFP 경고/복귀는 Monitor 웹훅(가스/압력 편차와 동일 채널)으로 전달
                msg = ev.message or ""
                if "REFP_WARN" in msg or "REFP 정상 복귀" in msg:
                    with contextlib.suppress(Exception):
                        import re
                        # 원본 메시지:
                        #   "⚠ REFP_WARN: REFP=5.3W ≥ 5.0W (공정은 계속 진행)"
                        #   "REFP 정상 복귀: REFP=2.1W < 5.0W"
                        nums = re.findall(r"(\d+(?:\.\d+)?)\s*W", msg)
                        refp_w = float(nums[0]) if len(nums) >= 1 else 0.0
                        warn_w = float(nums[1]) if len(nums) >= 2 else 0.0
                        recovered = ("정상 복귀" in msg)
                        proc_name = str(
                            getattr(self.process_controller, "current_params", {}).get("process_name")
                            or f"CH{self.ch}"
                        )
                        self._process_monitor.notify_refp_warning(
                            refp_w, warn_w,
                            recovered=recovered,
                            process_name=proc_name,
                        )
                            
            elif k == "power":
                with contextlib.suppress(Exception):
                    fwd = float(ev.forward or 0.0)
                    ref = float(ev.reflected or 0.0)
                    self._dl_fire_and_forget(self.data_logger.log_rfpulse_power, fwd, ref)
                    self._display_rf(fwd, ref)   # ← 추가: 화면 갱신
            elif k == "target_reached":
                self.process_controller.on_rf_pulse_target_reached()
            elif k == "command_failed":
                self.process_controller.on_rf_pulse_failed(
                    ev.reason or "unknown",
                    code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                    meta={
                        "cmd": getattr(ev, "cmd", None),
                        "kind": k,
                        "ch": self.ch,
                    },
                )
            elif k == "power_off_finished":
                self.process_controller.on_rf_pulse_off_finished()

    async def _pump_dcpulse_events(self) -> None:
        if not self.dc_pulse:
            return
        async for ev in self.dc_pulse.events():
            try:
                k = ev.kind
                if k == "status":
                    self.append_log(f"DCPulse{self.ch}", ev.message or "")
                
                elif k == "telemetry":
                    # 장비 내부 폴링 결과(P/V/I)를 화면/로거에 반영
                    P = getattr(ev, "power",   None)
                    V = getattr(ev, "voltage", None)
                    I = getattr(ev, "current", None)

                    # 혹시 dict 형태로 올 수도 있으니 보강
                    if (P is None or V is None or I is None) and hasattr(ev, "eng"):
                        eng = getattr(ev, "eng") or {}
                        P = P if P is not None else float(eng.get("P_W", 0.0))
                        V = V if V is not None else float(eng.get("V_V", 0.0))
                        I = I if I is not None else float(eng.get("I_A", 0.0))

                    # on_telemetry가 이미 DataLogger에 기록했다면 중복 방지
                    if not callable(getattr(self.data_logger, "log_dcpulse_power", None)):
                        try:
                            self.data_logger.log_dc_power(float(P or 0.0), float(V or 0.0), float(I or 0.0))
                        except Exception:
                            pass

                    self._display_dc(P, V, I)
                    self.append_log(
                        f"DCPulse{self.ch}",
                        f"[telemetry] P={float(P or 0):.1f} W, V={float(V or 0):.2f} V, I={float(I or 0):.3f} A"
                    )

                elif k == "arc_threshold_reached":
                    # ✅ 런타임 레벨 중복 방지 (dc_pulse 내부 가드의 2차 방어선)
                    if getattr(self, "_runtime_arc_notified", False):
                        continue
                    self._runtime_arc_notified = True

                    soft = int(getattr(ev, "power", 0) or 0)
                    hard = int(getattr(ev, "voltage", 0) or 0)
                    proc_name = str(
                        getattr(self.process_controller, "current_params", {}).get("process_name")
                        or f"CH{self.ch}"
                    )
                    msg = (
                        f"⚠️ CH{self.ch} DC Pulse Arc 경고\n"
                        f"공정: {proc_name}\n"
                        f"Soft Arc: {soft}회  Hard Arc: {hard}회  합계: {soft + hard}회"
                    )
                    self.append_log(f"DCPulse{self.ch}", msg)
                    with contextlib.suppress(Exception):
                        self._process_monitor.notify_arc_warning(
                            soft_arc=soft,
                            hard_arc=hard,
                            process_name=proc_name,
                        )

                elif k == "command_confirmed":
                    cmd = (ev.cmd or "").upper()

                    # 시작 step은 OUTPUT_ON 확인으로 완료
                    if cmd.startswith("OUTPUT_ON"):
                        self.process_controller.on_dc_pulse_target_reached()

                    # 중간 power 변경 step은 REF_POWER 확인으로 완료
                    elif cmd.startswith("REF_POWER"):
                        self.process_controller.on_dc_pulse_set_confirmed()

                    elif cmd.startswith("OUTPUT_OFF"):
                        self.process_controller.on_dc_pulse_off_finished()

                elif k == "command_failed":
                    why_raw = ev.reason or "unknown"
                    why = str(why_raw).lower()
                    cmd = (ev.cmd or "").upper()

                    self.append_log(f"DCPulse{self.ch}", f"CMD FAIL: {cmd} ({why_raw})")

                    # ✅ 내부 진단/복구용 명령 실패는 공정 실패로 승격하지 않는다.
                    if cmd.startswith("READ_") or cmd == "FAULT_RESET":
                        continue

                    # ✅ OUTPUT_OFF 실패는 별도 안전 경고
                    if cmd.startswith("OUTPUT_OFF"):
                        alert = (
                            f"⚠️ CH{self.ch} DC Pulse OUTPUT OFF 실패 - 출력 상태 미확인 "
                            f"(켜져 있을 가능성 있음). 장비 패널의 HV/OUTPUT 상태를 즉시 확인하세요."
                        )

                        self.append_log(f"DCPulse{self.ch}", alert)

                        try:
                            warns = getattr(self, "_run_warnings", None)
                            if isinstance(warns, list):
                                warns.append("DC Pulse OUTPUT OFF 실패(출력 상태 미확인)")
                        except Exception:
                            pass

                        if self.chat:
                            with contextlib.suppress(Exception):
                                self.chat.notify_error_with_src("DCPulse", alert)
                                if hasattr(self.chat, "flush"):
                                    self.chat.flush()

                        self.process_controller.on_dc_pulse_failed(
                            "OUTPUT_OFF 미확인(출력 상태 미확인)",
                            code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                            meta={
                                "cmd": cmd,
                                "kind": k,
                                "ch": self.ch,
                                "safety": "output_state_unconfirmed",
                            },
                        )
                        continue

                    # AUTO_STOP은 진짜 공정 실패
                    if cmd == "AUTO_STOP" or "target_failed" in why:
                        if self.chat:
                            with contextlib.suppress(Exception):
                                self.chat.notify_error_with_src(
                                    "DCPulse",
                                    "세트포인트 이탈(연속) 또는 P=0W 감지 → 전체 공정 중단"
                                )
                                if hasattr(self.chat, "flush"):
                                    self.chat.flush()

                    self.process_controller.on_dc_pulse_failed(
                        why_raw,
                        code=getattr(ev, "code", None) or getattr(ev, "error_code", None),
                        meta={
                            "cmd": cmd,
                            "kind": k,
                            "ch": self.ch,
                        },
                    )

            except Exception as e:
                # 펌프 루프 자체가 죽지 않도록 방어
                self.append_log(f"DCPulse{self.ch}", f"[pump] 예외 발생: {e!r}")

    async def _pump_oes_events(self) -> None:
        async for ev in self.oes.events():
            try:
                k = getattr(ev, "kind", None)
                if k == "status":
                    self.append_log(f"OES{self.ch}", ev.message or ""); continue
                if k in ("data", "spectrum", "frame"):
                    x = getattr(ev, "x", None)
                    if x is None: x = getattr(ev, "wavelengths", None)
                    if x is None: x = getattr(ev, "lambda_axis", None)

                    y = getattr(ev, "y", None)
                    if y is None: y = getattr(ev, "intensities", None)
                    if y is None: y = getattr(ev, "counts", None)

                    if x is not None and y is not None:
                        x_list = x.tolist() if hasattr(x, "tolist") else list(x)
                        y_list = y.tolist() if hasattr(y, "tolist") else list(y)
                        self._post_update_oes_plot(x_list, y_list)
                    else:
                        self.append_log(f"OES{self.ch}", f"경고: 데이터 필드 없음: kind={k}")
                    continue

                elif k == "finished":
                    # ✅ cb_oes_run()에서 _oes_active=True로 올린 '이번 런'만 처리
                    if not getattr(self, "_oes_active", False):
                        self.append_log(f"OES{self.ch}", "이전 런 잔여 'finished' 이벤트 무시")
                        continue

                    ok = bool(getattr(ev, "success", False))
                    out_csv = getattr(ev, "out_csv", None) or getattr(ev, "csv_path", None)
                    msg = getattr(ev, "message", None) or ("측정 완료" if ok else "measure failed")

                    if ok:
                        if out_csv:
                            self.append_log(f"OES{self.ch}", f"{msg} (csv={out_csv})")
                        else:
                            self.append_log(f"OES{self.ch}", msg)
                        self._oes_active = False
                        self.process_controller.on_oes_ok()
                    else:
                        # 실패면 stderr tail도 남겨서 원인 추적이 가능하게
                        err_tail = getattr(ev, "error", None)
                        if out_csv:
                            self.append_log(f"OES{self.ch}", f"측정 실패: {msg} (csv={out_csv}) → 종료 절차로 전환")
                        else:
                            self.append_log(f"OES{self.ch}", f"측정 실패: {msg} → 종료 절차로 전환")

                        if err_tail:
                            _t = str(err_tail)
                            if len(_t) > 2000:
                                _t = _t[:2000] + "..."
                            self.append_log(f"OES{self.ch}", f"워커 stderr tail:\n{_t}")

                        self._oes_active = False
                        # 실패 시 다음 런에서 init을 다시 시도하도록 캐시 무효화
                        self._oes_initialized = False
                        self.process_controller.on_oes_failed(
                            "OES",
                            msg,
                            meta={
                                "kind": k,
                                "success": ok,
                                "csv": out_csv,
                                "stderr_tail": err_tail,
                                "ch": self.ch,
                            },
                        )
                    continue

                self.append_log(f"OES{self.ch}", f"알 수 없는 이벤트: {ev!r}")
            except Exception as e:
                self.append_log(f"OES{self.ch}", f"이벤트 처리 예외: {e!r}")
                continue

    # ------------------------------------------------------------------
    # 백그라운드 시작/보장
    def _ensure_task_alive(self, name: str, coro_factory: Callable[[], Coroutine[Any, Any, Any]]) -> None:
        """
        ✅ 핵심 구조 변경
        - Pump.*(이벤트 펌프)는 '상주(keep-alive)'로 유지
        - 공정 종료 cleanup에서 cancel 대상으로 넣지 않음(= _bg_tasks에 넣지 않음)

        왜?
        - 기존 구조는 finished 처리 중 _stop_device_watchdogs()에서 _bg_tasks를 대량 cancel하는데,
        이 cancel이 꼬이면 다음 공정 StartAfterPreflight가 생성/스케줄되지 않는 문제가 생김.
        """
        if not hasattr(self, "_keepalive_tasks"):
            self._keepalive_tasks: dict[str, asyncio.Task] = {}

        t = self._keepalive_tasks.get(name)
        if t and not t.done():
            return

        loop = self._loop
        try:
            running = asyncio.get_running_loop()
        except RuntimeError:
            running = None

        def _create() -> None:
            # ✅ 추가: 예약 실행 시점에 “이미 살아있는 task”가 생겼으면 중복 생성 금지
            exist = self._keepalive_tasks.get(name)
            if exist and not exist.done():
                return
                
            try:
                task = loop.create_task(coro_factory(), name=name)
            except Exception as e:
                tb = "".join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
                self.append_log(f"Task{self.ch}", f"[{name}] create_task failed:\n{tb}")
                return

            def _done(tsk: asyncio.Task) -> None:
                # 끝나면 dict에서 제거 (다음 ensure에서 재생성)
                with contextlib.suppress(Exception):
                    if self._keepalive_tasks.get(name) is tsk:
                        self._keepalive_tasks.pop(name, None)
                if tsk.cancelled():
                    return
                with contextlib.suppress(Exception):
                    exc = tsk.exception()
                    if exc:
                        tb2 = "".join(traceback.format_exception(type(exc), exc, exc.__traceback__)).rstrip()
                        self.append_log(f"Task{self.ch}", f"[{name}] crashed:\n{tb2}")

            task.add_done_callback(_done)
            self._keepalive_tasks[name] = task

        if running is loop:
            _create()
        else:
            loop.call_soon_threadsafe(_create)

    def _ensure_background_started(self) -> None:
        # 🔒 실패 등으로 자동 연결 차단 중이면 아무 것도 올리지 않음
        if not getattr(self, "_auto_connect_enabled", True):
            return
        if getattr(self, "_ensuring_bg", False):
            return
        self._ensuring_bg = True
        try:
            self._ensure_devices_started()
            sel = getattr(self, "_run_select", None) or {}

            self._ensure_task_alive("Pump.PC", self._pump_pc_events)
            self._ensure_task_alive(f"Pump.MFC.{self.ch}", self._pump_mfc_events)  # 항상
            self._ensure_task_alive(f"Pump.IG.{self.ch}",  self._pump_ig_events)   # 항상

            if self.rga:
                self._ensure_task_alive(f"Pump.RGA.{self.ch}", self._pump_rga_events)

            # 연속 DC/RF는 PLC 경유 제어라 기존 그대로(변경 없음)
            if self.dc_power:
                self._ensure_task_alive(f"Pump.DC.{self.ch}", self._pump_dc_events)
            if self.dc_power2:
                self._ensure_task_alive(f"Pump.DC2.{self.ch}", self._pump_dc2_events)
            if self.rf_power:
                self._ensure_task_alive(f"Pump.RF.{self.ch}", self._pump_rf_events)

            # 펄스 펌프는 선택된 경우에만
            if self.dc_pulse and sel.get("dc_pulse", False):
                self._ensure_task_alive(f"Pump.DCPulse.{self.ch}", self._pump_dcpulse_events)
            if self.rf_pulse and sel.get("rf_pulse", False):
                self._ensure_task_alive(f"Pump.RFPulse.{self.ch}", self._pump_rfpulse_events)

            self._ensure_task_alive(f"Pump.OES.{self.ch}", self._pump_oes_events)

            self._bg_started = True
        finally:
            self._ensuring_bg = False

    # ──────────────────────────────────────────────────────────────
    # 디바이스 start/connect 보장(중복 호출 안전)
    # ──────────────────────────────────────────────────────────────
    def _ensure_devices_started(self) -> None:
        """MFC/IG는 start(), PLC는 connect()로 워치독/하트비트까지 기동."""
        # ✅ DevStart 태스크가 살아있으면 재생성하지 않음
        t = getattr(self, "_devstart_task", None)
        if isinstance(t, asyncio.Task) and (not t.done()):
            self._devices_started = True
            return

        if getattr(self, "_devices_started", False):
            return

        task = self._spawn_detached(self._start_devices_task(), store=True, name=f"DevStart.CH{self.ch}")
        if task is None:
            # create_task 실패(또는 다른 스레드에서 예약만 된 경우) → 다음 주기에 재시도 가능하게 둠
            self._devices_started = False
            self.append_log("MAIN", "[DevStart] task 생성 실패/지연 → 다음 주기에 재시도")
            return

        self._devstart_task = task
        self._devices_started = True

    async def _start_devices_task(self) -> None:
        # ✅ 개발자 모드: PLC/MFC/IG 자동 연결 생략 (연결 실패·재시도 로그 스팸 방지)
        #    Start 시 preflight는 그대로 동작하므로 시작 흐름 테스트에는 영향 없음
        try:
            from lib import config_common as _cc
            if getattr(_cc, "DEV_MODE", False):
                self.append_log("MAIN", "[DEV] 개발자 모드: 장치 자동 연결 생략")
                return
        except Exception:
            pass

        async def _maybe_start_or_connect(obj, label: str, *, log: bool = True):
            if not obj:
                return
            try:
                if self._is_dev_connected(obj):        # ★ 이미 연결됨
                    if log:
                        self.append_log(label, "already connected → skip")
                    return
                
                meth = getattr(obj, "start", None) or getattr(obj, "connect", None)
                if not callable(meth):
                    if log:
                        self.append_log(label, "start/connect 메서드 없음 → skip")
                    return
                
                # ⬇️ start/connect가 동기(blocking)여도 이벤트루프를 막지 않게: to_thread + timeout
                timeout_s = float(getattr(self.cfg.mod, "DEVICE_START_TIMEOUT_S", 20.0))

                try:
                    # meth() 호출 자체를 백그라운드 스레드에서 수행
                    res = await asyncio.wait_for(asyncio.to_thread(meth), timeout=timeout_s)
                except asyncio.TimeoutError:
                    if log:
                        self.append_log(label, f"{getattr(meth, '__name__', 'start/connect')} timeout({timeout_s}s) → skip")
                    return
                except Exception as e:
                    if log:
                        self.append_log(label, f"{getattr(meth, '__name__', 'start/connect')} 호출 실패: {e!r}")
                    return

                # meth가 coroutine을 “반환”하는 타입이면 여기서 await
                if inspect.isawaitable(res):
                    try:
                        await asyncio.wait_for(res, timeout=timeout_s)
                    except asyncio.TimeoutError:
                        if log:
                            self.append_log(label, f"{getattr(meth, '__name__', 'start/connect')} await timeout({timeout_s}s) → skip")
                        return

                if log:
                    self.append_log(label, f"{getattr(meth, '__name__', 'start/connect')} 호출 완료")

            except Exception as e:
                try:
                    name = meth.__name__  # type: ignore[attr-defined]
                except Exception:
                    name = "start/connect"
                if log:
                    self.append_log(label, f"{name} 실패: {e!r}")

        sel = getattr(self, "_run_select", None) or {}

        # PLC는 공유 → 소유자만 로그 출력 (비소유자는 연결 시도하되 로그 무음)
        await _maybe_start_or_connect(self.plc, "PLC", log=self._owns_plc)

        # 나머지는 기존대로 각 챔버에서 로그 출력
        await _maybe_start_or_connect(self.mfc, "MFC")
        await _maybe_start_or_connect(self.ig,  "IG")

        # 펄스 장비는 '이번 런에서 선택된 경우에만' 연결 시도
        if self.dc_pulse and sel.get("dc_pulse", False):
            await _maybe_start_or_connect(self.dc_pulse, "DCPulse")
        if self.rf_pulse and sel.get("rf_pulse", False):
            await _maybe_start_or_connect(self.rf_pulse, "RFPulse")

    # ------------------------------------------------------------------
    # 표시/입력/상태
    def _display_rf(self, for_p: Optional[float], ref_p: Optional[float]) -> None:
        if for_p is None or ref_p is None:
            self.append_log("MAIN", "for.p/ref.p 비어있음"); return
        self._set("forP_edit", f"{for_p:.2f}")
        self._set("refP_edit", f"{ref_p:.2f}")

    def _display_dc(self, power: Optional[float], voltage: Optional[float], current: Optional[float],
                    *, unit: int = 1) -> None:
        if power is None or voltage is None or current is None:
            self.append_log("MAIN", "P/V/I 비어있음"); return

        # DC2 미도입(=dc_power2 None)이면 기존 단일 표시 그대로 유지 (CH1 DC Pulse 경로 포함)
        if self.dc_power2 is None:
            self._set("Power_edit",   f"{power:.1f}")
            self._set("Voltage_edit", f"{voltage:.1f}")
            self._set("Current_edit", f"{current:.3f}")
            return

        u = 2 if int(unit) == 2 else 1
        self._dc_disp[u] = (float(power), float(voltage), float(current))
        self._render_dc_display()

    def _render_dc_display(self) -> None:
        """DC1|DC2 병기 렌더. 값이 없는 유닛은 '-' 로 표시."""
        d1 = self._dc_disp.get(1)
        d2 = self._dc_disp.get(2)

        def _pair(i: int, fmt: str) -> str:
            a = format(d1[i], fmt) if d1 else "-"
            b = format(d2[i], fmt) if d2 else "-"
            return f"{a}|{b}"

        self._set("Power_edit",   _pair(0, ".0f"))
        self._set("Voltage_edit", _pair(1, ".0f"))
        self._set("Current_edit", _pair(2, ".2f"))

    def _set_pulse_radio(self, leaf: str, checked: bool) -> None:
        """펄스 라디오 전용 세터.
        exclusive QButtonGroup의 라디오는 setChecked(False) 직접 호출이 무효이므로
        (Qt 사양) setExclusive 토글 트릭으로 해제한다. 그룹이 없으면 일반 setChecked 폴백.
        """
        btn = self._u(leaf)
        if btn is None:
            return
        grp = getattr(self.ui, f"{self.prefix}pulsePower_group", None) if getattr(self, "ui", None) else None
        with contextlib.suppress(Exception):
            if checked or grp is None:
                btn.setChecked(bool(checked))
                return
            try:
                grp.setExclusive(False)
                btn.setChecked(False)
            finally:
                grp.setExclusive(True)

    def _wire_pulse_radio_toggle(self) -> None:
        """(CH1/CH2 공통) 체크된 펄스 라디오를 다시 클릭하면 해제되도록 배선."""
        for leaf in ("rfPulsePower_checkbox", "dcPulsePower_checkbox"):
            btn = self._u(leaf)
            if btn is None or not hasattr(btn, "pressed"):
                continue
            # pressed는 토글 '이전' 상태를 캡처, clicked에서 이전에 이미 체크였으면 해제
            btn.pressed.connect(lambda b=btn: setattr(b, "_was_checked_before_click", bool(b.isChecked())))
            def _on_clicked(_checked=False, b=btn):
                if getattr(b, "_was_checked_before_click", False):
                    grp = getattr(self.ui, f"{self.prefix}pulsePower_group", None) if getattr(self, "ui", None) else None
                    try:
                        if grp is not None:
                            grp.setExclusive(False)
                        b.setChecked(False)
                    finally:
                        if grp is not None:
                            grp.setExclusive(True)
            btn.clicked.connect(_on_clicked)

    @staticmethod
    def _pulse_endpoint_of(dev) -> Optional[str]:
        """펄스 드라이버의 실효 엔드포인트 'host:port'. 해석 실패 시 None(가드는 fail-open)."""
        try:
            fn = getattr(dev, "_resolve_endpoint", None)
            if callable(fn):
                host, port = fn()
                return f"{str(host).strip()}:{int(port)}"
        except Exception:
            pass
        return None

    def _on_process_status_changed(self, running: bool) -> None:
        # ✅ 공정 종료 시 이 챔버가 점유한 펄스 엔드포인트 클레임을 일괄 해제
        #    (정상/실패/중단/preflight 실패 모두 이 함수를 지나므로 해제 누락 불가)
        if not running:
            with contextlib.suppress(Exception):
                runtime_state.release_pulse_endpoints(self.ch)

        # ✅ Start/Stop 버튼은 '상태와 무관하게 항상 활성화' (사용자 요구)
        b_start = self._u("Start_button"); b_stop = self._u("Stop_button")
        if b_start:
            with contextlib.suppress(Exception):
                b_start.setEnabled(True)
        if b_stop:
            with contextlib.suppress(Exception):
                b_stop.setEnabled(True)

        # ★ PLC 소유권 콜백은 running 값이 실제로 바뀐 경우에만 호출
        prev = getattr(self, "_last_running_state", None)
        if prev is None or prev != running:
            cb = getattr(self, "_notify_plc_owner", None)
            if callable(cb):
                try:
                    cb(self.ch if running else None)
                except Exception:
                    pass

        self._last_running_state = running

    # === 외부 공개: 현재 챔버 공정 실행 여부 ===
    @property
    def is_running(self) -> bool:
        try:
            return bool(self.process_controller.is_running)
        except Exception:
            return False

    def _apply_process_state_message(self, message: str) -> None:
        if getattr(self, "_last_state_text", None) == message:
            return
        self._last_state_text = message
        if self._w_state:
            self._w_state.setPlainText(message)

    def _fmt_hms(self, seconds: float) -> str:
        if seconds < 0:
            seconds = 0
        s = int(seconds)
        h, m, sec = s // 3600, (s % 3600) // 60, s % 60
        return f"{h:02d}:{m:02d}:{sec:02d}"

    def _set_state_text(self, text: str) -> None:
        self._last_state_text = str(text)
        if self._w_state:
            try:
                self._w_state.setPlainText(self._last_state_text)
            except Exception:
                pass

    # ------------------------------------------------------------------
    # 파일 로딩 / UI 반영
    def _connect_my_buttons(self) -> None:
        if not self._has_ui():
            self._set_default_ui_values()  # 필요 없으면 생략 가능
            return

        btn = self._u("Start_button")
        if btn: btn.clicked.connect(self._handle_start_clicked)

        btn = self._u("Stop_button")
        if btn: btn.clicked.connect(self._handle_stop_clicked)

        btn = self._u("processList_button")
        if btn:
            btn.clicked.connect(lambda: self._spawn_detached(self._handle_process_list_clicked_async()))

        if self._w_log:
            self._w_log.setMaximumBlockCount(2000)

        # ✅ Dep.Rate / Thickness 입력 시 Process Time 자동계산
        dep_w   = self._u("depRate_edit")
        thick_w = self._u("thickness_edit")
        if dep_w:
            dep_w.textChanged.connect(self._recalc_process_time)
        if thick_w:
            thick_w.textChanged.connect(self._recalc_process_time)

        self._set_default_ui_values()

    def _recalc_process_time(self) -> None:
        """Dep.Rate 또는 Thickness 변경 시 Process Time 자동계산."""
        try:
            dep   = float(self._get_text("depRate_edit")   or 0)
            thick = float(self._get_text("thickness_edit") or 0)
            if dep > 0 and thick > 0:
                t_min = thick / dep / 60.0
                # 시그널 루프 방지: processTime_edit 값이 이미 같으면 갱신 안 함
                current = self._get_text("processTime_edit")
                new_val = f"{t_min:.3f}"
                if current != new_val:
                    self._set("processTime_edit", new_val)
        except Exception:
            pass

    async def _handle_process_list_clicked_async(self) -> None:
        start_dir = (
            getattr(self, "_last_process_list_dir", "")
            or str(self.cfg._get("PROCESS_LIST_START_DIR", r"\\VanaM_NAS\VanaM_toShare"))
        )

        file_path = await self._aopen_file(
            caption=f"CH{self.ch} 프로세스 리스트 파일 선택",
            start_dir=start_dir,
            name_filter="Recipe Files (*.csv *.xlsx);;CSV Files (*.csv);;Excel Files (*.xlsx);;All Files (*)"
        )

        if file_path:
            with contextlib.suppress(Exception):
                self._last_process_list_dir = str(Path(file_path).parent)

        if not file_path:
            self.append_log("File", "파일 선택 취소")
            return

        self.append_log("File", f"선택된 파일: {file_path}")

        # ✅ NAS 파일 open을 executor로 분리 (CSV / XLSX 겸용)
        loop = asyncio.get_running_loop()
        def _load_csv_rows():
            rows = []

            def _push(row: dict):
                name = (str(row.get('Process_name') or row.get('#') or '').strip()
                        or f"공정 {len(rows)+1}")
                row['Process_name'] = name
                rows.append(row)

            if str(file_path).lower().endswith(".xlsx"):
                # ── Excel 레시피: 1행=헤더, 이후 행=공정 (첫 번째 시트만 사용, '범례' 등 추가 시트 무시) ──
                from openpyxl import load_workbook
                wb = load_workbook(file_path, read_only=True, data_only=True)
                try:
                    ws = wb.worksheets[0]
                    it = ws.iter_rows(values_only=True)
                    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
                    for vals in it:
                        row = {}
                        empty = True
                        for h, v in zip(header, vals):
                            if not h:
                                continue
                            s = "" if v is None else str(v).strip()
                            if s:
                                empty = False
                            row[h] = s
                        if not empty:
                            _push(row)
                finally:
                    wb.close()
            else:
                with open(file_path, mode='r', encoding='utf-8-sig', newline='') as csvfile:
                    reader = csv.DictReader(csvfile)
                    for row in reader:
                        _push(row)
            return rows

        try:
            rows = await asyncio.wait_for(
                loop.run_in_executor(None, _load_csv_rows),
                timeout=30.0
            )
        except asyncio.TimeoutError:
            self.append_log("File", f"CSV 로드 30초 timeout (NAS 응답 지연): {file_path}")
            return
        except Exception as e:
            self.append_log("File", f"파일 처리 오류: {e}")
            return

        self.process_queue = [cast(RawParams, r) for r in rows]
        self.current_process_index = -1
        if not self.process_queue:
            self.append_log("File", "파일에 공정이 없습니다.")
            return

        self.append_log("File", f"총 {len(self.process_queue)}개 공정 읽음.")
        self._update_ui_from_params(self.process_queue[0])

    def _update_ui_from_params(self, params: RawParams) -> None:
        if self._w_log:
            if getattr(self, "process_queue", None):
                total = len(self.process_queue); current = getattr(self, "current_process_index", -1) + 1
                self.append_log("UI", f"[CH{self.ch}] 자동 공정 ({current}/{total}) 준비: '{params.get('Process_name','')}'")
            else:
                self.append_log("UI", f"[CH{self.ch}] 단일 공정 UI 업데이트: '{params.get('process_note','')}'")

        _set = self._set

        # --- Pulse UI ---
        if self.ch == 1:
            # ✅ CH1: Pulse 입력(Edit/Freq/Duty)은 dcPulse* 1세트를 공용으로 사용한다.
            #    RF/DC 선택은 라디오(rfPulsePower_checkbox / dcPulsePower_checkbox)로 표시한다.
            def _is_true(v) -> bool:
                return str(v).strip().upper() in ("T", "TRUE", "1", "Y", "YES")

            def _pos_num(v) -> bool:
                try:
                    s = str(v).strip()
                    if s == "" or s.lower() == "nan":
                        return False
                    return float(s) > 0.0
                except Exception:
                    return False

            use_rf_flag = _is_true(params.get("use_rf_pulse", "F"))
            use_dc_flag = _is_true(params.get("use_dc_pulse", "F"))

            rf_power = params.get("rf_pulse_power", "0")
            rf_freq  = str(params.get("rf_pulse_freq", "")).strip()
            rf_duty  = str(params.get("rf_pulse_duty_cycle") or params.get("rf_pulse_duty") or "").strip()

            dc_power = params.get("dc_pulse_power", "0")
            dc_freq  = str(params.get("dc_pulse_freq", "")).strip()
            dc_duty  = str(params.get("dc_pulse_duty_cycle") or params.get("dc_pulse_duty") or "").strip()

            rf_requested = use_rf_flag or _pos_num(rf_power) or (rf_freq not in ("", "0")) or (rf_duty not in ("", "0"))
            dc_requested = use_dc_flag or _pos_num(dc_power) or (dc_freq not in ("", "0")) or (dc_duty not in ("", "0"))

            # ✅ 동시 사용 허용: 체크박스는 각자 요청대로 표시, 공유 입력칸은 RF 우선 값 표시
            self._set_pulse_radio("rfPulsePower_checkbox", rf_requested)
            self._set_pulse_radio("dcPulsePower_checkbox", dc_requested)
            if rf_requested and dc_requested:
                self.append_log("UI", "[CH1] RF/DC Pulse 동시 레시피 → 입력칸(공유)에는 RF 값 표시, 실행은 각자 값 사용")

            if rf_requested:
                power, freq, duty = rf_power, rf_freq, rf_duty
            elif dc_requested:
                power, freq, duty = dc_power, dc_freq, dc_duty
            else:
                power, freq, duty = "0", "", ""

            _set("dcPulsePower_edit", power)
            _set("dcPulseFreq_edit",      "" if str(freq).strip() in ("", "0", "nan") else str(freq).strip())
            _set("dcPulseDutyCycle_edit", "" if str(duty).strip() in ("", "0", "nan") else str(duty).strip())

        else:
            # CH2: 입력칸은 rfPulse* 1세트를 RF/DC 공유(CH1과 동일 UX), 선택은 라디오로 표시
            use_rfp = params.get('use_rf_pulse', 'F') == 'T'
            use_dcp = params.get('use_dc_pulse', 'F') == 'T'
            self._set_pulse_radio("rfPulsePower_checkbox", use_rfp)
            self._set_pulse_radio("dcPulsePower_checkbox", use_dcp)
            if use_rfp and use_dcp:
                self.append_log("UI", "[CH2] RF/DC Pulse 동시 레시피 → 입력칸(공유)에는 RF 값 표시, 실행은 각자 값 사용")

            if use_rfp:
                pw  = str(params.get('rf_pulse_power', '0'))
                frq = str(params.get('rf_pulse_freq', '')).strip()
                dty = str(params.get('rf_pulse_duty_cycle') or params.get('rf_pulse_duty') or '').strip()
            elif use_dcp:
                pw  = str(params.get('dc_pulse_power', '0'))
                frq = str(params.get('dc_pulse_freq', '')).strip()
                dty = str(params.get('dc_pulse_duty_cycle') or params.get('dc_pulse_duty') or '').strip()
            else:
                pw, frq, dty = "0", "", ""

            _set("rfPulsePower_edit",     pw)
            _set("rfPulseFreq_edit",      '' if frq in ('', '0') else frq)
            _set("rfPulseDutyCycle_edit", '' if dty in ('', '0') else dty)

        # DC-Power
        _set("dcPower_checkbox", params.get('use_dc_power', 'F') == 'T')
        _set("dcPower_edit", params.get('dc_power', '0'))
        _set("dcPower2_checkbox", params.get('use_dc_power2', 'F') == 'T')
        _set("dcPower2_edit", params.get('dc_power2', '0'))

        # RF-Power
        _set("rfPower_checkbox", params.get('use_rf_power', 'F') == 'T')
        _set("rfPower_edit",     params.get('rf_power', '0'))
        
        # ✅ Integration Time 입력칸을 'Process Name' 입력으로 재활용
        #    (CSV 자동 공정: Process_name 표시 / UI 수동 공정: 사용자가 입력)
        _set("integrationTime_edit", params.get('Process_name', params.get('process_note', '')))
        _set("arFlow_edit", params.get('Ar_flow', '0'))
        _set("o2Flow_edit", params.get('O2_flow', '0'))
        _set("n2Flow_edit", params.get('N2_flow', '0'))
        _set("workingPressure_edit", params.get('working_pressure', '0'))
        _set("basePressure_edit", params.get('base_pressure', '0'))
        _set("shutterDelay_edit", params.get('shutter_delay', '0'))

        # ✅ dep.rate: CSV 헤더가 "dep.rate"(점) 또는 "dep_rate"(밑줄) 둘 다 지원
        _dep_rate_val = str(params.get('dep_rate', '') or params.get('dep.rate', '') or '')
        _set("depRate_edit",   _dep_rate_val)
        _set("thickness_edit", str(params.get('thickness', '') or ''))
        # ✅ processTime_edit: 명시적 값이 있을 때만 덮어씀
        #    (없으면 위 thickness_edit textChanged → _recalc_process_time이 계산한 값 유지)
        _pt = params.get('process_time', 0)
        if float(_pt or 0) > 0:
            _set("processTime_edit", _pt)

        _set("G1_checkbox", params.get('gun1', 'F') == 'T')
        _set("G2_checkbox", params.get('gun2', 'F') == 'T')
        _set("G3_checkbox", params.get('gun3', 'F') == 'T')
        _set("Ar_checkbox", params.get('Ar', 'F') == 'T')
        _set("O2_checkbox", params.get('O2', 'F') == 'T')
        _set("N2_checkbox", params.get('N2', 'F') == 'T')
        _set("mainShutter_checkbox", params.get('main_shutter', 'F') == 'T')
        _set("powerSelect_checkbox", params.get('power_select', 'F') == 'T')

        # ---- CH1: 단일 타겟 위젯에 한 번만 세팅 ----
        if self.ch == 1:
            name = str(params.get('G1 Target', '')).strip()
            _set("g1Target_name", name)
            if not name:
                asyncio.ensure_future(self._load_gun_targets())
        else:
            _set("g1Target_name", str(params.get('G1 Target', '')).strip())
            _set("g2Target_name", str(params.get('G2 Target', '')).strip())
            _set("g3Target_name", str(params.get('G3 Target', '')).strip())
            if not any([
                str(params.get('G1 Target', '')).strip(),
                str(params.get('G2 Target', '')).strip(),
                str(params.get('G3 Target', '')).strip(),
            ]):
                asyncio.ensure_future(self._load_gun_targets())

    def _set(self, leaf: str, v: Any) -> None:
        w = self._u(leaf)
        if w is None:
            return
        try:
            if hasattr(w, "setChecked"):
                w.setChecked(bool(v))
                return

            if hasattr(w, "setValue"):
                try:
                    w.setValue(v if isinstance(v, (int, float)) else float(str(v)))
                except Exception:
                    pass
                else:
                    return

            s = str(v)
            if hasattr(w, "setPlainText"):
                w.setPlainText(s)
                return
            if hasattr(w, "setText"):
                w.setText(s)
                return
        except Exception as e:
            self.append_log("UI", f"_set('{leaf}') 실패: {e!r}")

    # ✅ Gate(밸브) 인터락: 시작하려는 챔버의 Gate가 CLOSED인지 확인
    async def _check_gate_closed_before_start(self) -> bool:
        """
        Gate lamp 기준 상태 판정.
        - closed: CLOSE_LAMP=True & OPEN_LAMP=False
        - open  : OPEN_LAMP=True  & CLOSE_LAMP=False
        - moving_or_unknown: 둘 다 False
        - invalid_both_true: 둘 다 True (배선/맵/PLC 로직 이상 가능)

        시작 조건:
        - 'closed'일 때만 True
        - 그 외는 모두 False(시작 차단)
        """
        self._gate_fail_reason: str = "unknown"

        if not getattr(self, "plc", None):
            self._gate_fail_reason = "PLC 연결 없음"
            self.append_log("MAIN", f"[CH{self.ch}] PLC 없음 → Gate 상태 확인 불가 → 시작 차단")
            return False

        ch = int(getattr(self, "ch", 0) or 0)
        if ch not in (1, 2):
            self._gate_fail_reason = f"잘못된 CH 번호 ({ch})"
            self.append_log("MAIN", f"[CH{self.ch}] 잘못된 CH={ch} → 시작 차단")
            return False

        open_key = f"G_V_{ch}_OPEN_LAMP"
        close_key = f"G_V_{ch}_CLOSE_LAMP"

        gate_retry_count = int(self.cfg._get("CHAMBER_GATE_RECHECK_COUNT", 5))
        gate_read_timeout_s = float(self.cfg._get("CHAMBER_GATE_READ_TIMEOUT_S", 0.6))
        gate_retry_interval_s = float(self.cfg._get("CHAMBER_GATE_RECHECK_INTERVAL_S", 0.2))

        # (선택) 전이 상태(moving)일 때 잠깐만 재확인(짧게)
        for _ in range(gate_retry_count):  # 최대 gate_retry_count 회 재시도
            try:
                open_lamp = await asyncio.wait_for(self.plc.read_bit(open_key), timeout=gate_read_timeout_s)
                close_lamp = await asyncio.wait_for(self.plc.read_bit(close_key), timeout=gate_read_timeout_s)
                open_lamp = bool(open_lamp)
                close_lamp = bool(close_lamp)

            except KeyError as e:
                self._gate_fail_reason = f"PLC 주소맵 키 없음 ({e})"
                self.append_log("MAIN", f"[CH{self.ch}] PLC 주소맵에 gate lamp 키 없음: {e} → 시작 차단")
                return False
            except Exception as e:
                self.append_log("MAIN", f"[CH{self.ch}] Gate lamp 읽기 실패: {type(e).__name__}: {e} → 재시도 중")
                await asyncio.sleep(gate_retry_interval_s)
                continue

            if close_lamp and (not open_lamp):
                # ✅ 정상: CLOSED
                return True

            if open_lamp and (not close_lamp):
                # ❌ OPEN
                self._gate_fail_reason = "Gate OPEN 상태"
                self.append_log("MAIN", f"[CH{self.ch}] Gate 상태=OPEN (open={open_lamp}, close={close_lamp})")
                return False

            if open_lamp and close_lamp:
                # ❌ 비정상(둘 다 TRUE)
                self._gate_fail_reason = "Gate lamp 이상 (OPEN/CLOSE 모두 True)"
                self.append_log("MAIN", f"[CH{self.ch}] Gate lamp 이상(OPEN/CLOSE 모두 TRUE) (open={open_lamp}, close={close_lamp})")
                return False

            # 둘 다 False면 moving/unknown → 잠깐 기다렸다가 재확인
            await asyncio.sleep(gate_retry_interval_s)

        self._gate_fail_reason = "Gate 이동 중 / 상태 불명 (재시도 초과)"
        self.append_log("MAIN", f"[CH{self.ch}] Gate 상태=moving_or_unknown (OPEN/CLOSE 모두 FALSE) → 시작 차단")
        return False

    async def _check_main_valve_open_before_start(self) -> bool:
        """
        Main Valve 열림 여부 판정 (PLC 미수정 — 프로그램에서 재구성).
        래더상 M_V_OUT = M_V_SW AND M_V_인터락 (일반 코일)이므로,
        두 코일을 읽어 둘 다 True면 실제 밸브 OUT=True(열림)와 동일하다.
        인터락 조건: AIR · GAUGE_A · ROTARY_PUMP · FORELINE · TURBO_START · /VENT
        → 둘 다 True면 '터보 가동 + 벤트 아님 + 압력 낮음'까지 함께 보장.

        시작 조건:
        - sw=True & interlock=True 일 때만 True (그 외 시작 차단)
        """
        self._mv_fail_reason: str = "unknown"

        if not getattr(self, "plc", None):
            self._mv_fail_reason = "PLC 연결 없음"
            self.append_log("MAIN", f"[CH{self.ch}] PLC 없음 → Main Valve 상태 확인 불가 → 시작 차단")
            return False

        ch = int(getattr(self, "ch", 0) or 0)
        if ch not in (1, 2):
            self._mv_fail_reason = f"잘못된 CH 번호 ({ch})"
            self.append_log("MAIN", f"[CH{self.ch}] 잘못된 CH={ch} → 시작 차단")
            return False

        sw_key = f"M_V_{ch}_SW"
        itlk_key = f"M_V_{ch}_인터락"

        mv_retry_count = int(self.cfg._get("CHAMBER_MAINVALVE_RECHECK_COUNT", 5))
        mv_read_timeout_s = float(self.cfg._get("CHAMBER_MAINVALVE_READ_TIMEOUT_S", 0.6))
        mv_retry_interval_s = float(self.cfg._get("CHAMBER_MAINVALVE_RECHECK_INTERVAL_S", 0.2))

        last_sw = False
        last_itlk = False

        for _ in range(mv_retry_count):
            try:
                sw = await asyncio.wait_for(self.plc.read_bit(sw_key), timeout=mv_read_timeout_s)
                itlk = await asyncio.wait_for(self.plc.read_bit(itlk_key), timeout=mv_read_timeout_s)
                sw = bool(sw)
                itlk = bool(itlk)
            except KeyError as e:
                self._mv_fail_reason = f"PLC 주소맵 키 없음 ({e})"
                self.append_log("MAIN", f"[CH{self.ch}] PLC 주소맵에 Main Valve 키 없음: {e} → 시작 차단")
                return False
            except Exception as e:
                self.append_log("MAIN", f"[CH{self.ch}] Main Valve 읽기 실패: {type(e).__name__}: {e} → 재시도 중")
                await asyncio.sleep(mv_retry_interval_s)
                continue

            if sw and itlk:
                # ✅ M_V_OUT = SW AND 인터락 = True → Main Valve 열림
                return True

            last_sw, last_itlk = sw, itlk
            # 아직 열림 아님(전이 중일 수 있음) → 잠깐 후 재확인
            await asyncio.sleep(mv_retry_interval_s)

        if not last_itlk:
            self._mv_fail_reason = "Main Valve 인터락 불충족 (벤트 상태이거나 터보/러프펌프/압력 조건 미충족)"
        elif not last_sw:
            self._mv_fail_reason = "Main Valve가 열려있지 않음 (인터락 충족, 밸브 미개방)"
        else:
            self._mv_fail_reason = "Main Valve 상태 불명"
        self.append_log("MAIN", f"[CH{self.ch}] Main Valve 미개방 (sw={last_sw}, interlock={last_itlk}) → 시작 차단")
        return False
    
    async def _pulse_reconnect_safe(
        self,
        dev: Any,
        label: str,
        host: str,
        port: int,
        *,
        timeout_s: float | None = None,
    ) -> None:
        """
        STOP 직후 다음 런에서 pulse reconnect 중 내부 watchdog cancel이
        CancelledError로 새는 경우를 '조용한 stage cancel'이 아니라
        '명시적인 reconnect 실패'로 바꿔준다.
        """
        if timeout_s is None:
            timeout_s = float(self.cfg._get("CHAMBER_PULSE_RECONNECT_TIMEOUT_S", 2.0))

        if not dev or not hasattr(dev, "set_endpoint_reconnect"):
            return

        try:
            self.append_log("MAIN", f"[Runner] {label} reconnect 시작: {host}:{port}")
            await asyncio.wait_for(dev.set_endpoint_reconnect(host, port), timeout=timeout_s)
            self.append_log("MAIN", f"[Runner] {label} reconnect 완료")
        except asyncio.TimeoutError:
            raise RuntimeError(f"{label} reconnect timeout({timeout_s:.1f}s)")
        except asyncio.CancelledError:
            # STOP / stage-switch로 의도적으로 취소된 경우는 그대로 올려보낸다.
            if getattr(self, "_expected_stage_cancel_task", None) is asyncio.current_task():
                raise

            self.append_log(
                "MAIN",
                f"[Runner] {label} reconnect 중 예상치 못한 CancelledError 발생"
            )
            raise RuntimeError(f"{label} reconnect cancelled unexpectedly")
        except Exception as e:
            raise RuntimeError(f"{label} reconnect failed: {e!r}")

    async def _start_after_preflight(self, params: NormParams, run_gen: int) -> None:
        try:
            # ✅ 더 최신 Start가 들어오면(세대 불일치) 이 태스크는 조용히 종료
            if int(getattr(self, "_active_run_gen", 0)) != int(run_gen):
                return

            # ✅ 이전 런 cleanup 제한 플래그를 클리어 (정상 종료 후 다음 런 시작을 위해)
            self._pending_device_cleanup = False

            # ------------------------------------------------------------
            # TEST MODE : preflight/인터락/chuck/장비연결 전부 스킵
            # ------------------------------------------------------------
            if bool(params.get("test_mode", False)):
                # ✅ TEST MODE에서는 장비 자동연결/워치독을 절대 올리지 않음
                self._auto_connect_enabled = False
                self._run_select = None

                time_str = str(params.get("time", "")).strip()
                dur_s = float(params.get("test_duration_sec", 0.0) or 0.0)

                if dur_s <= 0 and time_str:
                    dur_s = self._parse_duration_seconds(time_str.lower())
                    params["test_duration_sec"] = dur_s

                if dur_s <= 0:
                    try:
                        dur_s = float(params.get("process_time", 0.0)) * 60.0
                    except Exception:
                        dur_s = 0.0

                dur_s = max(1.0, float(dur_s))

                # ✅ 카드/로그용 분 단위
                params.setdefault("process_time", round(dur_s / 60.0, 3))
                params.setdefault("process_note", params.get("Process_name") or "TEST")

                note = params.get("process_note") or "TEST"
                self.append_log("MAIN", f"[TEST MODE] '{note}' 장비 제어 스킵 / {dur_s:.1f}s 시뮬레이션")
                self._host_report_start(True, f"TEST MODE: {time_str or f'{dur_s:.0f}s'}")

                # ✅ 상태 RUNNING (UI/상태/구글챗 흐름은 정상 공정과 동일)
                self._on_process_status_changed(True)

                # ✅ 더 최신 Start가 들어오면(세대 불일치) start_process 호출 금지
                if int(getattr(self, "_active_run_gen", 0)) != int(run_gen):
                    return

                # ✅ 핵심: ProcessController가 TEST MODE(DELAY) 시퀀스로 실행
                self.process_controller.start_process(params)
                return
            # ------------------------------------------------------------
            
            # ✅ REAL MODE부터 여기서 장비 연결/백그라운드 허용
            self._auto_connect_enabled = True

            # ✅ 이번 런에서 실제로 사용할 펄스만 표시(IG/MFC는 항상 연결이므로 제외)
            def _pos(v) -> bool:
                try:
                    return float(v) > 0.0
                except Exception:
                    return False

            # ✅ "요청" 판정(체크박스 + 값)
            rf_requested = (
                bool(params.get("use_rf_pulse", False))
                or _pos(params.get("rf_pulse_power", 0.0))
                or (params.get("rf_pulse_freq") is not None)
                or (params.get("rf_pulse_duty") is not None)
            )

            dc_requested = (
                bool(params.get("use_dc_pulse", False))
                or _pos(params.get("dc_pulse_power", 0.0))
                or (params.get("dc_pulse_freq") is not None)
                or (params.get("dc_pulse_duty") is not None)
            )

            use_dc_pulse = bool(dc_requested) and self.supports_dc_pulse
            use_rf_pulse = bool(rf_requested) and self.supports_rf_pulse

            self._run_select = {
                "dc_pulse": use_dc_pulse,
                "rf_pulse": use_rf_pulse,
            }

            # ✅ Pulse는 각 장비의 "자기 포트" 설정을 그대로 사용한다. (덮어쓰기 금지)
            # - DC Pulse  : cfg.DCPULSE_TCP
            # - RF Pulse  : RFPulseAsync 내부 설정(또는 cfg.RFPULSE_TCP가 있으면 그 값)

            pulse_reconnect_timeout_s = float(self.cfg._get("CHAMBER_PULSE_RECONNECT_TIMEOUT_S", 2.0))

            # (선택) DC는 필요시 재연결만 수행
            if use_dc_pulse and self.dc_pulse and hasattr(self.dc_pulse, "set_endpoint_reconnect"):
                host, port = self.cfg.DCPULSE_TCP
                await self._pulse_reconnect_safe(
                    self.dc_pulse,
                    "DC-Pulse",
                    host,
                    port,
                    timeout_s=pulse_reconnect_timeout_s,
                )

            # ✅ RF는 DCPULSE_TCP로 절대 덮어쓰지 않는다.
            #    - RFPulseAsync가 내부적으로 RF 포트를 알고 있으면: 아무 것도 안 해도 됨
            #    - cfg.RFPULSE_TCP가 준비되어 있으면: 그 값으로만 reconnect
            if use_rf_pulse and self.rf_pulse and hasattr(self.rf_pulse, "set_endpoint_reconnect"):
                rf_tcp = getattr(self.cfg, "RFPULSE_TCP", None)
                if rf_tcp:
                    host, port = rf_tcp
                    await self._pulse_reconnect_safe(
                        self.rf_pulse,
                        "RF-Pulse",
                        host,
                        port,
                        timeout_s=pulse_reconnect_timeout_s,
                    )
                # else: RFPulseAsync 내부 설정을 그대로 사용

            self._ensure_background_started()

            # ✅ Start 클릭 시점에 OES init을 백그라운드로 '미리' 수행
            self._kick_oes_init_background(force=False)

            self._on_process_status_changed(True)

            # ✅ 공유 펄스 장비 가드: '같은 엔드포인트'를 다른 챔버가 사용 중이면 시작 거부
            #    - 엔드포인트가 다르면(향후 장비 증설) 동시 실행 허용
            #    - 같은 챔버의 재클레임은 허용, 해석 실패 시 가드 스킵(fail-open)
            #    - 해제는 _on_process_status_changed(False)에서 일괄 수행
            _pulse_claims: list[tuple[str, str]] = []
            if use_dc_pulse and self.dc_pulse:
                _ep = self._pulse_endpoint_of(self.dc_pulse)
                if _ep:
                    _pulse_claims.append(("DC-Pulse", _ep))
            if use_rf_pulse and self.rf_pulse:
                _ep = self._pulse_endpoint_of(self.rf_pulse)
                if _ep:
                    _pulse_claims.append(("RF-Pulse", _ep))
            for _kind_nm, _ep in _pulse_claims:
                _ok_claim, _owner = runtime_state.claim_pulse_endpoint(_ep, _kind_nm, self.ch)
                if not _ok_claim:
                    _o_ch = _owner.get("ch") if isinstance(_owner, dict) else "?"
                    _o_kind = _owner.get("kind") if isinstance(_owner, dict) else "?"
                    _msg = f"{_kind_nm} 시작 불가: 동일 주소({_ep}) 장비를 CH{_o_ch}({_o_kind}) 공정이 사용 중입니다."
                    self.append_log("MAIN", _msg)
                    self._post_warning("공유 장비 사용 중", _msg)
                    with contextlib.suppress(Exception):
                        self._host_report_start(False, _msg)
                    with contextlib.suppress(Exception):
                        runtime_state.set_error("chamber", self.ch, _msg)
                        runtime_state.mark_finished("chamber", self.ch)
                    self._on_process_status_changed(False)   # ← 부분 클레임도 여기서 해제됨
                    return

            timeout_no_pulse = float(self.cfg._get("CHAMBER_PREFLIGHT_TIMEOUT_S", 8.0))
            timeout_with_pulse = float(self.cfg._get("CHAMBER_PREFLIGHT_TIMEOUT_WITH_PULSE_S", 10.0))
            timeout = timeout_with_pulse if (use_dc_pulse or use_rf_pulse) else timeout_no_pulse
            
            ok, failed = await self._preflight_connect(params, timeout_s=timeout)

            if not ok:
                fail_list = ", ".join(failed) if failed else "알 수 없음"
                self.append_log("MAIN", f"필수 장비 연결 실패: {fail_list} → 시작 중단")
                self._post_critical(
                    "장비 연결 실패",
                    "다음 장비 연결을 확인하지 못했습니다:\n"
                    f" - {fail_list}\n\n케이블/전원/포트 설정 확인 후 재시도"
                )

                # ✅ 자동 재연결 자체 차단 (이후 _ensure_background_started 가 장치 start 못 올리도록)
                self._auto_connect_enabled = False

                # ✅ 이미 올라가 있던 워치독/연결 태스크 완전 정지
                try:
                    await self._stop_device_watchdogs(light=False)
                except Exception:
                    pass

                # (선택) 폴링 상태도 명시적으로 내려줌 — 없어도 무방
                with contextlib.suppress(Exception):
                    if not self._skip_mfc_finalize_due_to_pc():
                        self.mfc.set_process_status(False)
                with contextlib.suppress(Exception):
                    if hasattr(self.ig, "set_process_status"): self.ig.set_process_status(False)
                with contextlib.suppress(Exception):
                    if self.dc_pulse and hasattr(self.dc_pulse, "set_process_status"):
                        self.dc_pulse.set_process_status(False)

                self._on_process_status_changed(False)

                # ✅ 실패 종료 상태를 runtime_state에 명확히 기록
                try:
                    runtime_state.set_error("chamber", self.ch, f"preflight connect failed: {fail_list}")
                    runtime_state.mark_finished("chamber", self.ch)
                except Exception:
                    pass

                raise RuntimeError(f"preflight connect failed: {fail_list}")
            
            # ✅ Gate(밸브) 인터락: 시작하려는 챔버의 Gate가 열려 있으면 공정 시작을 차단
            # - Plasma Cleaning은 Gate를 열고 진행하므로, Gate Open 상태면 Sputter Start 금지
            ok_gate = await self._check_gate_closed_before_start()
            if not ok_gate:
                _gate_reason = getattr(self, '_gate_fail_reason', 'Gate 상태 이상')
                self.append_log("MAIN", f"[CH{self.ch}] Gate 체크 실패 ({_gate_reason}) → 공정 시작 차단")

                # ↓ 이 블록 추가
                if self.chat:
                    with contextlib.suppress(Exception):
                        _gate_note = params.get("process_note") or params.get("Process_name") or "알 수 없음"
                        self.chat.notify_error_event(
                            f"CH{self.ch}",
                            "E301",
                            f"Gate 체크 실패: {_gate_reason} → 공정 시작 차단 (공정: {_gate_note})",
                        )
                        self.chat.flush()

                # ✅ 이미 올라간 연결/펌프/점유를 정리(특히 RF-Pulse 점유 해제 목적)
                self._auto_connect_enabled = False
                with contextlib.suppress(Exception):
                    await self._stop_device_watchdogs(light=False)

                self._on_process_status_changed(False)
                with contextlib.suppress(Exception):
                    runtime_state.set_error("chamber", self.ch, f"gate check failed: {_gate_reason}")
                    runtime_state.mark_finished("chamber", self.ch)

                raise RuntimeError(f"gate check failed: {_gate_reason}")
            
            # ✅ Main Valve 인터락: Main Valve가 열려 있을 때만 공정 시작 허용
            # - PLC 미수정. 래더상 M_V_OUT = M_V_SW AND M_V_인터락 → 두 코일을 읽어 재구성.
            # - 벤트/터보정지/압력높음이면 인터락이 깨져 차단됨.
            ok_mv = await self._check_main_valve_open_before_start()
            if not ok_mv:
                _mv_reason = getattr(self, '_mv_fail_reason', 'Main Valve 상태 이상')
                self.append_log("MAIN", f"[CH{self.ch}] Main Valve 체크 실패 ({_mv_reason}) → 공정 시작 차단")

                if self.chat:
                    with contextlib.suppress(Exception):
                        _mv_note = params.get("process_note") or params.get("Process_name") or "알 수 없음"
                        self.chat.notify_error_event(
                            f"CH{self.ch}",
                            "E301",
                            f"Main Valve 체크 실패: {_mv_reason} → 공정 시작 차단 (공정: {_mv_note})",
                        )
                        self.chat.flush()

                self._auto_connect_enabled = False
                with contextlib.suppress(Exception):
                    await self._stop_device_watchdogs(light=False)

                self._on_process_status_changed(False)
                with contextlib.suppress(Exception):
                    runtime_state.set_error("chamber", self.ch, f"main valve check failed: {_mv_reason}")
                    runtime_state.mark_finished("chamber", self.ch)

                raise RuntimeError(f"main valve check failed: {_mv_reason}")
            
            # ★ 추가: 공정 시작 직전 Chuck 위치 선행 설정
            self._run_chuck_position = str(params.get("chuck_position") or "").strip().lower()
            self._run_warnings = []

            ok_chuck = await self._set_chuck_position_if_needed(params)
            if not ok_chuck:
                pos = self._run_chuck_position
                warn = f"Chuck 위치 이동 실패 (target='{pos}')" if pos else "Chuck 위치 이동 실패"
                self.append_log("MAIN", f"⚠️ {warn} → 공정은 계속 진행")
                self._run_warnings.append(warn)
                # ✅ 여기서 실패처리/return/큐 fail 처리/critical/error 기록 전부 하지 않음

            if not params.get("chuck_position"):
                if self.plc:
                    actual = await self._read_chuck_pos(self.ch)
                    if actual != "unknown":
                        params["chuck_position"] = actual

            self._last_polling_targets = None
            self.append_log("MAIN", "장비 연결 확인 완료 → 공정 시작")

            with contextlib.suppress(Exception):
                runtime_state.mark_started("chamber", self.ch)

            self.process_controller.start_process(params)

        except Exception as e:
            note = params.get("process_note", "알 수 없는")
            msg = f"오류: '{note}' 시작 실패. ({e})"
            self.append_log("MAIN", msg)
            self._post_critical("오류", msg)

            # ✅ Host start 요청이 걸려있으면 즉시 실패 응답(대기/타임아웃 방지)
            with contextlib.suppress(Exception):
                self._host_report_start(False, msg)

            # ✅ 자동 재연결 차단 + 워치독/점유 정리(RF-Pulse 점유 해제 포함)
            self._auto_connect_enabled = False
            with contextlib.suppress(Exception):
                await self._stop_device_watchdogs(light=False)

            # ✅ 예외로 비정상 종료 → error 표시(정상 종료만 idle)
            with contextlib.suppress(Exception):
                runtime_state.set_error("chamber", self.ch, msg)
                runtime_state.mark_finished("chamber", self.ch)

            self._on_process_status_changed(False)
            raise RuntimeError(msg)

    def _kick_oes_init_background(self, *, force: bool = True) -> None:
        """Start 버튼을 눌렀을 때 OES init을 백그라운드로 미리 수행한다."""
        if not getattr(self, "oes", None):
            return

        t = getattr(self, "_oes_init_task", None)
        if t and not t.done():
            return

        async def _run():
            try:
                self.append_log(f"OES{self.ch}", "[OES] init (background) begin")
                self._ensure_background_started()

                # ✅ 메인 timeout은 워커 내부 timeout(25s)보다 길어야 함
                oes_init_timeout_s = float(self.cfg._get("CHAMBER_OES_INIT_TIMEOUT_S", 30.0))

                ok = await self.oes.initialize_device(
                    timeout_s=oes_init_timeout_s,
                    force=force,
                )
                self._oes_initialized = bool(ok)

                if ok:
                    self.append_log(f"OES{self.ch}", f"[OES] init (background) OK (timeout={oes_init_timeout_s:.1f}s)")
                else:
                    err = getattr(self.oes, "_init_error", None) or "unknown"
                    self.append_log(
                        f"OES{self.ch}",
                        f"[OES] init (background) FAIL: {err} (timeout={oes_init_timeout_s:.1f}s)"
                    )

            except Exception as e:
                self._oes_initialized = False
                self.append_log(f"OES{self.ch}", f"[OES] init (background) EXC: {type(e).__name__}: {e}")

        self._oes_init_task = self._spawn_detached(_run(), store=True, name=f"OES.init.{self.ch}")

    async def _wait_device_connected(self, dev: object, name: str, timeout_s: float) -> bool:
        """장비 연결 여부가 True가 될 때까지 대기(타임아웃).
        ⚠️ qasync(QEventLoop) 환경에서 loop.time()이 간헐적으로 정지/비정상 값을 반환하는 케이스가 있어,
        프리플라이트가 영원히 끝나지 않는 현상을 막기 위해 time.monotonic() 기반으로 구현한다.
        """
        t0 = time.monotonic()
        while True:
            if self._is_dev_connected(dev):
                return True

            if (time.monotonic() - t0) >= float(timeout_s):
                self.append_log(name, f"연결 확인 실패(타임아웃 {timeout_s:.1f}s)")
                return False

            await asyncio.sleep(0.2)

    async def _preflight_connect(self, params: Mapping[str, Any], timeout_s: float = 8.0) -> tuple[bool, list[str]]:
        need: list[tuple[str, object]] = [("PLC", self.plc), ("MFC", self.mfc), ("IG", self.ig)]

        use_dc_pulse = bool(params.get("use_dc_pulse", False))
        use_rf_pulse = bool(params.get("use_rf_pulse", False))

        if use_dc_pulse and self.dc_pulse:
            need.append(("DC-Pulse", self.dc_pulse))
        if use_rf_pulse and self.rf_pulse:
            need.append(("RF-Pulse", self.rf_pulse))

        stop_evt = asyncio.Event()
        prog_task = asyncio.create_task(self._preflight_progress_log(need, stop_evt))

        try:
            results = await asyncio.gather(
                *[self._wait_device_connected(dev, name, timeout_s) for name, dev in need],
                return_exceptions=False
            )
        finally:
            stop_evt.set()
            with contextlib.suppress(Exception):
                await prog_task

        failed = [name for (name, _), ok in zip(need, results) if not ok]
        ok = len(failed) == 0
    
        # ✅ 여기서만 Host로 성공/실패 신호를 보낸다
        if ok:
            self._host_report_start(True, "preflight OK")
        else:
            self._host_report_start(False, "장비 연결 실패: " + ", ".join(failed))

        return ok, failed
    
    async def _read_chuck_pos(self, ch: int) -> str:
        try:
            up  = bool(await self.plc.read_bit(f"Z{ch}_UP_LOCATION"))
            mid = bool(await self.plc.read_bit(f"Z{ch}_MID_LOCATION"))
            dn  = bool(await self.plc.read_bit(f"Z{ch}_DOWN_LOCATION"))
            if int(up) + int(mid) + int(dn) == 1:
                return "up" if up else ("mid" if mid else "down")
        except Exception:
            pass
        return "unknown"
    
    async def _set_chuck_position_if_needed(self, params: Mapping[str, Any]) -> bool:
        """
        레시피에 chuck_position 값이 있으면(공란 제외) 공정 시작 전에 1회만 Chuck 위치를 조정.

        handlers.py 의 chuck_up/chuck_down 과 동일한 구조:
        - Z_M_P_{CH}_SW (Z-POWER) ON 유지
        - 방향 스위치(Z_M_P_{CH}_CW/MID/CCW_SW) ON 유지
        - Z{CH}_*_LOCATION 램프를 폴링해서 목표 위치 도달 여부 확인
        - 타임아웃/예외 시에도 스위치는 반드시 OFF
        """
        pos = str(params.get("chuck_position") or "").strip().lower()
        if not pos:
            # 공란이면 스킵
            return True

        ch = 1 if int(getattr(self, "ch", 1)) != 2 else 2

        # 허용 값 체크
        if pos not in ("up", "mid", "down"):
            self.append_log("PLC", f"[CH{self.ch}] 알 수 없는 chuck_position='{pos}' → 스킵")
            return True

        # POWER / 방향 스위치 / 위치 램프 매핑 (handlers.py와 동일한 구조)
        power_sw = f"Z_M_P_{ch}_SW"
        if pos == "up":
            move_sw = f"Z_M_P_{ch}_CW_SW"
            lamp_bit = f"Z{ch}_UP_LOCATION"
        elif pos == "mid":
            move_sw = f"Z_M_P_{ch}_MID_SW"
            lamp_bit = f"Z{ch}_MID_LOCATION"
        else:  # "down"
            move_sw = f"Z_M_P_{ch}_CCW_SW"
            lamp_bit = f"Z{ch}_DOWN_LOCATION"

        if not self.plc:
            self.append_log("PLC", f"[CH{self.ch}] PLC 미연결 상태 → Chuck 제어 불가")
            return False

        timeout_s = float(self.cfg._get("CHAMBER_CHUCK_MOVE_TIMEOUT_S", 60.0))
        power_on_settle_s = float(self.cfg._get("CHAMBER_CHUCK_POWER_ON_SETTLE_S", 0.2))
        poll_interval_s = float(self.cfg._get("CHAMBER_CHUCK_POLL_INTERVAL_S", 0.3))

        async def _read_actual_pos() -> str:
            """3개 lamp 비트를 읽어 실제 위치 반환. 실패 시 'unknown'."""
            try:
                up  = bool(await self.plc.read_bit(f"Z{ch}_UP_LOCATION"))
                mid = bool(await self.plc.read_bit(f"Z{ch}_MID_LOCATION"))
                dn  = bool(await self.plc.read_bit(f"Z{ch}_DOWN_LOCATION"))
                if int(up) + int(mid) + int(dn) == 1:
                    return "up" if up else ("mid" if mid else "down")
            except Exception:
                pass
            return "unknown"

        try:
            # (A) 이미 목표 위치인지 먼저 한 번 확인
            try:
                already = bool(await self.plc.read_bit(lamp_bit))
            except Exception:
                already = False

            if already:
                self.append_log(
                    "PLC",
                    f"[CH{self.ch}] Chuck '{pos}' 이미 목표 위치 ({lamp_bit}=True) → 이동 생략",
                )
                # ✅ 이미 목표 위치 → params 그대로 유지 (pos == 실측값)
                return True

            # (B) POWER ON → MOVE ON
            self.append_log(
                "PLC",
                f"[CH{self.ch}] Chuck '{pos}' 이동 시작: {power_sw} → {move_sw} → {lamp_bit} 폴링",
            )
            # ✅ 상태창 표시
            self._set_state_text(f"Chuck {pos.upper()} 이동 중…")

            await self.plc.write_switch(power_sw, True)
            await asyncio.sleep(power_on_settle_s)
            await self.plc.write_switch(move_sw, True)

            # (C) 램프 폴링 (최대 timeout_s)
            deadline = time.monotonic() + timeout_s
            while time.monotonic() < deadline:
                try:
                    ok = bool(await self.plc.read_bit(lamp_bit))
                except Exception:
                    ok = False

                if ok:
                    # 성공: 스위치 OFF
                    with contextlib.suppress(Exception):
                        await self.plc.write_switch(move_sw, False)
                        await self.plc.write_switch(power_sw, False)
                    self.append_log(
                        "PLC",
                        f"[CH{self.ch}] Chuck '{pos}' 이동 성공 ({lamp_bit}=True)",
                    )
                    # ✅ 성공: lamp 확인으로 이미 pos 검증됨 → params 그대로 유지
                    return True

                await asyncio.sleep(poll_interval_s)

            # (D) 타임아웃: 스위치 OFF 후 실패 반환
            with contextlib.suppress(Exception):
                await self.plc.write_switch(move_sw, False)
                await self.plc.write_switch(power_sw, False)

            actual = await _read_actual_pos()
            self.append_log(
                "PLC",
                f"[CH{self.ch}] Chuck '{pos}' 타임아웃({int(timeout_s)}s) — 실측 위치: {actual}",
            )
            # ✅ 실패: 실측 위치로 덮어씀
            params["chuck_position"] = actual  # type: ignore[index]
            return False

        except Exception as e:
            # (E) 예외 시에도 스위치 OFF 보장
            with contextlib.suppress(Exception):
                try:
                    await self.plc.write_switch(move_sw, False)
                    await self.plc.write_switch(power_sw, False)
                except Exception:
                    pass

            actual = await _read_actual_pos()
            self.append_log(
                "PLC",
                f"[CH{self.ch}] Chuck '{pos}' 이동 중 예외: {e!r} — 실측 위치: {actual}",
            )
            # ✅ 실패: 실측 위치로 덮어씀
            params["chuck_position"] = actual  # type: ignore[index]
            return False

    # ------------------------------------------------------------------
    # Start/Stop (개별 챔버)
    # ------------------------------------------------------------------
    def _handle_start_clicked(self, _checked: bool = False):
        """
        ✅ Runner 기반 시작 처리

        - 기존(레거시): Start 클릭 시 프리플라이트/시작을 detached task로 흩뿌리고,
        process_controller finished 이벤트 펌프에서 cleanup/다음 공정까지 직접 수행
        → 공정 종료 직후 다음 공정 시작 레이스/태스크 누수로 “프리플라이트 멈춤”이 재발 가능

        - 현재: Start 클릭은 Runner 큐에 START/START_QUEUE 명령만 enqueue.
        실제 프리플라이트/공정 시작/종료 정리/다음 공정 진행은 Runner가 순차 처리
        """
        try:
            self._ensure_runner_started()

            remain = runtime_state.remaining_cooldown("chamber", self.ch, cooldown_s=60.0)
            if remain > 0.0:
                secs = int(remain + 0.999)
                self._host_report_start(False, f"cooldown {remain:.0f}s remaining")
                self._post_warning("대기 필요", f"이전 공정 종료 후 1분 대기 필요합니다.\n{secs}초 후에 시작하십시오.")
                return
            
            # ✅ 이전 cleanup이 timeout으로 끝났고 강제 복구도 실패한 상태라면 Start를 막는다.
            if bool(getattr(self, "_pending_device_cleanup", False)):
                self._host_report_start(False, "cleanup pending")
                self._post_warning(
                    "정리 미완료",
                    "이전 공정의 장치 정리가 완전히 끝나지 않았습니다(타임아웃).\n"
                    "STOP을 한 번 더 눌러 강제 복구를 재시도하거나, 필요하면 프로그램을 재시작하세요.\n"
                    "(정리 로그에 남아있는 장치 이름이 원인입니다.)",
                )
                return

            # Runner가 바쁘면 중복 Start 금지
            if getattr(self, "_runner_state", "IDLE") != "IDLE":
                self._host_report_start(False, f"runner busy: {getattr(self,'_runner_state','')}")
                self._post_warning("대기 중", "이전 공정/정리가 아직 진행 중입니다. 잠시 후 다시 시도하세요.")
                return

            if runtime_state.is_running("chamber", self.ch):
                self._host_report_start(False, "this chamber already running")
                self._post_warning("실행 오류", f"CH{self.ch}는 이미 다른 공정이 실행 중입니다.")
                return

            if self.process_controller.is_running:
                self._host_report_start(False, "process controller busy")
                self._post_warning("실행 오류", "다른 공정이 실행 중입니다.")
                return

            self._auto_connect_enabled = True

            # (1) 파일 기반 자동 공정
            if getattr(self, "process_queue", None):
                self.append_log("MAIN", f"[CH{self.ch}] 파일 기반 자동 공정 시작")
                self.current_process_index = -1

                # Runner state는 Runner 내부에서만 변경
                self._runner_put(_RunnerCmd(kind="START_QUEUE"))
                return

            # (2) 단일 공정
            vals = self._validate_single_run_inputs()
            if vals is None:
                self._host_report_start(False, "invalid inputs")
                return

            base_pressure = float(self._get_text("basePressure_edit") or 1e-5)
            working_pressure = float(self._get_text("workingPressure_edit") or 0.0)
            shutter_delay = float(self._get_text("shutterDelay_edit") or 0.0)

            _dep_rate_txt  = (self._get_text("depRate_edit")   or "").strip()
            _thickness_txt = (self._get_text("thickness_edit") or "").strip()
            _dep_rate  = float(_dep_rate_txt)  if _dep_rate_txt  else None
            _thickness = float(_thickness_txt) if _thickness_txt else None

            self._recalc_process_time()   # ← 추가: 위젯 자동계산 강제 반영
            process_time = float(self._get_text("processTime_edit") or 0.0)
            process_name = (self._get_text("integrationTime_edit") or "").strip()
            process_note = process_name if process_name else f"Single CH{self.ch}"

            params: dict[str, Any] = {
                "base_pressure": base_pressure,
                "integration_time": 60,
                "working_pressure": working_pressure,
                "shutter_delay": shutter_delay,
                "process_time": process_time,
                "dep_rate":    _dep_rate,      
                "thickness":   _thickness,      
                "process_note": process_note,
                "Process_name": process_note,  # (다른 코드 참조가 있어 유지)
                **vals,
                "t0_pressed_wall": datetime.now().isoformat(timespec="seconds"),
                "t0_pressed_ns":   time.monotonic_ns(),
            }

            # vals에 "G1_target_name"이 있으므로 "G1 Target" 키도 동기화
            params.setdefault("G1 Target",  params.get("G1_target_name", ""))
            params.setdefault("G2 Target",  params.get("G2_target_name", ""))
            params.setdefault("G3 Target",  params.get("G3_target_name", ""))

            with contextlib.suppress(Exception):
                if not getattr(self, "_log_file_path", None):
                    self._open_run_log(params)

            errs = self._validate_norm_params(cast(NormParams, params))
            if errs:
                self._host_report_start(False, "; ".join(errs))
                self._post_warning("입력값 확인", "\n".join(f"- {e}" for e in errs))
                return

            self.append_log("MAIN", "입력 검증 통과 → Runner START")

            # ✅ Runner만 _runner_state를 소유하게 한다(여기서 선점 금지)
            self._runner_put(_RunnerCmd(kind="START", params=cast(NormParams, params)))

        except Exception as e:
            tb = "".join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
            self.append_log("MAIN", f"_handle_start_clicked 예외:\n{tb}")
            self._host_report_start(False, f"exception: {e!r}")
            with contextlib.suppress(Exception):
                self._post_critical("실행 오류", "공정 시작 준비 중 내부 오류가 발생했습니다.\n로그를 확인하세요.")

    def _handle_stop_clicked(self, _checked: bool = False):
        self.request_stop_all(user_initiated=True)

    def start_presputter_from_ui(self) -> bool:
        """
        Pre-Sputter 자동 실행 진입점.
        'Start' 버튼을 누른 것과 동일한 경로로, 현재 UI 값(기본값/마지막값)으로 1회 실행한다.
        """
        if self.is_running:
            self.append_log("MAIN", f"[CH{self.ch}] PreSputter: 이미 공정 중입니다.")
            return False
        try:
            # 버튼 클릭과 동일 경로(쿨다운·검증·프리플라이트·로깅 모두 재사용)
            self._handle_start_clicked(False)
            self.append_log("MAIN", f"[CH{self.ch}] PreSputter 자동 시작 (UI 현재값)")
            return True
        except Exception as e:
            self.append_log("MAIN", f"[CH{self.ch}] PreSputter 시작 실패: {e!r}")
            return False

    def request_stop_all(self, user_initiated: bool):
        """
        ✅ 구조 변경:
        - STOP은 어떤 상태든 Runner로 일원화
        - '공정 중이지 않다'로 STOP이 막히던 케이스 제거
        """
        self._runner_put(_RunnerCmd(kind="STOP", user_initiated=bool(user_initiated)))


    # ======================= runner 메서드 =======================
    def _ensure_runner_started(self) -> None:
        """
        챔버당 1개의 Runner만 존재하도록 보장한다.

        핵심:
        - '검사/생성'을 반드시 이벤트루프 스레드에서 수행해야 레이스로 2개가 뜨지 않는다.
        - 이미 떠버린 중복 Runner가 있으면(all_tasks에서 name 기준) 자동 cancel해서 복구한다.
        """
        loop = self._loop
        runner_name = f"Runner.CH{self.ch}"

        def _start_in_loop() -> None:
            # 1) 이미 떠있는 중복 Runner 정리(자가 복구)
            try:
                me = getattr(self, "_runner_task", None)
                for t in asyncio.all_tasks():
                    if t is me:
                        continue
                    if not isinstance(t, asyncio.Task):
                        continue
                    if t.done():
                        continue
                    # 이름이 동일한 Runner가 2개 이상이면 중복
                    if getattr(t, "get_name", None) and t.get_name() == runner_name:
                        t.cancel()
                        with contextlib.suppress(Exception):
                            self.append_log("MAIN", f"[Runner] duplicate runner cancelled: {t!r}")
            except Exception:
                pass

            # 2) Runner가 없으면 1개만 생성
            t = getattr(self, "_runner_task", None)
            if isinstance(t, asyncio.Task) and (not t.done()):
                return

            try:
                # _ensure_runner_started() 안에서 runner 생성하는 부분만 교체
                token = uuid.uuid4().hex
                self._runner_token = token
                self._runner_task = loop.create_task(self._runner_main(token), name=runner_name)
            except Exception as e:
                tb = "".join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
                self.append_log("Task", f"[{runner_name}] create_task failed:\n{tb}")

        try:
            running = asyncio.get_running_loop()
        except RuntimeError:
            running = None

        if running is loop:
            _start_in_loop()
        else:
            with contextlib.suppress(Exception):
                loop.call_soon_threadsafe(_start_in_loop)


    def _runner_put(self, cmd: _RunnerCmd) -> None:
        loop = self._loop

        def _do_put() -> None:
            # ✅ put 직전에 루프 스레드에서 Runner 1개 보장(중복 정리 포함)
            self._ensure_runner_started()

            # ============================================================
            # ✅ 핵심: START/STOP 중복 enqueue 차단 (레이스/잔여 STOP 방지)
            # ============================================================
            if cmd.kind in ("START", "START_QUEUE"):
                if getattr(self, "_runner_cmd_start_enqueued", False):
                    self.append_log("MAIN", f"[Runner] drop duplicate {cmd.kind} (already enqueued)")
                    return
                self._runner_cmd_start_enqueued = True

            elif cmd.kind == "STOP":
                # STOP은 연타/중복 시그널이 들어오면 다음 START를 끊을 수 있으므로 1개만 유지
                # 단, _pending_device_cleanup(취소 타임아웃) 상태에서는 STOP 재시도가 필요하니 허용
                if getattr(self, "_runner_cmd_stop_enqueued", False) and (not getattr(self, "_pending_device_cleanup", False)):
                    self.append_log("MAIN", "[Runner] drop duplicate STOP (already enqueued)")
                    return
                self._runner_cmd_stop_enqueued = True

            try:
                self._cmd_q.put_nowait(cmd)
            except asyncio.QueueFull:
                # rollback
                if cmd.kind in ("START", "START_QUEUE"):
                    self._runner_cmd_start_enqueued = False
                elif cmd.kind == "STOP":
                    self._runner_cmd_stop_enqueued = False
                self.append_log("MAIN", f"[Runner] cmd queue full → drop: {cmd.kind}")

        try:
            running = asyncio.get_running_loop()
        except RuntimeError:
            running = None

        if running is loop:
            _do_put()
        else:
            with contextlib.suppress(Exception):
                loop.call_soon_threadsafe(_do_put)


    async def _runner_start_stage(
        self,
        kind: str,
        coro: Coroutine[Any, Any, Any],
        *,
        cancel_timeout: float = 10.0,
    ) -> None:
        """
        Runner 내부에서만 쓰는 '단일 stage task' 실행기.
        - 기존 stage는 cancel 후 "완전히 끝날 때까지" 기다려 레이스를 제거한다.

        ✅ 중요: stage cancel이 timeout이면(=아직 살아있을 수 있음) 새 stage를 시작하면 안 된다.
        (유령 stage + 동시 실행 → ADVANCE_QUEUE cancelled / 대기중 멈춤 유발)
        """
        await self._runner_cancel_stage(timeout=cancel_timeout, reason=f"start:{kind}")

        # cancel이 timeout이면 _runner_stage_task가 그대로 남아있다 → 새 stage 시작 금지
        t_prev = getattr(self, "_runner_stage_task", None)
        if isinstance(t_prev, asyncio.Task) and (not t_prev.done()):
            self.append_log("MAIN", f"[Runner] start_stage blocked: previous stage still alive kind={getattr(self, '_runner_stage_kind', None)}")
            # 시작은 막되, UI는 멈춰있는 것처럼 보이지 않게 IDLE로 표시(단, pending cleanup으로 Start는 차단됨)
            self._runner_state = "IDLE"
            return

        self._runner_stage_kind = kind
        t_stage = self._spawn_detached(
            coro,
            store=False,
            name=f"RunnerStage.{kind}.CH{self.ch}",
        )
        self._runner_stage_task = t_stage

        if isinstance(t_stage, asyncio.Task):
            def _clear_stage_done(task: asyncio.Task) -> None:
                with contextlib.suppress(Exception):
                    if getattr(self, "_runner_stage_task", None) is task:
                        self._runner_stage_task = None
                        self._runner_stage_kind = None

            t_stage.add_done_callback(_clear_stage_done)


    async def _runner_cancel_stage(
        self,
        *,
        timeout: float = 10.0,
        reason: str = "stage-switch",
    ) -> None:
        """
        stage task가 있으면 취소하고 종료까지(최대 timeout) 기다린다.

        ✅ 핵심:
        - timeout 내에 stage가 끝나지 않으면: 참조를 지우지 않는다(유령 stage 방지)
        - 대신 _pending_device_cleanup=True로 승격해서 다음 Start를 차단한다.
        """
        t = getattr(self, "_runner_stage_task", None)

        # stage가 없거나 이미 끝났으면 정리
        if not isinstance(t, asyncio.Task) or t.done():
            self._runner_stage_task = None
            self._runner_stage_kind = None
            return

        cur = asyncio.current_task()
        if t is cur:
            return

        kind = getattr(self, "_runner_stage_kind", None)
        state = getattr(self, "_runner_state", None)

        # Runner가 의도적으로 cancel하는 경우 표시(취소 핸들러에서 “예상된 취소”로 처리)
        self._expected_stage_cancel_task = t
        self._expected_stage_cancel_kind = kind

        self.append_log(
            "MAIN",
            f"[Runner] cancel stage begin kind={kind} state={state} reason={reason} timeout={timeout:.1f}s"
        )

        try:
            t.cancel()
            try:
                await asyncio.wait_for(t, timeout=timeout)
            except asyncio.TimeoutError:
                # ❗ stage가 아직 살아있을 수 있음 → 참조 유지 + 시작 차단
                self.append_log(
                    "MAIN",
                    f"[Runner] stage cancel TIMEOUT ({timeout:.1f}s) kind={kind} state={state} → pending cleanup"
                )
                self._pending_device_cleanup = True
                self._cleanup_timed_out = True
                with contextlib.suppress(Exception):
                    self._set_state_text("정리 지연(취소 타임아웃). STOP을 한 번 더 눌러 복구를 시도하세요.")
                return
            except asyncio.CancelledError:
                pass
            except Exception as e:
                self.append_log("MAIN", f"[Runner] stage cancel wait exception kind={kind}: {e!r} (ignored)")
        finally:
            if getattr(self, "_expected_stage_cancel_task", None) is t:
                self._expected_stage_cancel_task = None
                self._expected_stage_cancel_kind = None

        # 여기까지 왔으면 done이거나 거의 확실히 종료
        if t.done():
            self._runner_stage_task = None
            self._runner_stage_kind = None


    def _cancel_delay_task(self) -> None:
        """
        ✅ 레거시 호환: 'delay/countdown task'만 정리하는 안전 버전

        Runner 구조에서는 START_SINGLE/ADVANCE_QUEUE 같은 stage는 _runner_stage_task로 운용된다.
        그런데 _cancel_delay_task()가 stage를 cancel하면,
        '이전 공정의 늦은 aborted/reset/cleanup'이 '다음 공정 stage'까지 끊어버리는 레이스가 생긴다.

        따라서 이 함수는:
        - 과거 레거시 delay/countdown task가 남아있을 때만 cancel
        - ❌ _runner_stage_task(현재 stage)는 절대 cancel하지 않는다
        (stage 취소는 STOP/_runner_cancel_stage 또는 shutdown_fast에서 명시적으로 수행)
        """
        legacy_attrs = (
            "_delay_main_task",
            "_delay_countdown_task",
            "_delay_task",
            "_countdown_task",
        )

        cancelled_any = False
        for attr in legacy_attrs:
            t = getattr(self, attr, None)
            if isinstance(t, asyncio.Task) and (not t.done()):
                cancelled_any = True
                with contextlib.suppress(Exception):
                    t.cancel()
            # 참조 제거(다음 런에 영향 방지)
            with contextlib.suppress(Exception):
                if getattr(self, attr, None) is t:
                    setattr(self, attr, None)

        # (선택) 디버그 로그: 레거시가 실제로 있었을 때만
        if cancelled_any:
            self.append_log("MAIN", "[Legacy] cancel delay/countdown tasks")

    
    async def _runner_handle_stop(self, user_initiated: bool) -> None:
        """
        Runner STOP 처리(단일 진입점).

        원칙
        1) 현재 stage(preflight/쿨다운/딜레이/advance 등)가 있으면 먼저 취소한다.
        2) ProcessController가 RUNNING이면 장치 cleanup은 하지 않고 request_stop()만 보낸다.
           - finished 이벤트가 들어오면 AFTER_FINISH stage가 cleanup을 수행한다.
        3) RUNNING이 아니면 지금 즉시 heavy cleanup을 수행하고 UI를 Idle로 복구한다.

        ✅ 공정 로직/장비 파라미터는 바꾸지 않고,
           "정지 처리의 주체"만 Runner로 모으는 구조 변경이다.
        """
        self.append_log("MAIN", f"[Runner] STOP 요청(user={user_initiated}) state={getattr(self, '_runner_state', '')}")

        # 리스트 자동 진행은 여기서 끊는다(완료/실패/STOP 후 다음 공정으로 넘어가지 않게)
        self._runner_queue_mode = False

        # 정지 중 자동 재연결/백그라운드 재기동 방지
        self._auto_connect_enabled = False

        # stage 취소 (프리플라이트/딜레이/큐 advance 등)
        await self._runner_cancel_stage(timeout=10.0, reason="stop")

        # cancel timeout이면 stage가 아직 살아있을 수 있다 → 강제 복구 모드로 승격
        t_stage = getattr(self, "_runner_stage_task", None)
        if isinstance(t_stage, asyncio.Task) and (not t_stage.done()):
            self.append_log("MAIN", f"[Runner] STOP: stage cancel incomplete(kind={getattr(self,'_runner_stage_kind',None)}) → pending cleanup")
            self._pending_device_cleanup = True
            self._runner_state = "IDLE"
            with contextlib.suppress(Exception):
                self._set_state_text("정리 지연(취소 미완료). STOP을 한 번 더 눌러 복구를 시도하세요.")
            return

        # 공정이 실행 중이면: PC에 stop 요청만 보낸다(장치 정리는 finished 이후)
        if bool(getattr(self.process_controller, "is_running", False)):
            self.append_log("MAIN", "[Runner] STOP → process_controller.request_stop()")
            with contextlib.suppress(Exception):
                self._set_state_text("STOP 요청 중... (공정 종료 대기)")
            with contextlib.suppress(Exception):
                self.process_controller.request_stop()
            # 여기서는 상태를 IDLE로 만들지 않는다. finished가 오면 AFTER_FINISH가 정리한다.
            return

        # 공정 시작 전 상태(preflight/idle/delay)라면 지금 바로 정리한다.
        self.append_log("MAIN", "[Runner] STOP → 공정 시작 전 상태, 즉시 정리(cleanup)")
        with contextlib.suppress(Exception):
            self._apply_polling_targets({"mfc": False, "dc_pulse": False, "rf_pulse": False, "dc": False, "rf": False})

        # START stage에서 mark_started를 찍었을 수 있으므로, 여기서 running 상태가 남지 않게 마무리
        with contextlib.suppress(Exception):
            runtime_state.mark_finished("chamber", self.ch)

        with contextlib.suppress(Exception):
            await self._stop_device_watchdogs(light=False)

        with contextlib.suppress(Exception):
            self._clear_queue_and_reset_ui()

        self._runner_state = "IDLE"


    async def _runner_main(self, token: str) -> None:
        """
        Runner 메인 루프:
        - UI/Host/기타 경로에서 들어오는 START/START_QUEUE/STOP/PC_FINISHED를 '단일 지점'에서 순차 처리.
        - stage 실행은 _runner_start_stage()로 통일한다.
        """
        while True:
            cmd = await self._cmd_q.get()

            # ✅ 디바운스 플래그 해제: 이제부터는 Runner state/stage로 중복 방지
            if cmd.kind in ("START", "START_QUEUE"):
                self._runner_cmd_start_enqueued = False
            elif cmd.kind == "STOP":
                self._runner_cmd_stop_enqueued = False

            # ✅ get 직후에도 재확인: stale runner가 명령을 “먹어버리는” 걸 방지
            if getattr(self, "_runner_token", None) != token:
                with contextlib.suppress(Exception):
                    self._cmd_q.put_nowait(cmd)  # put back
                return

            try:
                if cmd.kind == "START":
                    # ✅ 중복 START 방지: stage가 이미 실행 중이면 무시
                    t_stage = getattr(self, "_runner_stage_task", None)
                    if isinstance(t_stage, asyncio.Task) and (not t_stage.done()):
                        self.append_log(
                            "MAIN",
                            f"[Runner] START ignored (stage running kind={getattr(self,'_runner_stage_kind',None)})"
                        )
                        continue

                    if not isinstance(cmd.params, dict):
                        self.append_log("MAIN", "[Runner] START: params missing → ignore")
                        continue

                    self._runner_queue_mode = False
                    self._runner_state = "PREFLIGHT"
                    await self._runner_start_stage("START_SINGLE", self._runner_stage_start_single(cmd.params))

                elif cmd.kind == "START_QUEUE":
                    # ✅ 중복 START_QUEUE 방지: stage가 이미 실행 중이면 무시
                    t_stage = getattr(self, "_runner_stage_task", None)
                    if isinstance(t_stage, asyncio.Task) and (not t_stage.done()):
                        self.append_log(
                            "MAIN",
                            f"[Runner] START_QUEUE ignored (stage running kind={getattr(self,'_runner_stage_kind',None)})"
                        )
                        continue

                    self._runner_queue_mode = True
                    self._runner_state = "COOLDOWN"
                    await self._runner_start_stage("ADVANCE_QUEUE", self._runner_stage_advance_queue(was_successful=True))

                elif cmd.kind == "PC_FINISHED":
                    prev_state = getattr(self, "_runner_state", "IDLE")

                    # ✅ RUNNING/STOPPING 흐름이 아닌데 finished가 오면 “stale 이벤트”로 보고 무시
                    if prev_state not in ("RUNNING", "STOPPING", "CLEANUP"):
                        self.append_log("MAIN", f"[Runner] PC_FINISHED ignored (state={prev_state})")
                        continue

                    ok = bool(cmd.ok)
                    detail = dict(cmd.detail or {})
                    self._runner_state = "CLEANUP"
                    await self._runner_start_stage("AFTER_FINISH", self._runner_stage_after_finish(ok, detail))

                elif cmd.kind == "STOP":
                    # STOP은 즉시 처리(현재 stage 취소/정리)
                    self._runner_state = "STOPPING"
                    await self._runner_handle_stop(cmd.user_initiated)

                else:
                    self.append_log("MAIN", f"[Runner] unknown cmd: {cmd.kind}")

            except Exception as e:
                tb = "".join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
                self.append_log("MAIN", f"[Runner] loop exception:\n{tb}")
                with contextlib.suppress(Exception):
                    self._clear_queue_and_reset_ui()
                self._runner_state = "IDLE"

    
    async def _runner_stage_start_single(self, params: NormParams) -> None:
        """
        단일 공정 실행(Runner Stage):
        - preflight → process_controller.start_process()까지 진행
        - 실패하면 여기서 cleanup 후 IDLE로 복귀
        """
        try:
            # 세대 번호 갱신 (기존 로직 유지)
            self._run_gen = int(getattr(self, "_run_gen", 0)) + 1
            gen = self._run_gen
            self._active_run_gen = gen

            # ✅ 기존 _safe_start_process와 동일하게 "시작" 상태를 먼저 찍어 둔다.
            # - preflight 중에도 다른 Start를 막고
            # - Host/외부에서 is_running 판정을 일관되게 하기 위함
            with contextlib.suppress(Exception):
                runtime_state.mark_started("chamber", self.ch)

            # ✅ 프리플라이트/시작은 "await"로 직접 실행 (detached로 던지지 않음)
            await self._start_after_preflight(params, gen)

            # 여기까지 왔으면 PC가 start_process를 호출한 상태(= RUNNING 진입)
            self._runner_state = "RUNNING"

        except asyncio.CancelledError:
            self.append_log("MAIN", "[Runner] START_SINGLE cancelled")

            # ✅ Runner가 stage 교체/STOP 처리 중 의도적으로 cancel한 경우:
            #    - 여기서 queue/UI/state를 건드리면 다음 stage(또는 STOP cleanup)가 망가질 수 있음
            if getattr(self, "_expected_stage_cancel_task", None) is asyncio.current_task():
                raise

            # ✅ 의도치 않은 cancel(외부 cancel/레이스)일 때만 최소 복구
            with contextlib.suppress(Exception):
                runtime_state.mark_finished("chamber", self.ch)
            self._runner_queue_mode = False
            self._runner_state = "IDLE"
            raise

        except Exception as e:
            self.append_log("MAIN", f"[Runner] START_SINGLE failed: {e!r}")
            # 실패 시 정리 + UI 리셋
            with contextlib.suppress(Exception):
                self._auto_connect_enabled = False
                await self._stop_device_watchdogs(light=False)
            with contextlib.suppress(Exception):
                self._clear_queue_and_reset_ui()
            self._runner_state = "IDLE"

    
    async def _runner_stage_after_finish(self, ok: bool, detail: dict[str, Any]) -> None:
        """
        process_controller finished 이후:
        - 장치 정리
        - (큐 모드 + 성공이면) 다음 공정으로 advance
        """
        try:
            stopped = bool(detail.get("stopped", False))
            is_test = bool(detail.get("test_mode", False))
            is_test_cancel = bool(is_test and stopped)

            # 1) 자동연결 차단 후 정리
            self._auto_connect_enabled = False

            if not is_test_cancel:
                self.append_log("MAIN", "[Runner] finished → device cleanup")
                await self._stop_device_watchdogs(light=False)
            else:
                self.append_log("MAIN", "[Runner] [TEST] STOP finished → cleanup skip")

            # 2) 공정 종료 후 공통 리셋
            self._last_polling_targets = None

            # 3) 큐 진행 여부 결정
            #    - STOP/stopped면 다음 공정으로 넘어가지 않음
            if self._runner_queue_mode and ok and (not stopped):
                self.append_log("MAIN", "[Runner] queue advance (ok)")
                self._runner_state = "COOLDOWN"

                # 중요:
                # AFTER_FINISH stage 내부에서 다음 stage를 직접 시작하면
                # 현재 _runner_stage_task가 아직 자기 자신(AFTER_FINISH)이라
                # "start_stage blocked: previous stage still alive kind=AFTER_FINISH"
                # 가 발생할 수 있다.
                #
                # 따라서 현재 stage가 완전히 끝난 뒤,
                # runner main 루프가 다음 START_QUEUE 명령을 처리하도록
                # 다음 틱에 enqueue만 하고 여기서는 바로 반환한다.
                self._soon(self._runner_put, _RunnerCmd(kind="START_QUEUE"))
                return

            # 큐 종료(실패/stop/마지막) → UI 정리
            self.append_log("MAIN", f"[Runner] queue end (ok={ok}, stopped={stopped}) → reset")
            self._runner_queue_mode = False
            with contextlib.suppress(Exception):
                self._clear_queue_and_reset_ui()
            self._runner_state = "IDLE"

        except asyncio.CancelledError:
            raise
        except Exception as e:
            self.append_log("MAIN", f"[Runner] AFTER_FINISH exception: {e!r}")
            with contextlib.suppress(Exception):
                self._clear_queue_and_reset_ui()
            self._runner_state = "IDLE"


    async def _runner_stage_advance_queue(self, was_successful: bool) -> None:
        """
        큐 모드(파일 기반 자동 공정) Runner Stage:
        - 다음 params 선택
        - delay step이면 대기 후 다음으로 계속
        - normal step이면 preflight → start_process 진입
        """
        try:
            if not was_successful:
                self.append_log("MAIN", "[Runner] prev failed → stop queue")
                self._runner_queue_mode = False
                with contextlib.suppress(Exception):
                    self._clear_queue_and_reset_ui()
                self._runner_state = "IDLE"
                return

            # ✅ delay step 재귀를 없애기 위해: 여기부터 루프로 돌린다.
            while True:
                # 다음 index
                self.current_process_index = int(getattr(self, "current_process_index", -1)) + 1
                q = list(getattr(self, "process_queue", []) or [])

                if self.current_process_index >= len(q):
                    self.append_log("MAIN", "[Runner] queue finished → reset")
                    self._runner_queue_mode = False
                    with contextlib.suppress(Exception):
                        self._clear_queue_and_reset_ui()
                    self._runner_state = "IDLE"
                    return

                # 다음 step params
                params = q[self.current_process_index]
                self._update_ui_from_params(params)

                # ✅ CSV에 G1/G2/G3 Target이 비어있으면 DB에서 로딩 완료 대기
                _needs_db = not any([
                    str(params.get("G1 Target", "") or params.get("G1_target_name", "")).strip(),
                    str(params.get("G2 Target", "") or params.get("G2_target_name", "")).strip(),
                    str(params.get("G3 Target", "") or params.get("G3_target_name", "")).strip(),
                ])
                if _needs_db:
                    try:
                        await asyncio.wait_for(self._load_gun_targets(), timeout=3.0)
                    except asyncio.TimeoutError:
                        self.append_log(
                            "TARGET",
                            f"[CH{self.ch}] Gun Target DB 3초 초과 — 빈 값으로 진행 (공정 진행)"
                        )
                    except Exception as e:
                        self.append_log(
                            "TARGET",
                            f"[CH{self.ch}] Gun Target DB 실패: {e!r} — 빈 값으로 진행"
                        )

                # ------------------------------
                # (A) delay step 처리
                #  ① 전용 'delay' 컬럼 우선: 10 → 10분, 90s / 10m / 1h30m 형식 지원
                #  ② (하위호환) Process_name의 'delay 10m' 패턴
                # ------------------------------
                name = str(params.get("Process_name") or params.get("process_note", "")).strip()
                duration_s = 0.0
                label = ""

                _dcol = str(params.get("delay", "") or "").strip()
                if _dcol:
                    try:
                        duration_s = float(_dcol) * 60.0          # 숫자만 → 분
                        label = f"{_dcol}분"
                    except Exception:
                        duration_s = float(self._parse_duration_seconds(_dcol.lower()))
                        label = _dcol
                else:
                    m = re.match(r"^\s*delay\s*(\d+)\s*([smhd]?)\s*$", name, re.IGNORECASE) if name else None
                    if m:
                        amount = int(m.group(1))
                        unit = (m.group(2) or "m").lower()
                        factor = {"s": 1.0, "m": 60.0, "h": 3600.0, "d": 86400.0}[unit]
                        duration_s = float(amount) * factor
                        unit_txt = {"s": "초", "m": "분", "h": "시간", "d": "일"}[unit]
                        label = f"{amount}{unit_txt}"

                if duration_s > 0:
                    disp = name if name else f"delay {label}"

                    # ✅ 외부 START_SPUTTER는 "큐 공정이 정상 수락되어 delay 단계에 진입"한 시점도 성공으로 본다.
                    #    delay-first recipe에서는 실제 preflight가 나중에 실행되므로,
                    #    여기서 host 응답 future를 먼저 success로 완료해 timeout fail을 막는다.
                    self._host_report_start(True, f"delay accepted: {label}")

                    self.append_log("Process", f"[Runner] '{disp}' 단계: {label} 대기 시작")
                    # 상태 표시 + 카운트다운(취소 가능: STOP이 오면 stage task cancel됨)
                    remain = int(duration_s)
                    while remain > 0:
                        self._set_state_text(f"지연 대기 중: {label} · 남은 시간 {self._fmt_hms(remain)}")
                        await asyncio.sleep(1)
                        remain -= 1

                    self.append_log("Process", f"[Runner] '{disp}' 지연 완료 → 다음 스텝")
                    await asyncio.sleep(0)
                    continue

                # ------------------------------
                # (B) TEST MODE marker 처리 (기존과 동일)
                # ------------------------------
                marker = str(params.get("#") or "").strip().lower()
                if marker == "test":
                    time_str = str(params.get("time") or "").strip()
                    test_duration_sec = self._parse_duration_seconds(time_str.lower())
                    params["test_mode"] = True
                    params["test_duration_sec"] = test_duration_sec
                    if test_duration_sec > 0:
                        params.setdefault("process_time", round(test_duration_sec / 60.0, 3))
                    params.setdefault("process_note", params.get("Process_name") or "TEST")

                norm = self._normalize_params_for_process(params)

                # 검증(기존 동일)
                errs = self._validate_norm_params(norm)
                if errs:
                    with contextlib.suppress(Exception):
                        if not getattr(self, "_log_file_path", None):
                            self._open_run_log(norm)
                    self.append_log("Validate", "CSV 공정 파라미터 오류:\n - " + "\n - ".join(errs))
                    self._runner_queue_mode = False
                    with contextlib.suppress(Exception):
                        self._clear_queue_and_reset_ui()
                    self._runner_state = "IDLE"
                    return

                # 다음 공정은 새 로그 파일로
                self._log_file_path = None
                with contextlib.suppress(Exception):
                    self._open_run_log(norm)

                # ------------------------------
                # (C) 쿨다운(기존 동일)
                # ------------------------------
                try:
                    remain = float(runtime_state.remaining_cooldown("chamber", self.ch, 60.0))
                except Exception:
                    remain = 0.0

                first_step = (self.current_process_index == 0)
                delay_s = max(0.0, float(remain))

                if delay_s > 0.0 and not first_step:
                    reason = "쿨다운 대기"
                    self.append_log("MAIN", f"[Runner] cooldown wait: {delay_s:.1f}s ({reason})")
                    r = int(delay_s)
                    while r > 0:
                        self._set_state_text(f"다음 공정 대기중 ({reason}) · 남은 시간 {self._fmt_hms(r)}")
                        await asyncio.sleep(1)
                        r -= 1

                # ------------------------------
                # (D) preflight → start (await로 직접)
                # ------------------------------
                self._runner_state = "PREFLIGHT"
                self._run_gen = int(getattr(self, "_run_gen", 0)) + 1
                gen = self._run_gen
                self._active_run_gen = gen

                with contextlib.suppress(Exception):
                    runtime_state.mark_started("chamber", self.ch)

                await self._start_after_preflight(norm, gen)
                self._runner_state = "RUNNING"
                return

        except asyncio.CancelledError:
            self.append_log("MAIN", "[Runner] ADVANCE_QUEUE cancelled")

            if getattr(self, "_expected_stage_cancel_task", None) is asyncio.current_task():
                raise

            with contextlib.suppress(Exception):
                runtime_state.mark_finished("chamber", self.ch)

            self._runner_queue_mode = False
            self._runner_state = "IDLE"
            raise

        except Exception as e:
            self.append_log("MAIN", f"[Runner] ADVANCE_QUEUE failed: {e!r}")
            self._runner_queue_mode = False
            with contextlib.suppress(Exception):
                self._auto_connect_enabled = False
                await self._stop_device_watchdogs(light=False)
            with contextlib.suppress(Exception):
                self._clear_queue_and_reset_ui()
            self._runner_state = "IDLE"
    # ======================= runner 메서드 =======================

    
    def _skip_mfc_finalize_due_to_pc(self) -> bool:
        """
        CH1 chamber 종료 시, PC가 mfc1(gas)을 공유 사용 중이면
        mfc 폴링/상태 리셋/cleanup을 모두 생략한다.
        (PC.mfc_gas는 main.py에서 항상 self.mfc1 = self.ch1.mfc 로 주입됨)
        CH2 chamber의 self.mfc는 mfc2이므로 공유 이슈 없음 → 항상 False.
        """
        try:
            if int(self.ch) != 1:
                return False
            # PC가 어느 챔버를 선택했든 mfc_gas는 mfc1을 쓰므로 둘 다 체크
            return bool(runtime_state.is_running("pc", 1)) or \
                   bool(runtime_state.is_running("pc", 2))
        except Exception:
            return False


    async def _stop_device_watchdogs(self, *, light: bool = False) -> None:
        if light:
            with contextlib.suppress(Exception):
                if not self._skip_mfc_finalize_due_to_pc():
                    self.mfc.set_process_status(False)
            if self.dc_pulse:
                with contextlib.suppress(Exception): self.dc_pulse.set_process_status(False)
            if self.rf_pulse:
                with contextlib.suppress(Exception): self.rf_pulse.set_process_status(False)
            if self.dc_power and hasattr(self.dc_power, "set_process_status"):
                with contextlib.suppress(Exception): self.dc_power.set_process_status(False)
            if self.dc_power2 and hasattr(self.dc_power2, "set_process_status"):
                with contextlib.suppress(Exception): self.dc_power2.set_process_status(False)
            if self.rf_power and hasattr(self.rf_power, "set_process_status"):
                with contextlib.suppress(Exception): self.rf_power.set_process_status(False)
            return

        # ✅ 이번 cleanup이 “완전히 끝났는지” 표시
        #    - False면 정상 정리 완료
        #    - True면 정상 정리 실패(→ 아래에서 Force Recovery 승격)
        self._cleanup_timed_out = False

        pending_cleanup_names: list[str] = []

        # ✅ heavy 시작 직후도 한 번 더 OFF
        with contextlib.suppress(Exception):
            if self._skip_mfc_finalize_due_to_pc():
                self.append_log("MFC", "PC 실행 중 → mfc 폴링/상태 리셋 생략(공유 자원 보호)")
            elif self.mfc and hasattr(self.mfc, "on_process_finished"):
                self.mfc.on_process_finished(False)
            elif self.mfc and hasattr(self.mfc, "set_process_status"):
                self.mfc.set_process_status(False)

        if self.dc_pulse:
            with contextlib.suppress(Exception): self.dc_pulse.set_process_status(False)
        if self.rf_pulse:
            with contextlib.suppress(Exception): self.rf_pulse.set_process_status(False)
        if self.dc_power and hasattr(self.dc_power, "set_process_status"):
            with contextlib.suppress(Exception): self.dc_power.set_process_status(False)
        if self.dc_power2 and hasattr(self.dc_power2, "set_process_status"):
            with contextlib.suppress(Exception): self.dc_power2.set_process_status(False)
        if self.rf_power and hasattr(self.rf_power, "set_process_status"):
            with contextlib.suppress(Exception): self.rf_power.set_process_status(False)

        loop = self._loop_from_anywhere()

        # 0) bg task cancel (timeout)
        try:
            current = asyncio.current_task()
            live = [t for t in getattr(self, "_bg_tasks", []) if t and not t.done() and t is not current]
            for t in live:
                with contextlib.suppress(Exception):
                    t.cancel()

            if live:
                try:
                    await asyncio.wait_for(asyncio.gather(*live, return_exceptions=True), timeout=5.0)
                except asyncio.TimeoutError:
                    self._cleanup_timed_out = True
                    with contextlib.suppress(Exception):
                        names = [getattr(t, "get_name", lambda: repr(t))() for t in live]
                    self.append_log("MAIN", f"⚠ bg task cancel timeout: {names!r}")
        finally:
            self._bg_tasks = []

        # 1) IG cancel (timeout 유지)
        try:
            if self.ig and hasattr(self.ig, "cancel_wait"):
                with contextlib.suppress(asyncio.TimeoutError):
                    await asyncio.wait_for(self.ig.cancel_wait(), timeout=2.0)
        except Exception:
            pass

        # 2) device cleanup (timeout)
        # 2-A) ✅ OES는 stop 요청만 던지고 기다리지 않는다.
        #      - one-shot 모드: 다음 공정에서 init을 다시 타도록 캐시 리셋
        #      - daemon 모드: OES daemon을 계속 유지하므로 캐시 유지
        if self.oes:
            with contextlib.suppress(Exception):
                if hasattr(self.oes, "stop_measurement"):
                    self._spawn_detached(
                        self.oes.stop_measurement(wait=False),
                        store=False,
                        name=f"Cleanup.OESAsync.CH{self.ch}.DETACHED",
                    )
                else:
                    self._spawn_detached(
                        self.oes.cleanup(),
                        store=False,
                        name=f"Cleanup.OESAsync.CH{self.ch}.DETACHED",
                    )

            # ✅ daemon 모드에서는 init 캐시를 깨지 않는다(상주 유지)
            if not bool(getattr(self.oes, "_daemon_enabled", False)):
                self._oes_initialized = False

        # 2-B) device cleanup (timeout) - ✅ OES는 제외하고 나머지만 기다림
        # ✅ PC가 mfc1(gas)을 공유 사용 중이면 self.mfc.cleanup()을 생략 (공유 자원 보호)
        _skip_mfc = self._skip_mfc_finalize_due_to_pc()
        if _skip_mfc:
            self.append_log("MFC", "PC 실행 중 → mfc cleanup 생략 (공유 자원 보호)")
        cleanup_tasks: list[asyncio.Task] = []
        for dev in (self.ig, self.mfc, self.dc_pulse, self.rf_pulse, self.dc_power, self.dc_power2, self.rf_power, self.rga):
            if dev is self.mfc and _skip_mfc:
                continue
            if dev and hasattr(dev, "cleanup"):
                try:
                    coro = dev.cleanup()
                except Exception:
                    continue
                try:
                    nm = getattr(dev, "NAME", None) or dev.__class__.__name__
                    cleanup_tasks.append(loop.create_task(coro, name=f"Cleanup.{nm}.CH{self.ch}"))
                except Exception:
                    cleanup_tasks.append(loop.create_task(coro))

        if cleanup_tasks:
            cleanup_timeout_s = 15.0  # ✅ OES 때문에 75초로 늘리지 않음
            done, pending = await asyncio.wait(cleanup_tasks, timeout=cleanup_timeout_s)
            if pending:
                self._cleanup_timed_out = True

                pn: list[str] = []
                with contextlib.suppress(Exception):
                    pn = [t.get_name() for t in pending]
                if not pn:
                    pn = [repr(t) for t in pending]
                pending_cleanup_names = list(pn)

                self.append_log("MAIN", f"⚠ device cleanup timeout: {pending_cleanup_names!r}")

                for t in pending:
                    with contextlib.suppress(Exception):
                        t.cancel()

                # ✅ cancel 했는데도 안 죽는 cleanup이 있으면 여기서 무한 대기 가능 → 2초로 끊음
                try:
                    await asyncio.wait_for(
                        asyncio.gather(*pending, return_exceptions=True),
                        timeout=2.0
                    )
                except asyncio.TimeoutError:
                    self._cleanup_timed_out = True
                    with contextlib.suppress(Exception):
                        pn2 = [getattr(t, "get_name", lambda: repr(t))() for t in pending]
                    self.append_log("MAIN", f"⚠ cleanup cancel timeout(2s): {pn2!r} (detached/leaked)")

        # 3) footer 먼저
        with contextlib.suppress(Exception):
            self._close_run_log()

        # 4) writer 완전 종료 (timeout)
        #    내부 NAS flush(6.0s) + close(2.0s+2.0s) 가 직렬 실행되므로,
        #    호출측은 그보다 큰 10.0s 로 잡아 정상 케이스에서 오발동 방지.
        try:
            t = loop.create_task(self._shutdown_log_writer(), name=f"ShutdownLogWriter.CH{self.ch}")
            try:
                await asyncio.wait_for(t, timeout=10.0)
            except asyncio.TimeoutError:
                self._cleanup_timed_out = True
                self.append_log("MAIN", "⚠ log writer shutdown timeout")
                with contextlib.suppress(Exception):
                    t.cancel()
                    try:
                        await asyncio.wait_for(asyncio.gather(t, return_exceptions=True), timeout=2.0)
                    except asyncio.TimeoutError:
                        self._cleanup_timed_out = True
                        self.append_log("MAIN", "⚠ log writer cancel timeout(2s): leaked")
        except Exception:
            self._cleanup_timed_out = True
            with contextlib.suppress(Exception):
                self.append_log("MAIN", "⚠ log writer shutdown exception")

        # 5) 파일 경로/버퍼 초기화
        self._log_file_path = None
        with contextlib.suppress(Exception):
            self._prestart_buf.clear()

        self._bg_started = False
        self._devices_started = False
        self._run_select = None

        # ------------------------------------------------------------------
        # ✅ 정상 정리(Graceful cleanup) 완료 여부 판정
        #    - _cleanup_timed_out == False : 정상 정리 완료 → 다음 공정 진행 허용
        #    - _cleanup_timed_out == True  : 정상 정리 실패 → 자동 강제 복구(Force Recovery) 승격 시도
        # ------------------------------------------------------------------
        if not getattr(self, "_cleanup_timed_out", False):
            self._pending_device_cleanup = False
            # 강제 복구 카운터는 정상 정리 성공 시 리셋
            self._force_recover_count = 0
            return

        # ✅ 정상 정리 실패: 자동 강제 복구 시도
        ok_force = False
        with contextlib.suppress(Exception):
            ok_force = await asyncio.wait_for(
                self._force_recover_after_cleanup_timeout(pending_cleanup_names),
                timeout=25.0,
            )

        if ok_force:
            # 강제 복구가 성공하면 Start 제한을 해제하고 다음 공정으로 진행 가능
            self._cleanup_timed_out = False
            self._pending_device_cleanup = False
            self._force_recover_count = 0
            self.append_log("MAIN", "🧯 강제 복구 성공 → Start 제한 해제")
        else:
            # 강제 복구도 실패하면 안전상 Start 제한 유지
            self._pending_device_cleanup = True
            self.append_log("MAIN", "⚠ cleanup 미완료(타임아웃) + 강제 복구 실패 → Start 제한 유지")

    async def _force_recover_after_cleanup_timeout(self, pending_cleanup_names: list[str]) -> bool:
        """
        강제 복구(Force Recovery)

        목적:
        - 정상 정리(Graceful cleanup)가 timeout/누수로 실패했을 때,
          프로그램 재시작 없이 다음 공정을 계속할 수 있도록
          "통신/핸들/대기"를 최대한 강하게 끊고 내부 상태를 리셋한다.

        성공 기준(현실적 기준):
        - 여기서 예외 없이 수행되고, 최소한 'Start를 막아야 할 이유'를 줄였다고 판단되면 True
        - 반복 실패가 누적되면 False(재시작 권고)
        """
        # 가드: 너무 자주 강제복구하면 오히려 누수/불안정 증가
        self._force_recover_count = int(getattr(self, "_force_recover_count", 0)) + 1
        if self._force_recover_count > 3:
            self.append_log("MAIN", f"⚠ 강제 복구 {self._force_recover_count}회 반복 → 프로그램 재시작 권고")
            return False

        self.append_log("MAIN", f"🧯 강제 복구 시작 (count={self._force_recover_count}, pending={pending_cleanup_names!r})")

        # 0) 자동 연결/폴링 강제 OFF (통신이 계속 발생하면 MOXA idle disconnect도 안 걸릴 수 있음)
        self._auto_connect_enabled = False
        with contextlib.suppress(Exception):
            self._apply_polling_targets({"mfc": False, "dc_pulse": False, "rf_pulse": False, "dc": False, "rf": False})

        # 1) starter task들 취소 (있다면)
        try:
            st = getattr(self, "_starter_threads", None)
            if isinstance(st, dict) and st:
                live = [t for t in st.values() if isinstance(t, asyncio.Task) and (not t.done())]
                for t in live:
                    with contextlib.suppress(Exception):
                        t.cancel()
                if live:
                    with contextlib.suppress(asyncio.TimeoutError):
                        await asyncio.wait_for(asyncio.gather(*live, return_exceptions=True), timeout=2.0)
                st.clear()
        except Exception:
            pass

        # 2) 장치 transport를 “가능한 만큼” 강제로 close (best-effort)
        dev_list = [
            ("IG", self.ig),
            ("MFC", self.mfc),
            ("DCPulse", self.dc_pulse),
            ("RFPulse", self.rf_pulse),
            ("DCPower", self.dc_power),
            ("DCPower2", self.dc_power2),
            ("RFPower", self.rf_power),
            ("RGA", self.rga),
            ("OES", self.oes),
        ]
        for name, dev in dev_list:
            if dev is None:
                continue
            with contextlib.suppress(Exception):
                await asyncio.wait_for(self._hard_close_device_transport(dev, name), timeout=3.0)

        # 3) 로그 writer가 누수(leaked)된 케이스는 executor/queue를 재생성해서 로그 기능을 살린다
        try:
            t = getattr(self, "_log_writer_task", None)
            if isinstance(t, asyncio.Task) and (not t.done()):
                with contextlib.suppress(Exception):
                    t.cancel()
                with contextlib.suppress(asyncio.TimeoutError):
                    await asyncio.wait_for(asyncio.gather(t, return_exceptions=True), timeout=1.5)

            ex = getattr(self, "_log_io_exec", None)
            if ex is not None:
                with contextlib.suppress(Exception):
                    ex.shutdown(wait=False, cancel_futures=True)  # type: ignore[arg-type]

            self._log_io_exec = ThreadPoolExecutor(max_workers=1, thread_name_prefix=f"LogIO.CH{self.ch}")
            self._log_q = asyncio.Queue(maxsize=4096)
            self._log_writer_task = None
        except Exception:
            self.append_log("MAIN", "⚠ 강제 복구: log writer 재초기화 실패(무시)")

        # 4) 내부 상태 리셋 (다음 런에서 start/start_devices를 다시 태우도록)
        self._bg_started = False
        self._devices_started = False

        self.append_log("MAIN", "🧯 강제 복구 완료(최대한 복구 시도)")
        return True
    
    async def _hard_close_device_transport(self, dev: Any, label: str) -> None:
        """
        cleanup()이 timeout/누수로 실패했을 때, 가능한 방식으로 transport를 강제 close 한다.
        - 장치 클래스마다 메서드명이 달라서, 존재하는 메서드를 best-effort로 호출한다.
        """
        # 1) 장치가 제공하는 강제 종료류 메서드 우선 시도
        methods = [
            "force_close",
            "close_transport",
            "disconnect",
            "close",
            "shutdown",
            "stop",
            "abort",
            "reset_connection",
        ]

        for m in methods:
            fn = getattr(dev, m, None)
            if callable(fn):
                try:
                    r = fn()
                    if inspect.isawaitable(r):
                        await r
                except Exception:
                    pass

        # 2) 내부 transport 후보 속성에서 close() 시도
        candidate_attrs = [
            "_ser", "ser", "_serial", "serial",
            "_client", "client",
            "_sock", "sock", "_socket", "socket",
            "_writer", "writer",
            "_transport", "transport",
        ]

        for a in candidate_attrs:
            obj = getattr(dev, a, None)
            if obj is None:
                continue

            close_fn = getattr(obj, "close", None)
            if callable(close_fn):
                with contextlib.suppress(Exception):
                    r = close_fn()
                    if inspect.isawaitable(r):
                        await r

            wc = getattr(obj, "wait_closed", None)
            if callable(wc):
                with contextlib.suppress(Exception):
                    r = wc()
                    if inspect.isawaitable(r):
                        await r

        with contextlib.suppress(Exception):
            self.append_log("MAIN", f"🧯 force-close attempted: {label}")

    async def shutdown_fast_async(self) -> None:
        self._shutting_down = True
        self._auto_connect_enabled = False

        self._cancel_delay_task()

        try:
            if self.ig and hasattr(self.ig, "cancel_wait"):
                with contextlib.suppress(Exception):
                    await asyncio.wait_for(self.ig.cancel_wait(), timeout=1.0)
        except Exception:
            pass

        loop = asyncio.get_running_loop()
        current = asyncio.current_task()

        # ✅ runner / stage task 먼저 정지
        runner_tasks = []
        for t in (
            getattr(self, "_runner_stage_task", None),
            getattr(self, "_runner_task", None),
        ):
            if isinstance(t, asyncio.Task) and (not t.done()) and t is not current:
                runner_tasks.append(t)

        for t in runner_tasks:
            with contextlib.suppress(Exception):
                t.cancel()

        if runner_tasks:
            with contextlib.suppress(Exception):
                await asyncio.gather(*runner_tasks, return_exceptions=True)

        self._runner_stage_task = None
        self._runner_stage_kind = None
        self._runner_task = None

        # ✅ 일반 bg task 정지
        live = [t for t in getattr(self, "_bg_tasks", []) if t and not t.done() and t is not current]
        for t in live:
            with contextlib.suppress(Exception):
                loop.call_soon(t.cancel)

        if live:
            with contextlib.suppress(Exception):
                await asyncio.gather(*live, return_exceptions=True)

        self._bg_tasks = []

        # ✅ keepalive task 정지
        keepalive = [
            t for t in getattr(self, "_keepalive_tasks", {}).values()
            if t and not t.done() and t is not current
        ]
        for t in keepalive:
            with contextlib.suppress(Exception):
                loop.call_soon(t.cancel)

        if keepalive:
            with contextlib.suppress(Exception):
                await asyncio.gather(*keepalive, return_exceptions=True)

        self._keepalive_tasks = {}

        self._bg_started = False
        self._devices_started = False
        self._run_select = None

        # ✅ 장치 정리
        tasks = []
        for dev in (self.ig, self.mfc, self.dc_pulse, self.rf_pulse, self.dc_power, self.dc_power2, self.rf_power, self.oes, self.rga):
            if not dev:
                continue
            try:
                if hasattr(dev, "cleanup_quick"):
                    tasks.append(dev.cleanup_quick())
                elif hasattr(dev, "cleanup"):
                    tasks.append(dev.cleanup())
            except Exception:
                pass

        if tasks:
            await asyncio.gather(*tasks, return_exceptions=True)

        # ✅ 로그 마무리
        with contextlib.suppress(Exception):
            self._close_run_log()

        with contextlib.suppress(Exception):
            await self._shutdown_log_writer()

        self._log_file_path = None
        with contextlib.suppress(Exception):
            self._prestart_buf.clear()


    def shutdown_fast(self) -> None:
        self._spawn_detached(self.shutdown_fast_async(), name=f"ShutdownFast.CH{self.ch}")

    # ------------------------------------------------------------------
    # 입력 검증 / 정규화 / delay 처리
    def _get_text(self, leaf: str) -> str:
        w = self._u(leaf)
        return w.toPlainText().strip() if w else ""

    def _validate_single_run_inputs(self) -> dict[str, Any] | None:
        if self.ch == 1:
            use_ar = bool(getattr(self._u("Ar_checkbox"), "isChecked", lambda: False)())
            use_o2 = bool(getattr(self._u("O2_checkbox"), "isChecked", lambda: False)())
            use_n2 = bool(getattr(self._u("N2_checkbox"), "isChecked", lambda: False)())
            if not (use_ar or use_o2 or use_n2):
                self._post_warning("선택 오류", "가스를 하나 이상 선택해야 합니다.")
                return None

            def _read_flow(name: str) -> float:
                txt = self._get_text(name) or "0"
                v = float(txt)
                if v < 0:
                    raise ValueError()
                return v

            try:
                ar_flow = _read_flow("arFlow_edit") if use_ar else 0.0
                o2_flow = _read_flow("o2Flow_edit") if use_o2 else 0.0
                n2_flow = _read_flow("n2Flow_edit") if use_n2 else 0.0
            except Exception:
                self._post_warning("입력값 확인", "가스 유량 입력을 확인하세요.")
                return None

            # ✅ CH1에서 RF-Pulse 임시 사용: 지원 여부(supports) 반영
            use_rf_pulse = self.supports_rf_pulse and bool(getattr(self._u("rfPulsePower_checkbox"), "isChecked", lambda: False)())
            use_dc_pulse = self.supports_dc_pulse and bool(getattr(self._u("dcPulsePower_checkbox"), "isChecked", lambda: False)())

            # 최소 1개 선택
            if not (use_rf_pulse or use_dc_pulse):
                if self.supports_rf_pulse:
                    self._post_warning("선택 오류", "CH1은 RF-Pulse(또는 DC-Pulse) 중 하나를 선택해야 합니다.")
                elif self.supports_dc_pulse:
                    self._post_warning("선택 오류", "CH1은 DC-Pulse를 선택해야 합니다.")
                else:
                    self._post_warning("선택 오류", "CH1에서 사용 가능한 파워 장치가 없습니다.")
                return None

            # 기본값
            dc_pulse_power = 0.0; dc_pulse_freq = None; dc_pulse_duty = None
            rf_pulse_power = 0.0; rf_pulse_freq = None; rf_pulse_duty = None

            if use_dc_pulse:
                # ---- 기존 DC-Pulse 검증 유지 ----
                try:
                    dc_pulse_power = float(self._get_text("dcPulsePower_edit") or "0")
                    if dc_pulse_power <= 0:
                        raise ValueError()
                except ValueError:
                    self._post_warning("입력값 확인", "DC-Pulse Target Power(W)를 확인하세요.")
                    return None

                txtf = self._get_text("dcPulseFreq_edit")
                if txtf:
                    try:
                        dc_pulse_freq = int(float(txtf))  # kHz
                        if dc_pulse_freq < 20 or dc_pulse_freq > 150:
                            raise ValueError()
                    except ValueError:
                        self._post_warning("입력값 확인", "DC-Pulse Freq(kHz)는 20..150 범위입니다.")
                        return None

                txtd = self._get_text("dcPulseDutyCycle_edit")
                if txtd:
                    try:
                        dc_pulse_duty = int(float(txtd))
                        if dc_pulse_duty < 1 or dc_pulse_duty > 99:
                            raise ValueError()
                    except ValueError:
                        self._post_warning("입력값 확인", "DC-Pulse Duty(%)는 1..99 범위")
                        return None

            if use_rf_pulse:
                # ---- CH2에서 쓰는 RF-Pulse 검증 로직을 CH1에도 동일 적용 ----
                try:
                    rf_pulse_power = float(self._get_text("dcPulsePower_edit") or "0")
                    if rf_pulse_power <= 0:
                        raise ValueError()
                except ValueError:
                    self._post_warning("입력값 확인", "RF Pulse Target Power(W)를 확인하세요.")
                    return None

                txtf = self._get_text("dcPulseFreq_edit")
                if txtf:
                    try:
                        rf_pulse_freq = int(float(txtf))  # kHz
                        if rf_pulse_freq < 1 or rf_pulse_freq > 100:
                            raise ValueError()
                    except ValueError:
                        self._post_warning("입력값 확인", "RF Pulse Freq(kHz)는 1..100 범위입니다.")
                        return None

                txtd = self._get_text("dcPulseDutyCycle_edit")
                if txtd:
                    try:
                        rf_pulse_duty = int(float(txtd))
                        if rf_pulse_duty < 1 or rf_pulse_duty > 99:
                            raise ValueError()
                    except ValueError:
                        self._post_warning("입력값 확인", "RF Pulse Duty(%) 1..99")
                        return None

            g1n = self._get_text("g1Target_name")
            g2n = self._get_text("g2Target_name")
            g3n = self._get_text("g3Target_name")

            return {
                "use_ms": bool(getattr(self._u("mainShutter_checkbox"), "isChecked", lambda: False)()),
                "use_g1": False, "use_g2": False, "use_g3": False,
                "use_ar": use_ar, "use_o2": use_o2, "use_n2": use_n2,
                "ar_flow": ar_flow, "o2_flow": o2_flow, "n2_flow": n2_flow,

                # CH1 단일공정에서는 연속파는 일단 미사용(필요하면 확장)
                "use_rf_power": False, "use_dc_power": False, "use_dc_power2": False,

                "use_dc_pulse": use_dc_pulse,
                "dc_pulse_power": dc_pulse_power,
                "dc_pulse_freq": dc_pulse_freq,
                "dc_pulse_duty": dc_pulse_duty,

                "use_rf_pulse": use_rf_pulse,
                "rf_pulse_power": rf_pulse_power,
                "rf_pulse_freq": rf_pulse_freq,
                "rf_pulse_duty": rf_pulse_duty,

                "G1_target_name": g1n, "G2_target_name": g2n, "G3_target_name": g3n,
                "use_power_select": bool(getattr(self._u("powerSelect_checkbox"), "isChecked", lambda: False)()),
            }

        elif self.ch == 2:
            # -----------------------------
            # ✅ CH2 단일 실행(UI) 검증/파라미터 생성
            # -----------------------------
            use_ar = bool(getattr(self._u("Ar_checkbox"), "isChecked", lambda: False)())
            use_o2 = bool(getattr(self._u("O2_checkbox"), "isChecked", lambda: False)())
            use_n2 = bool(getattr(self._u("N2_checkbox"), "isChecked", lambda: False)())

            if not (use_ar or use_o2 or use_n2):
                self._post_warning("선택 오류", "가스를 하나 이상 선택해야 합니다.")
                return None

            def _read_flow(name: str) -> float:
                txt = self._get_text(name) or "0"
                v = float(txt)
                if v < 0:
                    raise ValueError()
                return v

            try:
                ar_flow = _read_flow("arFlow_edit") if use_ar else 0.0
                o2_flow = _read_flow("o2Flow_edit") if use_o2 else 0.0
                n2_flow = _read_flow("n2Flow_edit") if use_n2 else 0.0
            except Exception:
                self._post_warning("입력값 확인", "가스 유량 입력을 확인하세요.")
                return None
            
            # ✅ CH2: 연속 파워(DC/RF)
            use_dc_power = bool(getattr(self._u("dcPower_checkbox"), "isChecked", lambda: False)())
            use_dc_power2 = bool(getattr(self._u("dcPower2_checkbox"), "isChecked", lambda: False)())
            use_rf_power = bool(getattr(self._u("rfPower_checkbox"), "isChecked", lambda: False)())

            dc_power = 0.0
            dc_power2 = 0.0
            rf_power = 0.0

            if use_dc_power:
                try:
                    dc_power = float(self._get_text("dcPower_edit") or "0")
                    if dc_power <= 0:
                        raise ValueError()
                except Exception:
                    self._post_warning("입력값 확인", "DC Power(W)를 확인하세요.")
                    return None

            if use_dc_power2:
                try:
                    dc_power2 = float(self._get_text("dcPower2_edit") or "0")
                    if dc_power2 <= 0:
                        raise ValueError()
                except Exception:
                    self._post_warning("입력값 확인", "DC2 Power(W)를 확인하세요.")
                    return None

            if use_rf_power:
                try:
                    rf_power = float(self._get_text("rfPower_edit") or "0")
                    if rf_power <= 0:
                        raise ValueError()
                except Exception:
                    self._post_warning("입력값 확인", "RF Power(W)를 확인하세요.")
                    return None

            # ✅ CH2: Pulse 선택(DC-Pulse / RF-Pulse)
            use_dc_pulse = self.supports_dc_pulse and bool(getattr(self._u("dcPulsePower_checkbox"), "isChecked", lambda: False)())
            use_rf_pulse = self.supports_rf_pulse and bool(getattr(self._u("rfPulsePower_checkbox"), "isChecked", lambda: False)())

            # 기본값
            dc_pulse_power = 0.0
            dc_pulse_freq = None
            dc_pulse_duty = None

            rf_pulse_power = 0.0
            rf_pulse_freq = None
            rf_pulse_duty = None

            if use_dc_pulse:
                # ---- DC-Pulse 검증 (입력칸은 RF와 공유: CH1과 동일 UX) ----
                try:
                    dc_pulse_power = float(self._get_text("rfPulsePower_edit") or "0")
                    if dc_pulse_power <= 0:
                        raise ValueError()
                except Exception:
                    self._post_warning("입력값 확인", "DC-Pulse Target Power(W)를 확인하세요.")
                    return None

                txtf = self._get_text("rfPulseFreq_edit")
                if txtf:
                    try:
                        dc_pulse_freq = int(float(txtf))
                        if dc_pulse_freq < 20 or dc_pulse_freq > 150:
                            raise ValueError()
                    except Exception:
                        self._post_warning("입력값 확인", "DC-Pulse Freq(kHz)는 20..150 범위입니다.")
                        return None

                txtd = self._get_text("rfPulseDutyCycle_edit")
                if txtd:
                    try:
                        dc_pulse_duty = int(float(txtd))
                        if dc_pulse_duty < 1 or dc_pulse_duty > 99:
                            raise ValueError()
                    except Exception:
                        self._post_warning("입력값 확인", "DC-Pulse Duty(%)는 1..99 범위")
                        return None

            if use_rf_pulse:
                # ---- RF-Pulse 검증 ----
                try:
                    rf_pulse_power = float(self._get_text("rfPulsePower_edit") or "0")
                    if rf_pulse_power <= 0:
                        raise ValueError()
                except Exception:
                    self._post_warning("입력값 확인", "RF-Pulse Target Power(W)를 확인하세요.")
                    return None

                txtf = self._get_text("rfPulseFreq_edit")
                if txtf:
                    try:
                        rf_pulse_freq = int(float(txtf))  # kHz
                        if rf_pulse_freq < 1 or rf_pulse_freq > 100:
                            raise ValueError()
                    except Exception:
                        self._post_warning("입력값 확인", "RF-Pulse Freq(kHz)는 1..100 범위입니다.")
                        return None

                txtd = self._get_text("rfPulseDutyCycle_edit")
                if txtd:
                    try:
                        rf_pulse_duty = int(float(txtd))
                        if rf_pulse_duty < 1 or rf_pulse_duty > 99:
                            raise ValueError()
                    except Exception:
                        self._post_warning("입력값 확인", "RF-Pulse Duty(%)는 1..99 범위입니다.")
                        return None

            # ✅ 최소 1개 파워 동작 선택 확인(정책)
            if not (use_dc_power or use_dc_power2 or use_rf_power or use_dc_pulse or use_rf_pulse):
                self._post_warning("선택 오류", "DC1/DC2/RF Power 또는 DC/RF Pulse 중 하나 이상 선택해야 합니다.")
                return None

            g1n = self._get_text("g1Target_name")
            g2n = self._get_text("g2Target_name")
            g3n = self._get_text("g3Target_name")

            return {
                "use_ms": bool(getattr(self._u("mainShutter_checkbox"), "isChecked", lambda: False)()),
                "use_g1": bool(getattr(self._u("G1_checkbox"), "isChecked", lambda: False)()),
                "use_g2": bool(getattr(self._u("G2_checkbox"), "isChecked", lambda: False)()),
                "use_g3": bool(getattr(self._u("G3_checkbox"), "isChecked", lambda: False)()),
                "use_ar": use_ar, "use_o2": use_o2, "use_n2": use_n2,
                "ar_flow": ar_flow, "o2_flow": o2_flow, "n2_flow": n2_flow,

                "use_dc_power": use_dc_power,
                "dc_power": dc_power,
                "use_dc_power2": use_dc_power2,
                "dc_power2": dc_power2,
                "use_rf_power": use_rf_power,
                "rf_power": rf_power,

                "use_dc_pulse": use_dc_pulse,
                "dc_pulse_power": dc_pulse_power,
                "dc_pulse_freq": dc_pulse_freq,
                "dc_pulse_duty": dc_pulse_duty,

                # CH2도 RF-Pulse 사용 가능
                "use_rf_pulse": use_rf_pulse,
                "rf_pulse_power": rf_pulse_power,
                "rf_pulse_freq": rf_pulse_freq,
                "rf_pulse_duty": rf_pulse_duty,

                # ✅ P.W select는 Start를 막지 않음: 값만 전달
                "use_power_select": bool(getattr(self._u("powerSelect_checkbox"), "isChecked", lambda: False)()),

                "G1_target_name": g1n, "G2_target_name": g2n, "G3_target_name": g3n,
            }

        # (방어) 혹시 모르는 값
        return None

    def _normalize_params_for_process(self, raw: RawParams) -> NormParams:
        def tf(v): return str(v).strip().upper() in ("T","TRUE","1","Y","YES")
        def fget(key, default="0"):
            try: return float(str(raw.get(key, default)).strip())
            except Exception: return float(default)
        def iget(key, default="0"):
            try: return int(float(str(raw.get(key, default)).strip()))
            except Exception: return int(default)
        def iget_opt(key):
            s = str(raw.get(key, '')).strip()
            return int(float(s)) if s != '' else None

        _g1_raw = str(raw.get("G1 Target", "") or raw.get("G1_target_name", "")).strip()
        _g2_raw = str(raw.get("G2 Target", "") or raw.get("G2_target_name", "")).strip()
        _g3_raw = str(raw.get("G3 Target", "") or raw.get("G3_target_name", "")).strip()

        # ✅ CSV/params에 없으면 UI 위젯에서 직접 읽기 (DB 로딩값 포함)
        if not _g1_raw:
            _g1_raw = self._get_text("g1Target_name")
        if not _g2_raw:
            _g2_raw = self._get_text("g2Target_name")
        if not _g3_raw:
            _g3_raw = self._get_text("g3Target_name")

        _use_g1 = bool(raw.get("use_g1", False)) or tf(raw.get("gun1", "F"))
        _use_g2 = bool(raw.get("use_g2", False)) or tf(raw.get("gun2", "F"))
        _use_g3 = bool(raw.get("use_g3", False)) or tf(raw.get("gun3", "F"))

        if self.ch == 1:
            g1t, g2t, g3t = _g1_raw, "", ""
        else:
            g1t = _g1_raw if _use_g1 else ""
            g2t = _g2_raw if _use_g2 else ""
            g3t = _g3_raw if _use_g3 else ""

        # ▼ 추가: chuck_position(up/mid/down, 공란이면 스킵)
        _pos = str(raw.get("chuck_position", "")).strip().lower()
        if _pos not in ("up", "mid", "down"):
            _pos = ""

        def _opt_int_from_keys(*keys):
            for k in keys:
                s = str(raw.get(k, "")).strip()
                if s != "":
                    try:
                        return int(float(s))
                    except Exception:
                        return None
            return None

        def _float_from_keys(default: str, *keys):
            for k in keys:
                s = str(raw.get(k, "")).strip()
                if s != "":
                    try:
                        return float(s)
                    except Exception:
                        break
            return float(default)

        if self.ch == 1:
            # ✅ CH1: RF/DC Pulse 둘 다 CSV에서 독립적으로 사용 (두 컬럼 유지)
            use_dc_pulse = tf(raw.get("use_dc_pulse", "F"))
            dc_pulse_power = _float_from_keys("0", "dc_pulse_power")
            dc_pulse_freq  = _opt_int_from_keys("dc_pulse_freq")
            dc_pulse_duty  = _opt_int_from_keys("dc_pulse_duty_cycle")

            use_rf_pulse = tf(raw.get("use_rf_pulse", "F"))
            rf_pulse_power = _float_from_keys("0", "rf_pulse_power")
            rf_pulse_freq  = _opt_int_from_keys("rf_pulse_freq")
            rf_pulse_duty  = _opt_int_from_keys("rf_pulse_duty_cycle")

        else:
            # ✅ CH2: 기존 키 그대로 사용
            use_dc_pulse = tf(raw.get("use_dc_pulse", "F"))
            dc_pulse_power = fget("dc_pulse_power", "0")
            dc_pulse_freq  = iget_opt("dc_pulse_freq")
            dc_pulse_duty  = iget_opt("dc_pulse_duty_cycle")

            use_rf_pulse = tf(raw.get("use_rf_pulse", "F"))
            rf_pulse_power = fget("rf_pulse_power", "0")
            rf_pulse_freq  = iget_opt("rf_pulse_freq")
            rf_pulse_duty  = iget_opt("rf_pulse_duty_cycle")

        # ✅ dep.rate: CSV 헤더가 "dep.rate"(점) 또는 "dep_rate"(밑줄) 둘 다 지원
        def fget_deprate() -> float | None:
            for key in ("dep_rate", "dep.rate"):
                s = str(raw.get(key, "")).strip()
                if s:
                    try:
                        v = float(s)
                        return v if v > 0 else None
                    except Exception:
                        pass
            return None

        _dep_rate  = fget_deprate()
        _thickness = fget("thickness", "0") or None

        # ✅ process_time 역산: CSV에서 비어있고 dep_rate + thickness가 모두 있으면 자동 계산
        _process_time = fget("process_time", "0")
        if _process_time <= 0.0 and _dep_rate and _thickness:
            try:
                _process_time = float(_thickness) / float(_dep_rate) / 60.0  # nm / (nm/s) / 60 = 분
                self.append_log("Params", f"process_time 자동 계산: {_thickness}nm ÷ {_dep_rate}nm/s = {_process_time:.3f}분")
            except Exception:
                pass

        res: NormParams = {
            "base_pressure":     fget("base_pressure", "1e-5"),
            "working_pressure":  fget("working_pressure", "0"),
            "process_time":      _process_time,
            "dep_rate":          _dep_rate,
            "thickness":         _thickness,
            "shutter_delay":     fget("shutter_delay", "0"),
            "integration_time":  iget("integration_time", "60"),
            "dc_power":          fget("dc_power", "0"),
            "dc_power2":         fget("dc_power2", "0"),
            "rf_power":          fget("rf_power", "0"),

            "use_dc_pulse":      use_dc_pulse,
            "dc_pulse_power":    dc_pulse_power,
            "dc_pulse_freq":     dc_pulse_freq,
            "dc_pulse_duty":     dc_pulse_duty,

            "use_rf_pulse":      use_rf_pulse,
            "rf_pulse_power":    rf_pulse_power,
            "rf_pulse_freq":     rf_pulse_freq,
            "rf_pulse_duty":     rf_pulse_duty,

            "use_rf_power":      tf(raw.get("use_rf_power", "F")),
            "use_dc_power":      tf(raw.get("use_dc_power", "F")),
            "use_dc_power2":     tf(raw.get("use_dc_power2", "F")),
            "use_ar":            tf(raw.get("Ar", "F")),
            "use_o2":            tf(raw.get("O2", "F")),
            "use_n2":            tf(raw.get("N2", "F")),
            "ar_flow":           fget("Ar_flow", "0"),
            "o2_flow":           fget("O2_flow", "0"),
            "n2_flow":           fget("N2_flow", "0"),
            "use_g1":            _use_g1,
            "use_g2":            _use_g2,
            "use_g3":            _use_g3,
            "use_ms":            tf(raw.get("main_shutter", "F")),
            "process_note":      raw.get("Process_name", raw.get("process_note", "")),
            "G1_target_name":    g1t, "G2_target_name": g2t, "G3_target_name": g3t,
            "G1 Target":         g1t, "G2 Target": g2t, "G3 Target": g3t,
            "use_power_select":  tf(raw.get("power_select", "F")),

            # ★ 추가
            "chuck_position":    _pos,
        }

        # ✅ (옵션) DC Pulse 중간 power 변경용 컬럼 전달
        # - 비어있으면 ProcessController가 기존처럼 무시함
        res["power_change_time"] = str(raw.get("power_change_time", "") or "").strip()
        res["change_power_value"] = str(raw.get("change_power_value", "") or "").strip()

        # ------------------------------------------------------------
        # ✅ Pulse 파라미터 정규화
        #  - 체크(use_*) 뿐 아니라 값(power/freq/duty)로도 "요청"을 판단
        #  - CH1/CH2 각자 독립된 RF Pulse 장비 사용 (CH1: port 4008, CH2: port 4005)
        # ------------------------------------------------------------
        def _pos(v) -> bool:
            try:
                return float(v) > 0.0
            except Exception:
                return False

        rf_requested = (
            bool(res.get("use_rf_pulse"))
            or _pos(res.get("rf_pulse_power", 0.0))
            or (res.get("rf_pulse_freq") is not None)
            or (res.get("rf_pulse_duty") is not None)
        )

        dc_requested = (
            bool(res.get("use_dc_pulse"))
            or _pos(res.get("dc_pulse_power", 0.0))
            or (res.get("dc_pulse_freq") is not None)
            or (res.get("dc_pulse_duty") is not None)
        )

        if self.ch == 1:
            # ✅ CH1: RF/DC Pulse 둘 다 허용 (동시 사용 허용 정책)
            #    - 체크(use_*) 뿐 아니라 값(power/freq/duty)로도 "요청"을 판단
            if rf_requested:
                res["use_rf_pulse"] = True
            if dc_requested:
                res["use_dc_pulse"] = True
            # (참고) CH1의 연속 파워 차단은 settings ch1의 SUPPORTS_* 게이트가 담당

        elif self.ch == 2:
            # ✅ CH2: 값(power/freq/duty)이 들어오면 "요청"으로 간주해 use_rf_pulse를 True로 정규화
            if rf_requested:
                if self.supports_rf_pulse:
                    res["use_rf_pulse"] = True
                else:
                    self.append_log("Params", "CH2: RF-Pulse 미지원 설정 → RF-Pulse 요청을 무시하고 OFF 처리합니다.")
                    res["use_rf_pulse"] = False
                    res["rf_pulse_power"] = 0.0
                    res["rf_pulse_freq"] = None
                    res["rf_pulse_duty"] = None

        # 🔒 CH1은 N2 라인이 없으므로 강제 무시
        if self.ch == 1:
            if res.get("use_n2") or (res.get("n2_flow", 0.0) or 0.0) > 0.0:
                self.append_log("Params", "CH1은 N2 미지원 → N2 설정을 무시합니다.")
            res["use_n2"] = False
            res["n2_flow"] = 0.0

        return res

    # --- delay 단계 ---
    def _graph_reset_safe(self) -> None:
        try:
            self.graph.reset()
        except Exception:
            self.append_log("Graph", "reset skipped (headless)")

    def _graph_clear_rga_plot_safe(self) -> None:
        try:
            self.graph.clear_rga_plot()
        except Exception:
            self.append_log("Graph", "clear_rga_plot skipped (headless)")

    def _graph_update_rga_safe(self, x, y) -> None:
        try:
            x_list = x.tolist() if hasattr(x, "tolist") else x
            y_list = y.tolist() if hasattr(y, "tolist") else y
            self.graph.update_rga_plot(x_list, y_list)
        except Exception as e:
            self.append_log("Graph", f"update_rga_plot skipped: {e!r}")

    def _safe_clear_oes_plot(self) -> None:
        try: self.graph.clear_oes_plot()
        except Exception as e:
            self.append_log("OES", f"그래프 초기화 실패(무시): {e!r}")

    def _post_update_oes_plot(self, x: Sequence[float], y: Sequence[float]) -> None:
        def _safe_draw():
            try:
                # ✅ 종료 중이면 그래프 갱신 자체를 버림
                if getattr(self, "_shutting_down", False):
                    return

                graph = getattr(self, "graph", None)
                if not graph or not _qt_is_valid(graph):
                    return

                # ✅ 최소한 OES 그래프 핵심 객체가 살아있는지 확인
                for obj_name in ("oes_series", "oes_axis_x", "oes_axis_y"):
                    obj = getattr(graph, obj_name, None)
                    if obj is None or not _qt_is_valid(obj):
                        return

                xx = x.tolist() if hasattr(x, "tolist") else list(x)
                yy = y.tolist() if hasattr(y, "tolist") else list(y)
                graph.update_oes_plot(xx, yy)

            except Exception as e:
                self.append_log("OES", f"그래프 업데이트 실패(무시): {e!r}")

        self._soon(_safe_draw)

    # ------------------------------------------------------------------
    # 폴링/상태
    def _apply_polling_targets(self, targets: TargetsMap) -> None:
        mfc_on = bool(targets.get('mfc', False))
        dcpl_on = bool(targets.get('dc_pulse', False))
        rfpl_on = bool(targets.get('rf_pulse', False))
        dc_on   = bool(targets.get('dc', False))
        dc2_on  = bool(targets.get('dc2', False))
        rf_on   = bool(targets.get('rf', False))

        # ✅ 어떤 폴링이라도 실제로 켜야 할 때 + 자동연결 허용 + 공정 실행 중일 때만 자동 기동
        if (mfc_on or dcpl_on or rfpl_on or dc_on or dc2_on or rf_on) \
                and self._auto_connect_enabled \
                and self.process_controller.is_running:
            self._ensure_background_started()

        # ★ 챔버 공정에서 MFC 폴링을 켜는 시점에 mask reset
        #   (PC가 분리해놨을 수 있으므로 안전한 default 둘 다 True로 복귀.
        #    이로써 PC + 챔버 동시 실행 시 챔버의 R5/R60 폴링 보장)
        if mfc_on:
            with contextlib.suppress(Exception):
                if hasattr(self.mfc, "set_poll_mask"):
                    self.mfc.set_poll_mask(gas=True, pressure=True)

        with contextlib.suppress(Exception):
            self.mfc.set_process_status(mfc_on)

        if self.dc_pulse:
            with contextlib.suppress(Exception):
                # ✅ True/False 모두 직접 전달(다른 장치들과 일관)
                self.dc_pulse.set_process_status(dcpl_on)

        if self.rf_pulse:
            with contextlib.suppress(Exception):
                self.rf_pulse.set_process_status(rfpl_on)

        if self.dc_power and hasattr(self.dc_power, "set_process_status"):
            with contextlib.suppress(Exception):
                self.dc_power.set_process_status(dc_on)

        if self.dc_power2 and hasattr(self.dc_power2, "set_process_status"):
            with contextlib.suppress(Exception):
                self.dc_power2.set_process_status(dc2_on)

        if self.rf_power and hasattr(self.rf_power, "set_process_status"):
            with contextlib.suppress(Exception):
                self.rf_power.set_process_status(rf_on)

    # ------------------------------------------------------------------
    # 로그
    def _cam_log(self, msg: str) -> None:
        """CameraRecorder(백그라운드 스레드) 메시지를 공정 로그 + 화면으로 보낸다.
        append_log가 내부 _soon으로 이미 스레드 마샬링하므로 직접 호출해도 안전하다."""
        self.append_log("CAM", msg)

    def append_log(self, source: str, msg: str) -> None:
        now_ui = datetime.now().strftime("%H:%M:%S")
        now_file = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        line_ui = f"[{now_ui}] [CH{self.ch}:{source}] {msg}"
        line_file = f"[{now_file}] [CH{self.ch}:{source}] {msg}\n"

        self._soon(self._enqueue_ui_log, line_ui)

        if not getattr(self, "_log_file_path", None):
            buf = getattr(self, "_prestart_buf", None)
            if buf is not None:
                self._soon(buf.append, line_file)
            return
        self._soon(self._log_enqueue_nowait, line_file)

    def _dl_fire_and_forget(self, fn, *args, **kwargs) -> None:
        """
        DataLogger처럼 NAS/파일 I/O 가능성이 있는 동기 함수를
        이벤트루프(=UI)에서 직접 돌리지 않기 위한 안전 래퍼.
        """
        async def _run():
            try:
                # ✅ blocking I/O는 thread로
                await asyncio.to_thread(fn, *args, **kwargs)
            except Exception:
                # DataLogger 실패는 공정을 죽이면 안 되므로 조용히 무시(필요 시 rate-limit 로그만)
                pass

        self._spawn_detached(_run(), name=f"DL.{getattr(fn, '__name__', 'call')}.CH{self.ch}")

    def _enqueue_ui_log(self, line: str) -> None:
        # UI 스레드에서 호출되도록 _soon을 통해 들어온다고 가정
        self._ui_log_buf.append(line)

    def _flush_ui_log_to_ui(self) -> None:
        w = getattr(self, "_w_log", None)
        if not w or not _qt_is_valid(w):
            self._w_log = None
            return
        if not self._ui_log_buf:
            return

        try:
            sb = w.verticalScrollBar()
        except Exception:
            return

        # ✅ 사용자가 이미 최하단을 보고 있을 때만 '바닥에 붙이는' 오토 스크롤 유지
        stick_to_bottom = True
        try:
            stick_to_bottom = (sb.value() >= (sb.maximum() - 2))
        except Exception:
            stick_to_bottom = True

        # 한 번에 몰아서 출력 (UI 작업 최소화)
        lines = []
        max_lines = 200  # 100~300 사이 추천
        while self._ui_log_buf and len(lines) < max_lines:
            s = self._ui_log_buf.popleft()
            if s is None:
                continue
            try:
                s = str(s)
            except Exception:
                continue
            s = s.rstrip("\r\n")
            if not s:
                continue
            lines.append(s)

        if not lines:
            return

        prefix = ""
        try:
            prefix = "" if w.document().isEmpty() else "\n"
        except Exception:
            prefix = "\n"

        text = prefix + "\n".join(lines)

        old_sb_val = None
        if not stick_to_bottom:
            with contextlib.suppress(Exception):
                old_sb_val = sb.value()

        # ✅ 여기서도 invalid/타이밍 이슈를 최대한 방어
        try:
            w.moveCursor(QTextCursor.MoveOperation.End)
            w.insertPlainText(text)
        except Exception:
            return

        if old_sb_val is not None:
            with contextlib.suppress(Exception):
                sb.setValue(old_sb_val)
            return

        if not getattr(self, "_log_autoscroll_pending", False):
            self._log_autoscroll_pending = True

            def _scroll_bottom():
                self._log_autoscroll_pending = False
                ww = getattr(self, "_w_log", None)
                if not ww or not _qt_is_valid(ww):
                    self._w_log = None
                    return
                with contextlib.suppress(Exception):
                    sbb = ww.verticalScrollBar()
                    sbb.setValue(sbb.maximum())
                    ww.ensureCursorVisible()

            QTimer.singleShot(0, _scroll_bottom)

    def _ensure_log_dir(self, root: Path) -> Path:
        nas_path = Path(root)
        try:
            nas_path.mkdir(parents=True, exist_ok=True)
        except Exception:
            w = getattr(self, "_w_log", None)
            if w and _qt_is_valid(w):
                with contextlib.suppress(Exception):
                    w.appendPlainText(
                        f"[Logger] NAS 폴더 접근 실패 → 실제 기록 시 CH{self.ch} 로컬 폴백 예정: {self._local_log_dir}"
                    )
        return nas_path

    def _open_run_log(self, params: Mapping[str, Any]) -> None:
        now_local = datetime.now()
        ts = now_local.strftime("%Y%m%d_%H%M%S")
        
        raw_name = str(params.get("process_note") or params.get("Process_name") or "").strip()
        if not raw_name:
            raw_name = "Untitled"
        
        safe_name = re.sub(r'[\\/:*?"<>|]+', "_", raw_name)
        safe_name = re.sub(r"\s+", " ", safe_name).strip()
        safe_name = safe_name.replace(" ", "_")
        safe_name = safe_name.strip(" .")
        safe_name = safe_name[:60] if safe_name else "Untitled"
        
        # ✅ NAS path.exists()는 동기지만, 호출하는 곳에서 이미 executor 안에 있다면 OK
        # 여기서 더 안전하게: 마이크로초까지 timestamp를 넣어 충돌 가능성 최소화
        ts_ms = now_local.strftime("%Y%m%d_%H%M%S_%f")[:-3]  # ms 단위
        base = (self._log_dir / f"CH{self.ch}_{safe_name}_{ts_ms}").with_suffix(".txt")
        
        # ✅ 즉시 사용 (path.exists() 루프 제거 — 충돌 확률 거의 0)
        self._log_file_path = base
        
        if not self._log_writer_task or self._log_writer_task.done():
            self._set_task_later("_log_writer_task", self._log_writer_loop, name=f"LogWriter.CH{self.ch}")
        
        name = (params.get("process_note") or params.get("Process_name") or f"Run CH{self.ch}")

        # ✅ 헤더도 큐로 기록(순서 보장)
        self._log_enqueue_nowait("# ==== Sputter Run ====\n")
        self._log_enqueue_nowait(f"# started_at = {datetime.now().isoformat()}\n")
        self._log_enqueue_nowait(f"# chamber = CH{self.ch}\n")
        self._log_enqueue_nowait(f"# process_name = {name}\n")
        if "process_time" in params:
            self._log_enqueue_nowait(f"# time_min = {float(params.get('process_time', 0) or 0):.2f}\n")
        self._log_enqueue_nowait("# ============================\n")

        # ✅ pre-start 버퍼도 헤더 뒤로 밀어 넣기
        with contextlib.suppress(Exception):
            for line in list(self._prestart_buf):
                self._log_enqueue_nowait(line)
            self._prestart_buf.clear()

        self.append_log("Logger", f"새 로그 파일 시작: {base.name}")

    def _close_run_log(self) -> None:
        """종료 마커만 큐에 넣고, 실제 flush/close는 _shutdown_log_writer()에서 처리."""
        # ★ 이미 정리되어 path가 없으면 END 마커 중복 enqueue 금지
        #   (cleanup 후 _clear_queue_and_reset_ui가 또 호출되는 경로에서
        #    fallback 폴더에 'END만 있는 drain 파일'이 생기는 문제 차단)
        if not getattr(self, "_log_file_path", None):
            return
        with contextlib.suppress(Exception):
            self._log_enqueue_nowait("# ==== END ====\n")

    def _log_enqueue_nowait(self, line: str | None) -> None:
        try:
            self._log_q.put_nowait(line)
        except asyncio.QueueFull:
            with contextlib.suppress(Exception):
                _ = self._log_q.get_nowait()
                self._log_q.put_nowait(line)

    def _log_write_sync(self, path: Path, text: str) -> None:
        """
        ✅ keep-handle 방식 (util/log_hub.py의 SessionTextAppender 사용)
        - path가 바뀌면: close → open 교체
        - 같은 path면: 파일 핸들 유지한 채 write + flush
        """
        self._run_log_appender.set_primary_path(path)
        self._run_log_appender.write(text)

    async def _log_writer_loop(self):
        try:
            while True:
                # ✅ 폴링(get_nowait+sleep) 대신 “대기”로 CPU 절약 + 안정성↑
                line = await self._log_q.get()

                # ✅ 한 번에 배치로 모아서 write 횟수/flush 횟수 줄이기
                batch = [line]
                for _ in range(300):  # 배치 크기(원하면 조절)
                    try:
                        batch.append(self._log_q.get_nowait())
                    except asyncio.QueueEmpty:
                        break

                text = "".join(batch)

                # ✅ 파일 경로가 없으면 버림(또는 prestart_buf로 보내도 됨)
                if not self._log_file_path:
                    continue

                # ✅ open/write/flush는 무조건 스레드로
                try:
                    loop = asyncio.get_running_loop()

                    # SessionTextAppender가 내부에서
                    # 1) NAS 기록
                    # 2) 실패 시 self._local_log_dir 로 자동 폴백
                    # 을 처리하므로, 여기서는 "직접 로컬 파일 쓰기"를 하지 않는다.
                    await asyncio.wait_for(
                        loop.run_in_executor(
                            self._log_io_exec,
                            self._log_write_sync,
                            self._log_file_path,
                            text,
                        ),
                        timeout=5.0,
                    )

                    # ✅ SessionTextAppender가 이번 write에서 fallback으로 전환됐는지 1회 알림
                    try:
                        switched = await loop.run_in_executor(
                            self._log_io_exec,
                            self._run_log_appender.consume_switched_flag,
                        )
                        if switched:
                            cur_path = getattr(self._run_log_appender, "current_path", None)
                            self._soon(
                                self._enqueue_ui_log,
                                f"[{datetime.now().strftime('%H:%M:%S')}] [CH{self.ch}:Logger] 세션 로그를 로컬 폴백으로 전환: {cur_path}"
                            )
                    except Exception:
                        pass

                except Exception as e:
                    # ✅ 여기서 정상 run 파일명으로 직접 로컬 append 하지 않는다.
                    #    그렇게 하면 '부분 로그', 'end만 있는 로그', '중복 로그'가 생긴다.
                    try:
                        local_dir = self._local_log_dir
                        local_dir.mkdir(parents=True, exist_ok=True)
                        emergency_path = local_dir / f"CH{self.ch}_{datetime.now():%Y%m%d_%H%M%S}_writer_pending.txt"

                        def _write_emergency() -> None:
                            with open(emergency_path, "a", encoding="utf-8", newline="") as fp:
                                fp.write("# [log-writer pending]\n")
                                if self._log_file_path:
                                    fp.write(f"# source_path = {self._log_file_path}\n")
                                fp.write(f"# reason = {e!r}\n")
                                fp.write(text)

                        await loop.run_in_executor(self._log_io_exec, _write_emergency)

                        self._soon(
                            self._enqueue_ui_log,
                            f"[{datetime.now().strftime('%H:%M:%S')}] [CH{self.ch}:Logger] 세션 로그 긴급 보관 파일 생성: {emergency_path}"
                        )
                    except Exception:
                        pass

        except asyncio.CancelledError:
            pass

    async def _shutdown_log_writer(self, path_override: Path | None = None):
        """
        - 로그 writer task를 종료
        - 큐에 남은 로그를 최종 기록
        - keep-handle(appender) 핸들을 확실히 close
        """
        async with self._log_shutdown_lock:
            loop = asyncio.get_running_loop()
            path = path_override if path_override is not None else self._log_file_path

            try:
                # 1) writer task 종료
                t = getattr(self, "_log_writer_task", None)
                self._log_writer_task = None
                if t:
                    t.cancel()
                    with contextlib.suppress(Exception):
                        await asyncio.wait_for(asyncio.gather(t, return_exceptions=True), timeout=3.0)

                # 2) 큐 drain
                drained: list[str] = []
                q = getattr(self, "_log_q", None)
                if q is not None:
                    while True:
                        try:
                            drained.append(q.get_nowait())
                        except asyncio.QueueEmpty:
                            break

                # 3) 남은 로그 최종 기록
                if drained:
                    text = "".join(drained)

                    def _write_local_drain() -> None:
                        local_dir = self._local_log_dir
                        local_dir.mkdir(parents=True, exist_ok=True)
                        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                        drain_path = local_dir / f"CH{self.ch}_{ts}_shutdown_drain.txt"
                        with open(drain_path, "a", encoding="utf-8", newline="") as fp:
                            fp.write(text)

                    if path:
                        try:
                            await asyncio.wait_for(
                                loop.run_in_executor(self._log_io_exec, self._log_write_sync, path, text),
                                timeout=6.0
                            )
                        except Exception as e:
                            def _write_local_same_name() -> None:
                                local_dir = self._local_log_dir
                                local_dir.mkdir(parents=True, exist_ok=True)

                                target = local_dir / Path(path).name
                                with open(target, "a", encoding="utf-8", newline="") as fp:
                                    fp.write(text)
                                with open(target, "a", encoding="utf-8", newline="") as fp:
                                    fp.write(f"# [shutdown] final flush failed: {e!r}\n")

                            with contextlib.suppress(Exception):
                                await asyncio.wait_for(
                                    loop.run_in_executor(self._log_io_exec, _write_local_same_name),
                                    timeout=2.0
                                )

                            w = getattr(self, "_w_log", None)
                            if w and _qt_is_valid(w):
                                with contextlib.suppress(Exception):
                                    w.appendPlainText(f"[Logger] ⚠ final flush failed → local fallback (reason={e!r})")
                    else:
                        # ★ path가 없는 시점에 drained가 END 마커 하나뿐이면
                        #   의미 있는 데이터가 없으므로 fallback drain 파일을 만들지 않는다.
                        meaningful = any(
                            line and line.strip() and "==== END ====" not in line
                            for line in drained
                        )
                        if not meaningful:
                            pass  # skip drain 파일 생성
                        else:
                            with contextlib.suppress(Exception):
                                await asyncio.wait_for(
                                    loop.run_in_executor(self._log_io_exec, _write_local_drain),
                                    timeout=2.0
                                )

                # 4) keep-handle 닫기
                app = getattr(self, "_run_log_appender", None)
                if app is not None:
                    with contextlib.suppress(Exception):
                        await asyncio.wait_for(loop.run_in_executor(self._log_io_exec, app.close), timeout=2.0)

            finally:
                with contextlib.suppress(Exception):
                    await asyncio.wait_for(
                        loop.run_in_executor(self._log_io_exec, self._run_log_appender.close),
                        timeout=2.0,
                    )

                self._log_file_path = None

                while True:
                    try:
                        self._log_q.get_nowait()
                    except asyncio.QueueEmpty:
                        break

    def _clear_queue_and_reset_ui(self) -> None:
        # 전역 runtime_state로 종료 시각을 기록하므로 로컬 타임스탬프는 불필요
        # ★ 추가: 남아 있을 수 있는 카운트다운 태스크 정리
        #self._cancel_delay_task()

        # 1) 리스트 공정 인덱스/큐 초기화
        self.current_process_index = -1
        
        # ★ 핵심: 리스트 공정 큐까지 비워서 다음 Start는 단일 공정(UI 기반)으로만 동작하게
        try:
            if hasattr(self, "process_queue"):
                # 남아 있는 CSV 공정 리스트 제거
                self.process_queue.clear()
        except Exception:
            # 혹시 구조가 꼬여 있어도 다음 런에 영향 없도록 빈 리스트로 강제 재설정
            self.process_queue = []
    
        # 2) UI 리셋
        self._reset_ui_after_process()

        # 3) 로그 파일 / 로그 writer 정리
        shutdown_path = self._log_file_path

        with contextlib.suppress(Exception):
            self._close_run_log()

        with contextlib.suppress(Exception):
            self._spawn_detached(self._shutdown_log_writer(path_override=shutdown_path))

        # 4) 프리스타트 버퍼 정리 (한 번만 호출해도 충분)
        with contextlib.suppress(Exception):
            self._prestart_buf.clear()

        # 5) 종료 관련 내부 플래그
        # - cleanup 타임아웃이면 Start를 막아야 안전(“정리 덜 끝났는데 idle” 방지)
        if not getattr(self, "_cleanup_timed_out", False):
            self._pending_device_cleanup = False
        else:
            self.append_log("MAIN", "⚠ cleanup timeout 상태 유지: Start 제한 유지")

    # ------------------------------------------------------------------
    # 기본 UI값/리셋
    def _set_default_ui_values(self) -> None:
        _set = self._set
        
        _set("integrationTime_edit", "")
        _set("workingPressure_edit", "2")
        _set("arFlow_edit", "20")
        _set("o2Flow_edit", "0")
        _set("n2Flow_edit", "0")
        _set("dcPower_edit", "130")
        _set("dcPower2_edit", "0")

        # DC-Pulse (라디오는 전용 세터: exclusive 그룹 해제 트릭 필요)
        self._set_pulse_radio("dcPulsePower_checkbox", False)
        _set("dcPulsePower_edit", "200")
        _set("dcPulseFreq_edit", "")
        _set("dcPulseDutyCycle_edit", "")

        # RF-Pulse (라디오는 전용 세터)
        self._set_pulse_radio("rfPulsePower_checkbox", False)
        _set("rfPulsePower_edit", "100")
        _set("rfPulseFreq_edit", "")
        _set("rfPulseDutyCycle_edit", "")

        # RF-Power
        _set("rfPower_checkbox", False)
        _set("rfPower_edit", "0")

        # ← 추가: 챔버별 기본 체크
        try:
            if self.ch == 1:
                _set("integrationTime_edit", "RF pulse_pre_30min")  # ★ 추가: Process Name (위젯은 alias로 intergrationTime_edit으로 매핑됨)
                _set("basePressure_edit", "5e-6")             # 5.00E-06
                _set("workingPressure_edit", "5")             # 5 mTorr
                _set("arFlow_edit", "20")                     # 20 sccm
                _set("Ar_checkbox", True)
                _set("rfPulsePower_checkbox", True)           # CH1: RF Pulse 사용 (기본)
                _set("dcPulsePower_checkbox", False)
                _set("dcPower_checkbox", False)
                _set("dcPower_edit", "0")                     # RF Pulse 운전이므로 DC=0
                _set("rfPulsePower_edit", "250")              # 250 W (CH1은 dcPulsePower_edit로 alias됨)
                _set("rfPulseFreq_edit", "20")                # 20 kHz
                _set("rfPulseDutyCycle_edit", "80")           # 80 %
                _set("G1_checkbox", True)                     # gun1 사용
                _set("mainShutter_checkbox", True)            # ★ 추가: main_shutter=T
                _set("shutterDelay_edit", "5")                # 5 초
                _set("processTime_edit", "25")                # ★ 변경: 주석 해제, 25분 기본값
            elif self.ch == 2:
                _set("basePressure_edit", "9e-6")
                _set("G2_checkbox", True)             # CH2: G2 사용
                _set("Ar_checkbox", True)             # CH2: Ar 가스
                _set("dcPower_checkbox", True)        # CH2: DC Power 사용
                _set("dcPower2_checkbox", False)      # DC2는 기본 미사용
                _set("dcPower2_edit", "0")
                _set("dcPulsePower_checkbox", False)
                _set("shutterDelay_edit", "5")
                #_set("processTime_edit", "25")
        except Exception:
            pass

        self._apply_ui_lockouts()

    def _apply_ui_lockouts(self) -> None:
        """챔버별로 선택/입력을 막아야 하는 UI를 비활성화한다."""
        if not getattr(self, "ui", None):
            return

        def _raw(leaf: str):
            return getattr(self.ui, f"{self.prefix}{leaf}", None)

        def _disable(leaf: str):
            w = _raw(leaf)
            if w is None:
                return
            try:
                w.setEnabled(False)
            except Exception:
                pass

        # ✅ CH2에서도 RF-Pulse 가능
        #    단, 입력칸(rfPulse*_edit)은 DC-Pulse와 공유하므로:
        #    - RF 미지원이어도 DC-Pulse를 지원하면 입력칸은 살리고 RF 라디오만 비활성화
        #    - RF/DC 둘 다 미지원일 때만 입력칸까지 비활성화
        if (self.ch == 2) and (not self.supports_rf_pulse):
            _disable("rfPulsePower_checkbox")
            if not self.supports_dc_pulse:
                for leaf in ("rfPulsePower_edit", "rfPulseFreq_edit", "rfPulseDutyCycle_edit"):
                    _disable(leaf)
        if (self.ch == 2) and (not self.supports_dc_pulse):
            _disable("dcPulsePower_checkbox")

    def _reset_ui_after_process(self):
        self._set_default_ui_values()

        # ✅ 타겟명 초기화 (공통 leaf 사용 → CH1은 단일 위젯으로 alias 매핑됨)
        for leaf in ("g1Target_name", "g2Target_name", "g3Target_name"):
            # CH1에선 세 leaf가 모두 같은 'gunTarget_name'으로 alias 되지만, 같은 위젯을 여러 번 비워도 무해
            self._set(leaf, "")

        for name in (
            "G1_checkbox","G2_checkbox","G3_checkbox","Ar_checkbox","O2_checkbox","N2_checkbox",
            "mainShutter_checkbox","dcPulsePower_checkbox","rfPulsePower_checkbox","dcPower_checkbox","dcPower2_checkbox","powerSelect_checkbox",
        ):
            w = self._u(name)
            if w is not None:
                with contextlib.suppress(Exception):
                    w.setChecked(False)
        
        # ← 추가: 챔버별 기본 체크 복원
        try:
            if self.ch == 1:
                self._u("Ar_checkbox") and self._u("Ar_checkbox").setChecked(True)
                self._u("rfPulsePower_checkbox") and self._u("rfPulsePower_checkbox").setChecked(True)
                self._u("G1_checkbox") and self._u("G1_checkbox").setChecked(True)
                self._u("mainShutter_checkbox") and self._u("mainShutter_checkbox").setChecked(True)  # ★ 추가
            elif self.ch == 2:
                self._u("G2_checkbox") and self._u("G2_checkbox").setChecked(True)
                self._u("Ar_checkbox") and self._u("Ar_checkbox").setChecked(True)
                self._u("dcPower_checkbox") and self._u("dcPower_checkbox").setChecked(True)
        except Exception:
            pass

        _s = self._u("processState_edit")
        if _s: _s.setPlainText("대기 중")

        # ★ DC 표시 버퍼 초기화 (다음 공정에서 미사용 유닛이 '-' 로 나오도록)
        self._dc_disp = {1: None, 2: None}

        for leaf in ("Power_edit","Voltage_edit","Current_edit","forP_edit","refP_edit"):
            w = self._u(leaf)
            if w: w.setPlainText("")

        self._on_process_status_changed(False)
        with contextlib.suppress(Exception):
            self.graph.reset()

        # ✅ 공정 종료 후 타겟 초기화됐으므로 DB에서 재로드
        asyncio.ensure_future(self._load_gun_targets())

    # ======= 서버 통신 api =======
    def _host_report_start(self, ok: bool, reason: str = "") -> None:
        fut = getattr(self, "_host_start_future", None)
        if fut is not None and not fut.done():
            fut.set_result((bool(ok), str(reason)))

    async def start_with_recipe_string(self, recipe: str) -> None:
        """
        Host 진입점(서버/원격 호출용):
        - UI Start 버튼과 동일한 시작 경로(_handle_start_clicked)를 사용한다.
        - 실제 프리플라이트/시작/정리/다음 공정 진행은 Runner(ChamberRuntime 내부)가 담당한다.
        - 여기서는 '시작 가드 통과 여부(=프리플라이트 진입/거절)' 결과만 Future로 짧게 대기한다.
        """
        loop = asyncio.get_running_loop()
        self._host_start_future = loop.create_future()

        s = (recipe or "").strip()
        if not s:
            # 현재 UI 값으로 단발 시작 (버튼과 동일 경로)
            self._handle_start_clicked(False)
        elif s.lower().endswith(".csv"):
            # ✅ 동기 NAS 호출을 executor로 분리 (asyncio loop block 방지)
            loop = asyncio.get_running_loop()
            
            def _load_csv_sync(path: str):
                if not os.path.exists(path):
                    raise RuntimeError(f"CSV 파일을 찾을 수 없습니다: {path}")
                rows = []
                with open(path, mode='r', encoding='utf-8-sig', newline='') as csvfile:
                    reader = csv.DictReader(csvfile)
                    for row in reader:
                        name = (row.get('Process_name') or row.get('#') 
                                or f"공정 {len(rows)+1}").strip()
                        row['Process_name'] = name
                        rows.append(row)
                return rows
            
            try:
                # NAS open + read를 thread pool에서 실행, 30초 timeout
                rows = await asyncio.wait_for(
                    loop.run_in_executor(None, _load_csv_sync, s),
                    timeout=30.0
                )
            except asyncio.TimeoutError:
                raise RuntimeError(f"CSV 로드 30초 timeout (NAS 응답 지연): {s}")
            
            self.process_queue = [cast(RawParams, r) for r in rows]
            self.current_process_index = -1
            if not self.process_queue:
                raise RuntimeError("CSV에 공정 데이터가 없습니다.")
            self._update_ui_from_params(self.process_queue[0])
            self.append_log("File", f"CSV 로드 완료: {s} (총 {len(self.process_queue)}개)")

            self._handle_start_clicked(False)
        else:
            raise RuntimeError("지원하지 않는 레시피 형식입니다. CSV 경로만 허용됩니다.")

        # ✅ 시작 가드(=프리플라이트 진입/거절) 결과만 짧게 대기
        host_start_wait_timeout_s = float(
            self.cfg._get("CHAMBER_HOST_START_WAIT_TIMEOUT_S", 10.0)
        )

        try:
            ok, reason = await asyncio.wait_for(
                self._host_start_future,
                timeout=host_start_wait_timeout_s,
            )
        except asyncio.TimeoutError:
            raise RuntimeError(
                f"preflight timeout ({host_start_wait_timeout_s:.1f}s) "
                "(start guard 또는 내부 대기로 인해 프리플라이트 미도달)"
            )
        finally:
            self._host_start_future = None

        if not ok:
            raise RuntimeError(reason)
        # ok면 그대로 반환 (공정은 내부에서 계속 진행)

    # ------------------------------------------------------------------
    # 유틸
    # ------------------------------------------------------------------
    def _parse_duration_seconds(self, s: str) -> float:
        """
        '10s', '1m', '1h30m', '2h' 형태 문자열을 초 단위로 변환.
        """
        if not s:
            return 0.0
        s = s.replace(" ", "").lower()
        pattern = r"(?:(\d+(?:\.\d+)?)h)?(?:(\d+(?:\.\d+)?)m)?(?:(\d+(?:\.\d+)?)s)?"
        m = re.match(pattern, s)
        if not m:
            try:
                return float(s) * 60.0  # 단위 없으면 분으로 처리
            except Exception:
                return 0.0
        h = float(m.group(1) or 0)
        m_ = float(m.group(2) or 0)
        s_ = float(m.group(3) or 0)
        return h * 3600 + m_ * 60 + s_

    def _spawn_detached(
        self,
        coro: Coroutine[Any, Any, Any],
        *,
        store: bool = False,
        name: str | None = None,
    ) -> asyncio.Task | None:
        """
        ✅ 개선점
        - 같은 이벤트루프 스레드에서 호출되면 즉시 create_task() 해서 Task를 반환
        - 다른 스레드면 call_soon_threadsafe로 예약하고 None 반환
        - create_task 실패/태스크 예외는 반드시 로그로 남김
        """
        loop = self._loop

        def _attach_done_log(t: asyncio.Task) -> None:
            def _done(task: asyncio.Task) -> None:
                if task.cancelled():
                    return
                try:
                    exc = task.exception()
                except Exception as e:
                    self.append_log(f"Task{self.ch}", f"[{name or 'task'}] exception() failed: {e!r}")
                    return
                if exc:
                    tb = "".join(traceback.format_exception(type(exc), exc, exc.__traceback__)).rstrip()
                    self.append_log(f"Task{self.ch}", f"[{name or 'task'}] crashed:\n{tb}")
            t.add_done_callback(_done)

        def _create_here() -> asyncio.Task | None:
            try:
                t = loop.create_task(coro, name=name)
            except Exception as e:
                with contextlib.suppress(Exception):
                    coro.close()
                tb = "".join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
                self.append_log(f"Task{self.ch}", f"[{name or 'task'}] create_task failed:\n{tb}")
                return None

            _attach_done_log(t)
            if store:
                with contextlib.suppress(Exception):
                    self._bg_tasks.append(t)
            return t

        # 같은 루프면 즉시 생성
        try:
            running = asyncio.get_running_loop()
        except RuntimeError:
            running = None

        if running is loop:
            return _create_here()

        # 다른 스레드면 예약(반환 없음)
        def _create_later() -> None:
            _create_here()

        with contextlib.suppress(Exception):
            loop.call_soon_threadsafe(_create_later)
        return None

    def _set_task_later(
        self,
        attr_name: str,
        coro_factory: Callable[[], Coroutine[Any, Any, Any]],
        *,
        name: str | None = None
    ) -> None:
        """UI/다른 스레드 어디서든 안전하게 task를 만들고, 중복 생성을 방지한다."""
        loop = self._loop

        def _create_and_set():
            # ✅ 추가: 예약 실행 시점에 이미 살아있으면 중복 생성 금지
            exist = getattr(self, attr_name, None)
            if isinstance(exist, asyncio.Task) and (not exist.done()):
                return

            # ✅ 추가: 코루틴은 loop 컨텍스트에서 생성
            try:
                coro = coro_factory()
            except Exception as e:
                tb = "".join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
                self.append_log("Task", f"[{name or attr_name}] coro_factory failed:\n{tb}")
                setattr(self, attr_name, None)
                return

            try:
                t = loop.create_task(coro, name=name)
            except Exception as e:
                with contextlib.suppress(Exception):
                    coro.close()
                tb = "".join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
                self.append_log("Task", f"[{name or attr_name}] create_task failed:\n{tb}")

                if attr_name in ("_log_writer_task",):
                    with contextlib.suppress(Exception):
                        self._set_state_text("내부 오류: 태스크 생성 실패(로그 확인)")
                    with contextlib.suppress(Exception):
                        self._post_critical(
                            "내부 오류",
                            "백그라운드 작업 생성에 실패했습니다.\n"
                            "프로그램이 멈춘 것처럼 보일 수 있습니다.\n\n"
                            "자세한 내용은 로그 파일(또는 터미널)을 확인해주세요.",
                        )
                setattr(self, attr_name, None)
                return

            setattr(self, attr_name, t)

            def _done(task: asyncio.Task):
                if task.cancelled():
                    return
                try:
                    exc = task.exception()
                except Exception as e2:
                    self.append_log("Task", f"[{name or attr_name}] exception() failed: {e2!r}")
                    return
                if exc:
                    tb2 = "".join(traceback.format_exception(type(exc), exc, exc.__traceback__)).rstrip()
                    self.append_log("Task", f"[{name or attr_name}] crashed:\n{tb2}")

                with contextlib.suppress(Exception):
                    if getattr(self, attr_name, None) is task:
                        setattr(self, attr_name, None)

            t.add_done_callback(_done)

        try:
            running = asyncio.get_running_loop()
        except RuntimeError:
            running = None

        if running is loop:
            loop.call_soon(_create_and_set)
        else:
            loop.call_soon_threadsafe(_create_and_set)

    def _loop_from_anywhere(self) -> asyncio.AbstractEventLoop:
        try: return asyncio.get_running_loop()
        except RuntimeError: return self._loop

    def _soon(self, fn: Callable[..., Any], *args: Any, **kwargs: Any) -> None:
        def _safe():
            try:
                # ✅ 종료 중이면 예약된 콜백도 실행하지 않음
                if getattr(self, "_shutting_down", False):
                    return
                fn(*args, **kwargs)
            except Exception as e:
                tb = ''.join(traceback.format_exception(type(e), e, e.__traceback__)).rstrip()
                self.append_log(f"CB{self.ch}", f"callback failed:\n{tb}")

        loop = getattr(self, "_loop", None)
        if loop is None:
            return

        # ✅ 종료 중이면 새 콜백 예약 자체를 막음
        if getattr(self, "_shutting_down", False):
            return

        with contextlib.suppress(Exception):
            if loop.is_closed():
                return

        try:
            running = asyncio.get_running_loop()
        except RuntimeError:
            running = None

        try:
            if running is loop:
                loop.call_soon(_safe)
            else:
                loop.call_soon_threadsafe(_safe)
        except RuntimeError:
            # loop 종료 직전이면 조용히 무시
            pass

    def _is_dev_connected(self, dev: object) -> bool:
        try:
            v = getattr(dev, "is_connected", None)
            if callable(v): return bool(v())
            if isinstance(v, bool): return v
        except Exception:
            pass
        try: return bool(getattr(dev, "_connected", False))
        except Exception: return False

    async def _preflight_progress_log(self, need: list[tuple[str, object]], stop_evt: asyncio.Event) -> None:
        """프리플라이트 진행 상황을 1초 간격으로 계속 로그에 남긴다.

        기존 구현은 wait_for(timeout=1.0)에서 TimeoutError가 한 번 발생하면 함수가 종료되어
        '연결 대기 중' 로그가 1회만 남고 이후 진행 상황이 보이지 않는 문제가 있었다.
        """
        try:
            while not stop_evt.is_set():
                missing = [name for name, dev in need if not self._is_dev_connected(dev)]
                txt = ", ".join(missing) if missing else "모두 연결됨"
                self.append_log("MAIN", f"연결 대기 중: {txt}")

                try:
                    await asyncio.wait_for(stop_evt.wait(), timeout=1.0)
                except asyncio.TimeoutError:
                    continue
        except asyncio.CancelledError:
            raise
        except Exception as e:
            self.append_log("MAIN", f"프리플라이트 진행 로그 예외: {e!r}")

    # --- UI 위젯 접근/부모/다이얼로그 관리 -----------------------------------
    def _alias_leaf(self, leaf: str) -> str:
        """CH1의 UI 위젯 이름과 공통 이름을 매핑.
        주의: 실제 속성 접근은 getattr(self.ui, f"{self.prefix}{name}") 이므로,
        여기서는 prefix(예: 'ch1_')를 절대 포함하지 않는다.
        """
        if self.ch != 1:
            return leaf
        return {
            "integrationTime_edit": "intergrationTime_edit",

            # CH1은 단일 타겟 위젯: ch1_gunTarget_name
            "g1Target_name": "gunTarget_name",
            "g2Target_name": "gunTarget_name",
            "g3Target_name": "gunTarget_name",

            # ✅ CH1: Pulse 입력(Edit/Freq/Duty)은 1세트를 공용(dcPulse*)으로 유지한다.
            #    - RF/DC 선택은 라디오(rfPulsePower_checkbox / dcPulsePower_checkbox)로 처리
            #    - 따라서 edit/freq/duty만 rfPulse* → dcPulse*로 alias 한다.
            "rfPulsePower_edit": "dcPulsePower_edit",
            "rfPulseFreq_edit": "dcPulseFreq_edit",
            "rfPulseDutyCycle_edit": "dcPulseDutyCycle_edit",

        }.get(leaf, leaf)

    def _u(self, name: str) -> Any | None:
        """prefix+name 위젯을 가져온다. 없으면 None."""
        name = self._alias_leaf(name)
        if not getattr(self, "ui", None):
            return None

        w = getattr(self.ui, f"{self.prefix}{name}", None)
        if w is None:
            return None
        if not _qt_is_valid(w):
            return None
        return w

    def _parent_widget(self) -> Any | None:
        """메시지/파일 다이얼로그의 합리적 부모 위젯을 찾는다."""
        for leaf in ("Start_button", "Stop_button", "processState_edit", "logMessage_edit"):
            w = self._u(leaf)
            if w is not None:
                try:
                    return w.window()
                except Exception:
                    return w
        return None

    async def _aopen_file(self, caption="CSV 선택", start_dir="", 
                        name_filter="CSV Files (*.csv);;All Files (*.*)") -> str:
        if not self._has_ui():
            self.append_log("File", "headless: 파일 선택 UI 생략"); return ""

        dlg = QFileDialog(self._parent_widget() or None, caption, "", name_filter)
        dlg.setFileMode(QFileDialog.ExistingFile)

        # ✅ Windows 쉘/COM(네이티브 파일 다이얼로그) 우회
        with contextlib.suppress(Exception):
            dlg.setOption(QFileDialog.Option.DontUseNativeDialog, True)
        with contextlib.suppress(Exception):
            dlg.setOption(QFileDialog.DontUseNativeDialog, True)
        with contextlib.suppress(Exception):
            dlg.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose, True)
            dlg.setWindowModality(Qt.WindowModality.WindowModal)

        loop = asyncio.get_running_loop()
        fut: asyncio.Future[str] = loop.create_future()

        def _done(result: int):
            # ✅ finished가 중복으로 들어와도 Future를 두 번 set_result하지 않게
            if fut.done():
                return
            try:
                if result == QDialog.Accepted and dlg.selectedFiles():
                    fut.set_result(dlg.selectedFiles()[0])
                else:
                    fut.set_result("")  # 취소
            finally:
                with contextlib.suppress(Exception):
                    dlg.finished.disconnect(_done)
                dlg.deleteLater()

        dlg.finished.connect(_done)
        dlg.open()
        if start_dir:
            QTimer.singleShot(50, lambda: dlg.setDirectory(start_dir) if _qt_is_valid(dlg) else None)

        try:
            return await fut
        finally:
            # await 중 취소/예외에도 다이얼로그 정리
            if _qt_is_valid(dlg):
                with contextlib.suppress(Exception):
                    dlg.close()
                with contextlib.suppress(Exception):
                    dlg.deleteLater()

    def _ensure_msgbox_store(self):
        if not hasattr(self, "_msg_boxes"):
            self._msg_boxes = []

    def _post_warning(self, title: str, text: str, auto_close_ms: int = 5000) -> None:
        if not self._has_ui():
            self.append_log("WARN", f"{title}: {text}"); return

        self._ensure_msgbox_store()
        box = QMessageBox(self._parent_widget() or None)
        box.setWindowTitle(title)
        box.setText(text)
        box.setIcon(QMessageBox.Warning)
        box.setStandardButtons(QMessageBox.Ok)
        box.setWindowModality(Qt.WindowModality.WindowModal)
        box.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose, True)

        self._msg_boxes.append(box)
        def _cleanup(_res: int):
            with contextlib.suppress(ValueError):
                self._msg_boxes.remove(box)
            box.deleteLater()
        box.finished.connect(_cleanup)

        # ✅ 기본 5초 자동 닫힘
        attach_autoclose(box, ms=auto_close_ms)

        box.open()

    def _post_critical(self, title: str, text: str, *, clear_status_to_idle: bool = False) -> None:
        if not self._has_ui():
            self.append_log("ERROR", f"{title}: {text}"); return

        self._ensure_msgbox_store()
        box = QMessageBox(self._parent_widget() or None)
        box.setWindowTitle(title)
        box.setText(text)
        box.setIcon(QMessageBox.Critical)
        box.setStandardButtons(QMessageBox.Ok)
        box.setWindowModality(Qt.WindowModality.WindowModal)
        box.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose, True)

        self._msg_boxes.append(box)
        def _cleanup(_res: int):
            with contextlib.suppress(ValueError):
                self._msg_boxes.remove(box)
            box.deleteLater()
        box.finished.connect(_cleanup)

        if clear_status_to_idle:
            def _ack_to_idle(_res: int):
                # OK 클릭 시만 idle로 (X로 닫으면 0인 경우가 많음)
                if int(_res) == int(QMessageBox.Ok):
                    with contextlib.suppress(Exception):
                        runtime_state.clear_error("chamber", self.ch)
            box.finished.connect(_ack_to_idle)

        box.open()

    def _has_ui(self) -> bool:
        try:
            return QApplication.instance() is not None and self._parent_widget() is not None
        except Exception:
            return False
        
    def _format_card_payload_for_chat(self, p: dict) -> dict:
        """
        구글챗 카드에 보내기 전에 보기 좋게 정리:
        - CH1: 단일 타겟 위젯(ch1_gunTarget_name) 반영, G2/G3 제거
        - 파워: 사용하지 않는 종류는 키 자체를 제거(카드에 안 보이게)
        """
        q = dict(p)

        # 기본 헤더 타이틀이 없으면 CHx Sputter로 보정
        q.setdefault("prefix", f"CH{self.ch} Sputter")

        # ── 1) CH1은 건 1개만 노출 ─────────────────────────────────────────────
        if self.ch == 1:
            # NormParams 쪽(G1_target_name/ G1 Target)과 UI 위젯(ch1_gunTarget_name) 모두 커버
            name = (q.get("G1_target_name")
                    or q.get("G1 Target")
                    or q.get("ch1_gunTarget_name")  # ← 보강: 실제 UI 필드명
                    or "").strip()
            if name:
                q["use_g1"] = True
                q["G1_target_name"] = name
            # G2/G3 관련 키 제거
            for key in ("use_g2", "use_g3",
                        "G2_target_name", "G3_target_name",
                        "G2 Target", "G3 Target"):
                q.pop(key, None)

        # ── 2) 파워는 '사용 중'인 것만 노출 ─────────────────────────────────────
        def _drop(keys: tuple[str, ...]):
            for k in keys:
                q.pop(k, None)

        if not bool(q.get("use_dc_pulse", False)):
            _drop(("dc_pulse_power", "dc_pulse_freq", "dc_pulse_duty", "dc_pulse_duty_cycle"))
        if not bool(q.get("use_rf_pulse", False)):
            _drop(("rf_pulse_power", "rf_pulse_freq", "rf_pulse_duty", "rf_pulse_duty_cycle"))
        if not bool(q.get("use_dc_power", False)):
            _drop(("dc_power",))
        if not bool(q.get("use_dc_power2", False)):
            _drop(("dc_power2",))
        if not bool(q.get("use_rf_power", False)):
            _drop(("rf_power",))

        return q
        
    # ============================= PLC 로그 소유 관리 =============================
    def set_plc_log_owner(self, owns: bool) -> None:
        """이 런타임이 PLC 로그의 현재 소유자인지 토글"""
        prev = getattr(self, "_owns_plc", False)
        self._owns_plc = bool(owns)

    def set_main_done_callback(self, callback: Callable) -> None:
        """Main Process 완료 시 main.py의 pending_log 조회용 콜백 설정."""
        self._main_done_callback = callback
    # ============================= PLC 로그 소유 관리 =============================

    # ============================= 입력값 검증 헬퍼 =============================
    def _validate_norm_params(self, p: NormParams) -> list[str]:
        errs: list[str] = []

        # 공통: 가스/유량
        if not (p.get("use_ar") or p.get("use_o2") or p.get("use_n2")):
            errs.append("가스를 하나 이상 선택해야 합니다.")

        # 🔧 None / "None" / 빈 문자열 등을 0으로 처리
        for k in ("ar_flow", "o2_flow", "n2_flow"):
            raw = p.get(k, 0)

            # None, "", "None" → 0 으로 간주
            if raw is None:
                v = 0.0
            else:
                s = str(raw).strip()
                if s == "" or s.upper() == "NONE":
                    v = 0.0
                else:
                    try:
                        v = float(s)
                    except (TypeError, ValueError):
                        # 이상한 값이면 0으로 처리하고, 에러 리스트에만 남김 (선택)
                        v = 0.0
                        errs.append(f"{k} 값이 숫자가 아니라 0으로 처리(raw={raw!r}).")

            if v < 0:
                errs.append(f"{k}는 음수 불가")

            # 이후에서 확실히 float 로 쓰도록 p에 다시 넣어줌
            p[k] = v

        if self.ch == 1:
            use_rf = bool(p.get("use_rf_pulse"))
            use_dc = bool(p.get("use_dc_pulse"))

            if not (use_rf or use_dc):
                errs.append("CH1은 RF Pulse 또는 DC Pulse 중 하나를 반드시 선택해야 합니다.")

            if use_rf:
                if p.get("rf_pulse_power", 0) <= 0:
                    errs.append("RF Pulse Target Power(W)는 0보다 커야 합니다.")
                f = p.get("rf_pulse_freq")
                d = p.get("rf_pulse_duty")
                if f is not None and not (1 <= f <= 100):
                    errs.append("RF Pulse Freq(kHz)는 1..100")
                if d is not None and not (1 <= d <= 99):
                    errs.append("RF Pulse Duty(%)는 1..99")

            if use_dc:
                if p.get("dc_pulse_power", 0) <= 0:
                    errs.append("DC Pulse Target Power(W)는 0보다 커야 합니다.")
                f = p.get("dc_pulse_freq")
                d = p.get("dc_pulse_duty")
                if f is not None and not (20 <= f <= 150):
                    errs.append("DC Pulse Freq(kHz)는 20..150")
                if d is not None and not (1 <= d <= 99):
                    errs.append("DC Pulse Duty(%)는 1..99")

        else:
            checked = int(p.get("use_g1", False)) + int(p.get("use_g2", False)) + int(p.get("use_g3", False))
            if checked == 0 or checked == 3:
                errs.append("G1~G3 중 1개 또는 2개만 선택")

            use_rf_pulse = bool(p.get("use_rf_pulse"))
            use_dc_pulse = bool(p.get("use_dc_pulse"))

            if use_rf_pulse:
                if not self.supports_rf_pulse:
                    errs.append("이 챔버는 RF-Pulse를 지원하지 않습니다.")
                else:
                    if p.get("rf_pulse_power", 0) <= 0:
                        errs.append("RF Pulse Target Power(W)는 0보다 커야 합니다.")
                    f = p.get("rf_pulse_freq")
                    d = p.get("rf_pulse_duty")
                    if f is not None and not (1 <= f <= 100):
                        errs.append("RF Pulse Freq(kHz)는 1..100")
                    if d is not None and not (1 <= d <= 99):
                        errs.append("RF Pulse Duty(%)는 1..99")

            # CH2는 (연속 DC/RF) 또는 (Pulse DC/RF) 중 하나 이상은 필요
            if not (p.get("use_dc_power") or p.get("use_dc_power2") or p.get("use_rf_power") or use_dc_pulse or use_rf_pulse):
                errs.append("CH2는 DC1/DC2/RF Power 또는 DC/RF Pulse 중 하나 이상 선택 필요")

            if p.get("use_dc_power") and p.get("dc_power", 0) < 0:
                errs.append("DC Target Power(W)는 0 이상이어야 합니다.")
            if p.get("use_dc_power2") and p.get("dc_power2", 0) < 0:
                errs.append("DC2 Target Power(W)는 0 이상이어야 합니다.")

        return errs
    # ============================= 입력값 검증 헬퍼 =============================  

    # Gun Target 자동 불러오기 (재고 DB 연동)
    async def _load_gun_targets(self) -> None:
        """
        재고 DB에서 타겟 정보 불러와 UI 입력칸에 자동 입력.
        실패 시 UI 로그에 기록 — 메인 공정에 영향 없음.
        사용자가 칸을 직접 수정하면 그 값이 공정에 사용됨.
        """
        try:
            from util.inventory_client import fetch_gun_targets
            targets = await fetch_gun_targets(self.ch)
        except Exception as e:
            self.append_log("TARGET", f"[CH{self.ch}] 타겟 정보 불러오기 실패: {e!r}")
            return

        if self.ch == 1:
            w = self._u("gunTarget_name")
            if w and not w.toPlainText().strip():
                val = targets.get("G1 Target", "")
                w.setPlainText(val)
                w.setToolTip(val)

        elif self.ch == 2:
            for gun_num in (1, 2, 3):
                w = self._u(f"g{gun_num}Target_name")
                if w and not w.toPlainText().strip():
                    val = targets.get(f"G{gun_num} Target", "")
                    w.setPlainText(val)
                    w.setToolTip(val)
