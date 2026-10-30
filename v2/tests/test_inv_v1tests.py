"""v1 시험(tests/ 93개 파일·904 함수 — 전수 목록 영역 test, v2/ch12/spec/inv_test.py)이 고정하던 기대값 가운데
v2 시험에 아직 옮겨지지 않은 것을 옮긴 곳.

규칙
- v1 을 import 하지 않는다(test_isolation). 기대값(문구·숫자)은 v1 시험·코드에서 그대로 옮기고 근거를 'v1 파일:줄' 로 적는다.
- 시험 이름 = test_<항목 번호>_<무엇>(예 test_t019_dc_pulse_overcurrent_card) — 항목은 inv_test.py 의 TEST-<번호>.
- 대역(fake)은 v2 모듈이 실제로 부르는 속성만 갖춘다. 장비·파일은 가상 시계·시뮬레이터·tmp_path 만 쓴다.
"""
from __future__ import annotations

import os
from pathlib import Path
from types import SimpleNamespace

import pytest


# ═════════════════════════ 공통 대역 ═════════════════════════
class _Eng:
    """ProcessEventsMixin 이 보는 엔진 자리 — chat·run_warnings·stats·busy."""

    def __init__(self):
        self.chats, self.run_warnings, self.busy, self.run_id, self.params = [], [], True, "r1", {}

        class _St:
            def __getattr__(self, _n):
                return lambda *a, **k: None
        self.stats = _St()

    def chat(self, text):
        self.chats.append(text)

    def ig_in_use_by_pc(self):
        return False


class _Runner:
    def __init__(self):
        self.calls = []

    def __getattr__(self, name):
        return lambda *a, **k: self.calls.append((name, a, k))


def _events_host(n=2):
    from v2.ch12.process_events import ProcessEventsMixin

    class H(ProcessEventsMixin):
        pass
    h = H()
    h.n, h.e, h.runner, h.owner = n, _Eng(), _Runner(), f"chamber{n}"
    h.logs, h.fails, h.matches = [], [], []
    h.dc_failed = {1: False, 2: False}
    h.log = lambda src, msg: h.logs.append((src, msg))
    h._fail = lambda src, why, **k: h.fails.append((src, why, k))
    h._match = lambda *a: h.matches.append(a)
    return h


def _ev(**kw):
    base = dict(cmd=None, detail=None, reason=None, data={}, code=None, owner=None)
    base.update(kw)
    return SimpleNamespace(**base)


# ═════════════ TEST-019 전원 이상 공정 중단 카드 문구(v1 tests/test_dc_overcurrent.py:411-468) ═════════════
def test_t019_dc_pulse_overcurrent_low_current_and_other_cards():
    """runtime/chamber_runtime.py:2755-2762 — AUTO_STOP 사유별 머리말 + 사유 원문, 공정 실패로 넘김(on_dc_pulse_failed)."""
    reason = "overcurrent: I=2.710A > 1.000A (설정 190W 기준, 1회 연속)"
    h = _events_host(2)
    h._ev_dcp("command_failed", _ev(cmd="AUTO_STOP", detail=reason))
    assert len(h.e.chats) == 1
    msg = h.e.chats[0]
    assert "과전류" in msg and reason in msg and "CH2" in msg
    assert msg == f"⚠️ CH2 DC Pulse 과전류 차단 → 전체 공정 중단 | {reason}"
    assert [c[0] for c in h.runner.calls] == ["step_failed"]
    h = _events_host(2)
    h._ev_dcp("command_failed", _ev(cmd="AUTO_STOP", detail="low_current: I <= 0.050A (3회 연속)"))
    assert "저전류" in h.e.chats[0] and "과전류" not in h.e.chats[0]
    h = _events_host(2)
    h._ev_dcp("command_failed", _ev(cmd="AUTO_STOP", detail="target_failed"))
    assert "세트포인트 이탈(연속) 또는 P=0W 감지" in h.e.chats[0] and "target_failed" in h.e.chats[0]


@pytest.mark.parametrize("unit", ["DC1", "DC2"])
def test_t019_dc_overcurrent_and_low_power_cards(unit):
    """chamber_runtime.py:2374-2440 — DC/DC2 target_failed: 과전류면 '과전류 차단', 저출력이면 '출력 이상'(과전류 문구 없음);
    공정 실패의 첫 인자는 원래 메시지 그대로, 실패 플래그(_dc_failed_flag) 켜짐."""
    m = "DC 과전류: meas=1.500A > 한계 1.000A (목표 300.0W 기준, 1회 연속) → 공정 중단 | I_scale=0.00097405 (raw≈1540)"
    h = _events_host(1)
    h._ev_dc("target_failed", _ev(detail=m, data={"unit": unit}))
    u = 1 if unit == "DC1" else 2
    assert h.dc_failed[u] is True and len(h.e.chats) == 1
    assert "과전류 차단" in h.e.chats[0] and m in h.e.chats[0]
    assert h.fails[0][0] == ("DC Power" if u == 1 else "DC2 Power") and h.fails[0][1] == m
    m2 = "저전력 연속 3/3 (meas=0.0W ≤ 1.0W)"
    h = _events_host(1)
    h._ev_dc("target_failed", _ev(detail=m2, data={"unit": unit}))
    assert "출력 이상" in h.e.chats[0] and "과전류" not in h.e.chats[0] and m2 in h.e.chats[0]


# ═════════════ TEST-040 next_step_message 는 읽기 전용(v1 tests/test_heater_cards.py:318) ═════════════
def test_t040_next_step_message_read_only():
    """controller/process_controller.py next_step_message — 지금 스텝 다음 것의 message, 마지막이면 ''. 위치는 그대로."""
    from v2.ch12.process_engine import ProcessEngine
    from v2.ch12.process_steps import ProcessStep
    steps = [ProcessStep(action="HEATER_WAIT", message="히터 승온 대기", value=200.0),
             ProcessStep(action="DELAY", message="다음 스텝", duration=1000)]
    fake = SimpleNamespace(runner=SimpleNamespace(steps=steps, idx=0))
    assert ProcessEngine._next_step_message(fake) == "다음 스텝" and fake.runner.idx == 0
    fake.runner.idx = 1
    assert ProcessEngine._next_step_message(fake) == ""


# ═════════════ TEST-053·167 엑셀 히터·두께 그룹 머리글(v1 tests/test_heater_process_logs.py:135, test_thickness_logs.py:74) ═════════════
def _ws_with(cols, grp, starts):
    from openpyxl import Workbook
    from v2.ch12.outputs import excel_style as XS
    wb = Workbook()
    ws = wb.active
    XS._build_headers(ws, cols, grp, starts)
    return wb, ws


def test_t053_t167_excel_heater_then_thickness_groups():
    """util/gdrive_logger.py:51-180 — CH2 = 38열 + Heater 5열(39~43, 'SP Temp…Avg DAC', 단위 °C·count) + Thickness 11열(44~54,
    그룹 머리 AR1 'Thickness', 'Material' 53열·'Factor Source' 54열). CH1 은 38열 그대로(두께 없을 때)."""
    from v2.ch12.outputs import excel_style as XS
    c1, _, _ = XS._main_cols(1)
    c2, g2, _ = XS._main_cols(2)
    c2t, g2t, s2t = XS._main_cols(2, thickness=True)
    assert len(c1) == 38 and len(c2) == 43 and len(c2t) == 54
    assert [c[0] for c in c2[38:]] == ["SP Temp", "Avg Temp", "Min Temp", "Max Temp", "Avg DAC"]
    assert [c[1] for c in c2[38:]] == ["°C", "°C", "°C", "°C", "count"]
    assert g2["Heater"] == [39, 40, 41, 42, 43] and g2t["Thickness"] == list(range(44, 55))
    assert c2t[52][0] == "Material" and c2t[53][0] == "Factor Source"
    wb, ws = _ws_with(c2t, g2t, s2t)
    assert ws["AM1"].value == "Heater" and ws["AR1"].value == "Thickness"
    assert "AM1:AQ1" in [str(r) for r in ws.merged_cells.ranges] and "AR1:BB1" in [str(r) for r in ws.merged_cells.ranges]


def test_t053_t167_old_files_get_only_new_headers():
    """gdrive_logger.py:276-318 _extend_headers — 옛 38열 CH2 파일 → 히터·두께 머리글만 덧붙임(기존 행 그대로), 히터까지(43열)
    파일 → 두께 11열, Material·Factor Source 없던 옛 Thickness 그룹(52열) → 그 2열만; CH1·머리글이 안 맞는 파일은 0."""
    from v2.ch12.outputs import excel_style as XS
    c1, g1, s1 = XS._main_cols(1)
    c2, g2, s2 = XS._main_cols(2)
    c2t, g2t, s2t = XS._main_cols(2, thickness=True)
    wb, ws = _ws_with(c1, g1, s1)                       # 옛 CH2 파일(38열) + 데이터 행 1개
    ws.cell(row=4, column=1, value="old row")
    assert XS._extend_headers(ws, c2, g2, s2) == 5
    assert [ws.cell(row=2, column=c).value for c in range(39, 44)] == ["SP Temp", "Avg Temp", "Min Temp", "Max Temp", "Avg DAC"]
    assert ws.cell(row=1, column=39).value == "Heater" and ws.cell(row=4, column=1).value == "old row"
    assert XS._extend_headers(ws, c2t, g2t, s2t) == 11 and ws.cell(row=1, column=44).value == "Thickness"
    assert ws.cell(row=4, column=44).value is None
    old_thk = c2t[:52]                                     # Material·Factor Source 없던 두께 그룹
    wb, ws = _ws_with(old_thk, *XS._group_index(old_thk))
    assert XS._extend_headers(ws, c2t, g2t, s2t) == 2
    assert ws.cell(row=2, column=53).value == "Material" and ws.cell(row=2, column=54).value == "Factor Source"
    wb, ws = _ws_with(c1, g1, s1)
    assert XS._extend_headers(ws, c1, g1, s1) == 0          # CH1 은 그대로
    ws.cell(row=2, column=5, value="손으로 고친 머리글")
    assert XS._extend_headers(ws, c2, g2, s2) == 0          # 앞부분이 다르면 손대지 않음


# ═════════════ TEST-004 Calib CSV 는 기록할 때만 열고 바로 닫음(v1 tests/test_calib_csv_release.py:36) ═════════════
def _open_paths():
    out = set()
    for fd in os.listdir("/proc/self/fd"):
        try:
            out.add(os.readlink(f"/proc/self/fd/{fd}"))
        except OSError:
            pass
    return out


@pytest.mark.skipif(not Path("/proc/self/fd").is_dir(), reason="리눅스 /proc 필요")
def test_t004_calib_csv_not_held_open(tmp_path):
    """controller/data_logger.py:156-163·472 — 시작·기록 뒤 파일 손잡이가 열려 있지 않다(G: 가 열린 파일을 올리지 않음).
    v2 는 D07(로컬 정본 먼저 → 사본 스레드가 G: 로) — 두 파일 모두 기록 뒤 닫혀 있다. 로그 '1행 기록 완료 (NAS)'."""
    from v2.ch12.outputs.base import QueueWriter
    from v2.ch12.outputs.calib_csv import CalibCsv
    from v2.ch12.outputs.paths import CopySender
    from v2.ch12.outputs.run_stats import RunStats
    from v2.tests.outputs_golden import Clock, make_root, v2_cfg
    from v2.tools import v1golden_outputs_scen as SC
    root = make_root(tmp_path)
    cfg = v2_cfg(root, SC.BASE_CFG)
    w, s, logs = QueueWriter("t"), CopySender(cfg, retry_s=0.05), []
    c = CalibCsv(1, cfg, log_func=logs.append, writer=w, sender=s)
    try:
        w.flush()
        assert not [p for p in _open_paths() if p.endswith("_log.csv")]
        for t in ("2026-10-30T11:00:00", "2026-10-30T11:10:00"):
            st = RunStats(now=Clock(t).now)
            st.apply_events([["start", dict(SC.P_CH1, t0_wall=t)]])
            c.record(st, True)
            w.flush()
            assert s.flush(2.0)
            assert not [p for p in _open_paths() if p.endswith("_log.csv") or p.endswith("_pending.csv")]
        files = list(root.rglob("Ch1_log.csv"))
        assert len(files) == 2 and all(f.read_bytes().count(b"\r\n") == 3 for f in files)
        assert any("1행 기록 완료 (NAS)" in x for x in logs)
    finally:
        w.stop()
        s.stop()
