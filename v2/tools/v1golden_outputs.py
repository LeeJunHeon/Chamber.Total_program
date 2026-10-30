"""v1 기록 출력 골든 만들기 — v1 코드를 importlib 로 실제로 돌려 같은 입력으로 만든 파일을 v2/tests/golden/outputs/ 에 저장.

실행: cd Chamber.Total_program && QT_QPA_PLATFORM=offscreen python -m v2.tools.v1golden_outputs
- v1 모듈은 문자열 이름으로 importlib 로만 부른다(v2 안에서 v1 import 금지 — v2/tests/test_isolation.py). 저장소 루트를
  sys.path 에 넣는다. PySide6·openpyxl·cv2·pymodbus 가 있어야 한다(v1 의존성).
- 시각: v1 모듈의 datetime(모듈 전역)을 고정 시계 FDT 로 바꾼다(바뀌는 칸을 맞추기 위해). 경로: lib.config_common 의 키를
  임시 폴더(한글 포함)로 바꾼다 — v1 은 호출 시점에 읽는다(DEC-033).
- 골든 = <이름>/scenario.json(입력·설정·기대 로그) + <이름>/files[_단계]/…(v1 이 만든 파일 그대로, 임시 뿌리 기준 상대 경로).
  v2 시험(v2/tests/test_outputs_*.py)은 v1 을 import 하지 않고 이 골든만 읽어 같은 입력으로 v2 를 돌리고 바이트 비교한다.
"""
from __future__ import annotations

import asyncio
import importlib
import json
import logging
import os
import shutil
import sys
import tempfile
import types
from datetime import datetime, timedelta
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
GOLDEN = REPO / "v2" / "tests" / "golden" / "outputs"
T0 = "2026-10-30T09:15:20.123000"
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from v2.tools import v1golden_outputs_scen as SC   # noqa: E402  (입력 정의 — v1 을 import 하지 않음)


def v1(name: str):
    return importlib.import_module(name)


# ───────────────────────── 고정 시계 ─────────────────────────
class Clock:
    def __init__(self, iso: str = T0) -> None:
        self.FDT = _make_fdt(self)
        self.t = self.FDT.fromisoformat(iso)
        self.mono = 1000.0

    def set(self, iso: str) -> None:
        self.t = self.FDT.fromisoformat(iso)

    def advance(self, s: float) -> None:
        d = self.t + timedelta(seconds=s)
        self.t = self.FDT(d.year, d.month, d.day, d.hour, d.minute, d.second, d.microsecond)
        self.mono += s


def _make_fdt(clock: "Clock"):
    class FDT(datetime):
        @classmethod
        def now(cls, tz=None):
            return clock.t
    return FDT


def patch_dt(mod, clock: Clock, *, as_module: bool = False) -> None:
    if as_module:
        setattr(mod, "datetime", types.SimpleNamespace(datetime=clock.FDT, date=datetime.date))
    else:
        setattr(mod, "datetime", clock.FDT)


def set_cfg(root: Path, cfg: dict) -> dict:
    """scenario cfg('@/' 로 시작하면 root 기준 경로)를 lib.config_common 에 넣는다."""
    cc = v1("lib.config_common")
    real = {}
    for k, v in cfg.items():
        if isinstance(v, str) and v.startswith("@/"):
            v = str(root / v[2:])
        setattr(cc, k, v)
        real[k] = v
    return real


def save_tree(root: Path, dest: Path, skip=()) -> list:
    out = []
    if dest.exists():
        shutil.rmtree(dest)
    for p in sorted(root.rglob("*")):
        if p.is_file() and not any(s in p.name for s in skip):
            rel = p.relative_to(root)
            (dest / rel).parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(p, dest / rel)
            out.append(rel.as_posix())
    return out


def write_scen(name: str, scen: dict) -> Path:
    d = GOLDEN / name
    d.mkdir(parents=True, exist_ok=True)
    (d / "scenario.json").write_text(json.dumps(scen, ensure_ascii=False, indent=1), encoding="utf-8")
    return d


def new_root() -> Path:
    return Path(tempfile.mkdtemp(prefix="v1gold_")) / "기록 시험"


# ───────────────────────── 시스템 로그·공정 txt ─────────────────────────
def gen_process_txt() -> None:
    for name, scen in SC.PROCESS_TXT.items():
        root = new_root()
        clock = Clock(scen["t0"])
        set_cfg(root, SC.BASE_CFG)
        cr = v1("runtime.chamber_runtime")
        sl = v1("util.system_log")
        patch_dt(cr, clock)
        patch_dt(sl, clock)
        log_hub = v1("util.log_hub")

        class RT:
            _open_run_log = cr.ChamberRuntime._open_run_log
            append_log = cr.ChamberRuntime.append_log
            append_side_log = cr.ChamberRuntime.append_side_log
            _close_run_log = cr.ChamberRuntime._close_run_log
            _log_enqueue_nowait = cr.ChamberRuntime._log_enqueue_nowait
            _log_write_sync = cr.ChamberRuntime._log_write_sync

            def __init__(self, ch):
                import collections
                self.ch = ch
                self._log_dir = root / SC.BASE_CFG["LOG_ROOT_DIR"][2:] / f"CH{ch}"
                self._log_dir.mkdir(parents=True, exist_ok=True)
                self._local_log_dir = root / SC.BASE_CFG["LOCAL_FALLBACK_ROOT"][2:] / f"CH{ch}"
                self._log_file_path = None
                self._prestart_buf = collections.deque(maxlen=1000)
                self._run_log_appender = log_hub.SessionTextAppender(fallback_dir=self._local_log_dir,
                                                                     encoding="utf-8")
                self._log_q = asyncio.Queue(maxsize=4096)
                self._log_writer_task = types.SimpleNamespace(done=lambda: False)
                self.ui_lines = []
                self._run_origin = "ui"
                self._run_origin_meta = None

            def _soon(self, fn, *a):
                fn(*a)

            def _enqueue_ui_log(self, line):
                self.ui_lines.append(line)

            def _set_task_later(self, *a, **k):
                pass

            def drain(self):          # v1 _log_writer_loop 의 묶음 쓰기(파일 경로는 쓰기 순간 값)
                parts = []
                while not self._log_q.empty():
                    parts.append(self._log_q.get_nowait())
                if parts and self._log_file_path:
                    self._log_write_sync(self._log_file_path, "".join(parts))

            def shutdown(self):
                self.drain()
                self._run_log_appender.close()
                self._log_file_path = None

        rt = RT(scen["ch"])
        for ev in scen["events"]:
            k = ev[0]
            if k == "t":
                clock.set(ev[1])
            elif k == "log":
                rt.append_log(ev[1], ev[2])
            elif k == "side":
                rt.append_side_log(ev[1], ev[2])
            elif k == "open":
                rt._run_origin = ev[2]
                rt._run_origin_meta = ev[3]
                rt._open_run_log(ev[1])
            elif k == "drain":
                rt.drain()
            elif k == "close":
                rt._close_run_log()
                rt.shutdown()
        scen["ui"] = rt.ui_lines
        d = write_scen(name, scen)
        scen["files"] = save_tree(root, d / "files")
        write_scen(name, scen)


def gen_system_log() -> None:
    root = new_root()
    clock = Clock()
    set_cfg(root, SC.BASE_CFG)
    sl = v1("util.system_log")
    patch_dt(sl, clock)
    scen = dict(SC.SYSTEM_LOG)
    for t, src, msg in scen["lines"]:
        clock.set(t)
        sl.system_log_append(src, msg)
    d = write_scen("system_log", scen)
    scen["files"] = save_tree(root, d / "files")
    write_scen("system_log", scen)


# ───────────────────────── Calib CSV·엑셀 ─────────────────────────
def _run_stats_v1(dl, events, clock):
    """RunStats 와 같은 이벤트를 v1 DataLogger 에 넣는다."""
    for ev in events:
        k, a = ev[0], ev[1:]
        if k == "t":
            clock.set(a[0])
        elif k == "start":
            dl.start_new_log_session(dict(a[0]))
        elif k == "ig":
            dl.log_ig_pressure(a[0])
        elif k == "shutter_open":
            dl.notify_shutter_open()
        elif k == "shutter_close":
            dl.notify_shutter_close()
        elif k == "dc":
            dl.log_dc_power(*a)
        elif k == "dc2":
            dl.log_dc2_power(*a)
        elif k == "dcpulse":
            dl.log_dcpulse_power(*a)
        elif k == "rf":
            dl.log_rf_power(*a)
        elif k == "rfpulse":
            dl.log_rfpulse_power(*a)
        elif k == "flow":
            dl.log_mfc_flow(*a)
        elif k == "pressure":
            dl.log_mfc_pressure(a[0])
        elif k == "heater":
            dl.log_heater(*a)
        elif k == "thickness":
            dl.set_thickness_record(a[0])
        elif k == "end":       # chamber_runtime.py:1841-1863, 1882-1886
            d = a[0]
            pp = dl.process_params
            pn = str(d.get("process_name") or "").strip()
            if pn:
                pp["process_name"] = pn
                pp["Process_name"] = pn
                pp["process_note"] = pn
            pp["result"] = d.get("result", "")
            pp["stopped"] = bool(d.get("stopped", False))
            if d.get("actual_shutter_delay_min") is not None:
                pp["shutter_delay"] = float(d["actual_shutter_delay_min"])
            if d.get("actual_process_time_min") is not None:
                pp["process_time"] = float(d["actual_process_time_min"])
            if d.get("soft_arc_count") is not None:
                pp["soft_arc_count"] = d["soft_arc_count"]
            if d.get("hard_arc_count") is not None:
                pp["hard_arc_count"] = d["hard_arc_count"]


def gen_calib() -> None:
    dlm = v1("controller.data_logger")
    for name, scen in SC.CALIB.items():
        root = new_root()
        clock = Clock(scen.get("t0", T0))
        cfg = set_cfg(root, dict(SC.BASE_CFG, **scen.get("cfg", {})))
        patch_dt(dlm, clock)
        cc = v1("lib.config_common")
        setattr(cc, f"LOCAL_FALLBACK_CH{scen['ch']}_DIR", str(Path(cfg["LOCAL_FALLBACK_ROOT"]) / f"CH{scen['ch']}"))
        calib_dir = Path(cfg["SPUTTER_CALIB_DB_DIR"])
        for rel, text in scen.get("pre_files", {}).items():
            p = root / rel
            p.parent.mkdir(parents=True, exist_ok=True)
            if text is None:
                p.mkdir(parents=True, exist_ok=True)
            else:
                p.write_bytes(text.encode("utf-8"))
        logs = []
        dl = dlm.DataLogger(ch=scen["ch"], csv_dir=calib_dir, log_func=logs.append,
                            thickness_columns=scen.get("thickness", False))
        d = write_scen(name, scen)
        for i, rec in enumerate(scen["records"]):
            for rel in rec.get("unblock", []):
                (root / rel).unlink()
                (root / rel).mkdir(parents=True)
            _run_stats_v1(dl, rec["events"], clock)
            dl.finalize_and_write_log(rec["ok"])
            if rec.get("stage"):
                save_tree(root, d / f"files_{rec['stage']}")
        scen["logs"] = logs
        scen["files"] = save_tree(root, d / "files")
        write_scen(name, scen)


def _xlsx_rows_v1(gl, ch, dl, rec):
    return gl._build_data(ch, dl, rec.get("operator", ""), rec.get("substrate", ""), rec.get("note", ""), None)


def gen_excel() -> None:
    gl = v1("util.gdrive_logger")
    dlm = v1("controller.data_logger")
    for name, scen in SC.EXCEL.items():
        root = new_root()
        clock = Clock(scen.get("t0", T0))
        cfg = set_cfg(root, dict(SC.BASE_CFG, **scen.get("cfg", {})))
        patch_dt(dlm, clock)
        patch_dt(gl, clock)
        main_dir = root / scen["main_dir"]
        fb = Path(cfg["LOCAL_FALLBACK_ROOT"])
        d = write_scen(name, scen)
        if scen.get("pre") == "phantom":       # 서식만 있는 행(max_row 함정)이 있는 기존 파일
            cols, gc, gs = gl._main_cols(scen["ch"], scen.get("thickness", False))
            wb, ws = gl._get_or_create_wb(main_dir / f"CH{scen['ch']}.xlsx", "Main Process", cols, gc, gs)
            from openpyxl.styles import PatternFill
            for r in range(4, 30):
                ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="F2F3F4")
            main_dir.mkdir(parents=True, exist_ok=True)
            wb.save(str(main_dir / f"CH{scen['ch']}.xlsx"))
        elif scen.get("pre") == "base_only":   # 히터 열이 없던 옛 CH2 파일 → 새 열 머리글만 채움
            wb, ws = gl._get_or_create_wb(main_dir / f"CH{scen['ch']}.xlsx", "Main Process", gl._COLS_MAIN,
                                          gl._GRP_COLS_MAIN, gl._GRP_STARTS_MAIN)
            main_dir.mkdir(parents=True, exist_ok=True)
            wb.save(str(main_dir / f"CH{scen['ch']}.xlsx"))
        elif scen.get("pre") == "broken":
            main_dir.mkdir(parents=True, exist_ok=True)
            (main_dir / f"CH{scen['ch']}.xlsx").write_bytes(b"not a zip file")
        elif scen.get("pre") == "blocked":     # G: 폴더 자리에 파일 → 저장 실패 → pending.xlsx
            main_dir.parent.mkdir(parents=True, exist_ok=True)
            main_dir.write_bytes(b"blocker")
        if scen.get("pre") in ("phantom", "base_only", "broken"):
            save_tree(root, d / "files_pre")
        for rec in scen["records"]:
            clock.set(rec.get("t", scen.get("t0", T0)))
            if rec.get("pc"):
                gl.save_pc_only(scen["ch"], rec["pc"], main_dir)
                continue
            dl = dlm.DataLogger(ch=scen["ch"], csv_dir=root / "calib_unused", log_func=None,
                                thickness_columns=scen.get("thickness", False))
            _run_stats_v1(dl, rec["events"], clock)
            rows = _xlsx_rows_v1(gl, scen["ch"], dl, rec)    # 날짜 = 세션 시작(Calib 기록 전에 — LOG-058)
            cols = gl._main_cols(scen["ch"], scen.get("thickness", False))
            gl._save_sync(scen["ch"], rows, main_dir, int(cfg.get("GDRIVE_ARC_ALERT_THRESH", 5000)), 20.0,
                          False, "", fb, "Main Process", *cols)
        shutil.rmtree(root / "calib_unused", ignore_errors=True)
        scen["files"] = save_tree(root, d / "files")
        write_scen(name, scen)


# ───────────────────────── PLC 코일 CSV ─────────────────────────
def gen_plc_coil() -> None:
    plc = v1("device.plc")
    scen = dict(SC.PLC_COIL)
    root = new_root()
    clock = Clock(scen["t0"])
    cfg = set_cfg(root, dict(SC.BASE_CFG, **scen["cfg"]))
    patch_dt(plc, clock)
    keys = list(plc.PLC_COIL_MAP.keys())
    steps = scen["steps"]
    logs = []

    class FakeTime:
        @staticmethod
        def perf_counter():
            return clock.mono

        @staticmethod
        def time():
            return clock.mono

    state = {"i": 0}

    async def fake_sleep(s):
        if s > 0:
            clock.advance(s)
            state["i"] += 1
            if state["i"] >= len(steps):
                fake._plc_coil_log_stop.set()

    real_asyncio = plc.asyncio
    plc.asyncio = types.SimpleNamespace(sleep=fake_sleep, Event=asyncio.Event, to_thread=real_asyncio.to_thread,
                                        create_task=real_asyncio.create_task, CancelledError=asyncio.CancelledError)
    plc.time = FakeTime

    class Fake:
        _coil_plan_idx = 0
        _consec_timeouts = 0
        _consec_timeouts_low = 0

        def is_connected(self):
            return steps[state["i"]]["kind"] != "disconnected"

        def _note_coil_plan_result(self, *a):
            pass

        def _coil_ranges_for_plan(self, addrs, idx):
            return [], [1, 2]

        def _snapshot_budget_s(self):
            return 4.0

        def log(self, fmt, *args):
            logs.append(fmt % args if args else fmt)

        async def snapshot_all_coils_fast(self, keys):
            st = steps[state["i"]]
            if st["kind"] == "empty":
                return {}
            on = set(st.get("on", []))
            return {k: (k in on) for k in keys}

        async def snapshot_regs_fast(self, keys):
            st = steps[state["i"]]
            if st.get("reg_fail"):
                raise TimeoutError("reg timeout")
            return {k: v for k, v in st.get("regs", {}).items()}

    fake = Fake()
    Fake._plc_coil_log_loop = plc.AsyncPLC._plc_coil_log_loop if hasattr(plc, "AsyncPLC") else None
    cls = next(c for c in vars(plc).values() if isinstance(c, type) and hasattr(c, "_plc_coil_log_loop"))
    Fake._plc_coil_log_loop = cls._plc_coil_log_loop

    def _backoff_hook():
        pass

    async def run():
        fake._plc_coil_log_stop = asyncio.Event()
        fake._plc_coil_log_interval = max(0.5, float(cfg["PLC_COIL_LOG_INTERVAL_S"]))
        fake._plc_coil_log_keys = keys
        fake._plc_reg_log_keys = list(scen["reg_keys"])
        fake._plc_coil_log_nas_dir = Path(cfg["PLC_COIL_LOG_NAS_DIR"])
        fake._plc_coil_log_local_dir = Path(cfg["PLC_COIL_LOG_LOCAL_DIR"])
        lh = v1("util.log_hub")
        fake._plc_coil_csv_writer = lh.DailyCsvListAppender(
            primary_dir=fake._plc_coil_log_nas_dir, fallback_dir=fake._plc_coil_log_local_dir,
            filename_builder=lambda dt: f"{dt.strftime('%Y%m%d')}.csv", encoding="utf-8-sig",
            retry_primary_every_s=10.0)

        orig_conn = Fake.is_connected

        def is_connected(self):
            st = steps[state["i"]]
            self._consec_timeouts = 1 if st["kind"] == "backoff" else 0
            return orig_conn(self)
        Fake.is_connected = is_connected
        await fake._plc_coil_log_loop()
        fake._plc_coil_csv_writer.close()

    asyncio.run(run())
    plc.asyncio = real_asyncio
    import time as _t
    plc.time = _t
    scen["coil_keys"] = keys
    scen["logs"] = logs
    d = write_scen("plc_coil", scen)
    scen["files"] = save_tree(root, d / "files")
    write_scen("plc_coil", scen)


# ───────────────────────── 두께·히터 ─────────────────────────
def gen_thickness_heater() -> None:
    tl = v1("lib.thickness_logger")
    hl = v1("lib.heater_logger")
    scen = dict(SC.THICKNESS)
    root = new_root()
    clock = Clock(scen["t0"])
    cfg = set_cfg(root, dict(SC.BASE_CFG, **scen["cfg"]))
    warns = []
    lg = tl.ThicknessCsvLogger(cfg["THICKNESS_LOG_DIR"], scen["ch"], warn=warns.append, clock=lambda: clock.mono)
    for ev in scen["events"]:
        if ev[0] == "t":
            clock.set(ev[1])
        elif ev[0] == "row":
            lg.write_row(dict(ev[1], process_running=True, process_name=scen["process_name"]), wall=clock.t)
        elif ev[0] == "close":
            lg.close()
    lg.close()
    scen["warns"] = warns
    d = write_scen("thickness", scen)
    scen["files"] = save_tree(root, d / "files")
    write_scen("thickness", scen)

    scen = dict(SC.HEATER)
    root = new_root()
    clock = Clock(scen["t0"])
    cfg = set_cfg(root, dict(SC.BASE_CFG, **scen["cfg"]))
    patch_dt(hl, clock, as_module=True)
    h = hl.HeaterCsvLogger("", clock=lambda: clock.mono)
    for ev in scen["events"]:
        if ev[0] == "t":
            clock.set(ev[1])
            clock.mono = ev[2]
        elif ev[0] == "start":
            h.start()
        elif ev[0] == "row":
            h.write_row(ev[1], **ev[2])
        elif ev[0] == "stop":
            h.stop()
    d = write_scen("heater", scen)
    scen["files"] = save_tree(root, d / "files")
    write_scen("heater", scen)


# ───────────────────────── 카메라·selfwatch·ERROR·덤프 ─────────────────────────
def gen_camera() -> None:
    cam = v1("util.camera_recorder")
    import numpy as np
    scen = dict(SC.CAMERA)
    root = new_root()
    clock = Clock(scen["t0"])
    set_cfg(root, dict(SC.BASE_CFG, **scen["cfg"]))
    patch_dt(cam, clock)
    n_frames = scen["frames"]
    logs = []
    rec = cam.CameraRecorder(camera_index=1, interval=1.0)

    class FakeCap:
        def __init__(self, idx):
            self.n = 0

        def isOpened(self):
            return True

        def set(self, *a):
            return True

        def read(self):
            self.n += 1
            if self.n > n_frames:
                rec._stop_event.set()
                return False, None
            img = np.zeros((6, 8, 3), dtype=np.uint8)
            img[:, :, 0] = self.n * 20
            img[2:4, 1:5, 1] = 200
            return True, img

        def release(self):
            pass

    class FakeTime:
        @staticmethod
        def time():
            return clock.mono

        @staticmethod
        def sleep(s):
            clock.advance(s)

    real_vc = cam.cv2.VideoCapture
    cam.cv2.VideoCapture = FakeCap
    cam.time = FakeTime
    rec._mode = scen["mode"]
    rec._mode_folder = scen["mode"]
    rec._log = logs.append
    rec._free_override = None
    rec._record_loop()
    cam.cv2.VideoCapture = real_vc
    import time as _t
    cam.time = _t
    scen["logs"] = logs
    d = write_scen("camera", scen)
    scen["files"] = save_tree(root, d / "files")
    write_scen("camera", scen)


def gen_selfwatch() -> None:
    sw = v1("util.self_watch")
    scen = dict(SC.SELFWATCH)
    root = new_root()
    clock = Clock(scen["t0"])
    set_cfg(root, dict(SC.BASE_CFG, **scen["cfg"]))
    patch_dt(sw, clock)
    w = sw.SelfWatch(None)
    reasons_all = []
    for t, sample in scen["samples"]:
        clock.set(t)
        s = {k: -1 for k in sw._CSV_FIELDS}
        s.update(sample)
        s["ts"] = clock.t.strftime("%Y-%m-%d %H:%M:%S")
        w._append_csv(dict(s))
        r = w.check_thresholds(s)
        reasons_all.append(r)
        if r:
            sw.write_dump(r, s, None)
    if w._io_exec is not None:
        w._io_exec.shutdown(wait=True)
    scen["reasons"] = reasons_all
    d = write_scen("selfwatch", scen)
    scen["files"] = save_tree(root, d / "files")
    write_scen("selfwatch", scen)


def _records(scen):
    out = []
    for r in scen["records"]:
        rec = logging.LogRecord(scen["logger"], getattr(logging, r["level"]), "x.py", 1, r["msg"], tuple(r["args"]),
                                (ValueError, ValueError(r["exc"]), None) if r.get("exc") else None)
        ts = datetime.fromisoformat(r["t"]).timestamp()
        rec.created = ts
        rec.msecs = int((ts - int(ts)) * 1000) + 0.0
        rec.process = 4242
        rec.thread = 777
        out.append(rec)
    return out


def gen_app_log() -> None:
    al = v1("util.app_logging")
    scen = dict(SC.APP_LOG)
    root = new_root()
    clock = Clock(scen["t0"])
    cfg = set_cfg(root, dict(SC.BASE_CFG, **scen["cfg"]))
    patch_dt(al, clock)
    h = al._DailyFileHandler(scen["logger"], Path(cfg["LOG_ROOT_DIR"]) / "ERROR")
    h.setFormatter(logging.Formatter(fmt="%(asctime)s.%(msecs)03d [%(levelname)s] [%(name)s] "
                                     "[pid=%(process)d tid=%(thread)d] %(message)s", datefmt="%Y-%m-%d %H:%M:%S"))
    for rec in _records(scen):
        h.emit(rec)
    h.close()
    scen["levels"] = [[s, m, al.classify_log_level(s, m)] for s, m in scen["classify"]]
    d = write_scen("app_log", scen)
    scen["files"] = save_tree(root, d / "files")
    write_scen("app_log", scen)


def gen_runtime_dump() -> None:
    rd = v1("util.runtime_dump")
    scen = dict(SC.RUNTIME_DUMP)
    root = new_root()
    set_cfg(root, dict(SC.BASE_CFG))
    p = rd.request_dump(ui=types.SimpleNamespace(), loop=None, log_root=root / "로그", extra_objects=scen["state"],
                        reason=scen["reason"])
    import time as _t
    for _ in range(200):
        if p.exists() and p.read_text(encoding="utf-8").endswith("[end]\n"):
            break
        _t.sleep(0.02)
    text = p.read_text(encoding="utf-8")
    snap = json.loads(text.split("=== SNAPSHOT(JSON) ===\n", 1)[1].split("\n\n=== THREAD TRACEBACK", 1)[0])
    scen["v1_name_pattern"] = p.name.replace(str(os.getpid()), "<pid>")
    scen["v1_snapshot_keys"] = list(snap.keys())
    scen["v1_meta_keys"] = list(snap["meta"].keys())
    scen["v1_ui"] = snap["ui"]
    scen["v1_asyncio"] = snap["asyncio"]
    scen["v1_objects"] = snap["objects"]
    scen["v1_frame"] = [text.split("\n", 1)[0], "=== THREAD TRACEBACK (faulthandler) ===" in text, text.endswith("\n[end]\n")]
    write_scen("runtime_dump", scen)


# ───────────────────────── Robot CSV ─────────────────────────
def gen_robot() -> None:
    hp = v1("util.host_process_log")
    for name, scen in SC.ROBOT.items():
        root = new_root()
        clock = Clock(scen["t0"])
        set_cfg(root, dict(SC.BASE_CFG, **scen.get("cfg", {})))
        cc = v1("lib.config_common")
        for k in scen.get("del_cfg", []):
            setattr(cc, k, "not-a-number")
        patch_dt(hp, clock)

        class FakeTime:
            @staticmethod
            def time():
                return clock.t.timestamp()

            @staticmethod
            def sleep(s):
                clock.advance(s)
        hp.time = FakeTime
        d = write_scen(name, scen)
        for rel, spec in scen.get("pre_files", {}).items():
            p = root / rel
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_bytes(spec["text"].encode("utf-8"))
            if "age_s" in spec:
                ts = clock.t.timestamp() - spec["age_s"]
                os.utime(p, (ts, ts))
        inst = {"h": hp.HostProcessLog()}
        inst["h"]._ensure_worker = lambda: None
        keys = {}
        for ev in scen["events"]:
            k, a = ev[0], ev[1:]
            h = inst["h"]
            if k == "t":
                clock.set(a[0])
            elif k == "request":
                kw = dict(a[1])
                if kw.get("received_at"):
                    kw["received_at"] = clock.FDT.fromisoformat(kw["received_at"])
                keys[a[0]] = h.request(**kw)
            elif k == "update":
                h.update(keys[a[0]], **a[1])
            elif k == "started":
                h.mark_started(keys[a[0]], clock.t, a[1])
            elif k == "finalize":
                keys[a[0] + "_ret"] = h.finalize(keys[a[0]], a[1], a[2], clock.t)
            elif k == "reject":
                h.reject(keys[a[0]], a[1])
            elif k == "flush":
                h._flush_once()
            elif k == "restart":
                inst["h"] = hp.HostProcessLog()
                inst["h"]._ensure_worker = lambda: None
            elif k == "recover":
                h.startup_recover()
            elif k == "unblock":
                p = root / a[0]
                p.unlink()
                p.mkdir(parents=True)
            elif k == "stage":
                save_tree(root, d / f"files_{a[0]}")
        import time as _t
        hp.time = _t
        scen["keys"] = {k: v for k, v in keys.items()}
        scen["files"] = save_tree(root, d / "files")
        write_scen(name, scen)


def main() -> None:
    GOLDEN.mkdir(parents=True, exist_ok=True)
    gen_system_log()
    gen_process_txt()
    gen_calib()
    gen_excel()
    gen_plc_coil()
    gen_thickness_heater()
    gen_camera()
    gen_selfwatch()
    gen_app_log()
    gen_runtime_dump()
    gen_robot()
    print("골든 작성:", GOLDEN)


if __name__ == "__main__":
    main()
