"""CH1/CH2 공정 엑셀의 열 정의·서식 — v1 util/gdrive_logger.py:36-460 을 그대로 옮긴 것(스타일 호출 순서까지 같게).

열 정의(_COLS_MAIN·_COLS_PC·_COLS_HEATER·_COLS_THICKNESS), 머리글 3행(_build_headers), 새 뒤쪽 열 머리글(_extend_headers),
파일 열기/새로 만들기·broken 처리(_get_or_create_wb), 다음 행 찾기(_next_data_row — max_row 함정), 이상치 색(_cell_bg),
데이터 행(_write_data_row — 값 없는 행은 서식도 안 씀). 쓰는 곳: v2/ch12/outputs/excel_log.py.
"""
from __future__ import annotations

import contextlib
import threading
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_SP_DIFF_RATIO = 0.10   # SP 대비 Avg 편차 강조 기준(util/gdrive_logger.py:34)

# ── 색상 ────────────────────────────────────────────────────────
_T = {
    "기본 정보"       : ("1C2833", "2C3E50", "ABB2B9"),
    "Plasma Cleaning" : ("2E4057", "4A6278", "BDC3C7"),
    "Main Process"    : ("2E4057", "4A6278", "BDC3C7"),  # ← 엑셀 포맷과 동일하게 변경
    "Heater"          : ("4E342E", "6D4C41", "D7CCC8"),  # CH2 히터(맨 뒤 그룹)
    "Thickness"       : ("004D40", "00695C", "B2DFDB"),  # 두께 모니터(히터 그룹 다음, SUPPORTS_THICKNESS 채널만)
}
_WHITE    = "FFFFFF"
_ROW_ODD  = "FFFFFF"
_ROW_EVEN = "F2F3F4"
_WARN_ARC = "FADBD8"   # 이상치 — 빨강으로 통일
_FG       = "1C2833"

# ── Main Process 시트 컬럼 정의 ─────────────────────────────────
_COLS_MAIN: List[Tuple] = [
    # 기본 정보 (13개, A~M)
    ("날짜",           "YYYY-MM-DD HH:MM:SS", 19, "기본 정보",    "timestamp"),
    ("담당자",         "",              8,  "기본 정보",    "operator"),
    ("Process Name",   "",             16,  "기본 정보",    "process_name"),
    ("비고",           "",             18,  "기본 정보",    "note"),
    ("기판",           "",             12,  "기본 정보",    "substrate"),
    ("Main Shutter",   "T/F",           9,  "기본 정보",    "main_shutter"),
    ("Power Select",   "T/F",           9,  "기본 정보",    "power_select"),
    ("G1 Target",      "",              9,  "기본 정보",    "G1 Target"),
    ("G2 Target",      "",              9,  "기본 정보",    "G2 Target"),
    ("G3 Target",      "",              9,  "기본 정보",    "G3 Target"),
    ("Dep.rate",       "nm/s",         10,  "기본 정보",    "dep_rate"),
    ("Thickness",      "nm",            9,  "기본 정보",    "thickness"),
    ("Chuck",          "up/mid/down",  11,  "기본 정보",    "chuck_position"),
    # Main Process (25개, N~AL)
    ("Shutter Delay",  "min",           8,  "Main Process", "shutter_delay"),
    ("Process Time",   "min",           8,  "Main Process", "process_time"),
    ("Base Pressure",  "Torr",         12,  "Main Process", "base_pressure"),
    ("SP Ar",          "sccm",          8,  "Main Process", "sp_ar"),
    ("Avg Ar",         "sccm",          8,  "Main Process", "avg_ar"),
    ("SP N2",          "sccm",          8,  "Main Process", "sp_n2"),
    ("Avg N2",         "sccm",          8,  "Main Process", "avg_n2"),
    ("SP O2",          "sccm",          8,  "Main Process", "sp_o2"),
    ("Avg O2",         "sccm",          8,  "Main Process", "avg_o2"),
    ("SP Pressure",    "mTorr",         9,  "Main Process", "sp_pressure"),
    ("Avg Pressure",   "mTorr",         9,  "Main Process", "avg_pressure"),
    ("Power Source",   "DC/RF/DCPulse/RFPulse", 19, "Main Process", "power_source"),
    ("SP Power",       "W",             8,  "Main Process", "sp_power"),
    ("Avg Power",      "W",             9,  "Main Process", "avg_power"),
    ("Avg For.p",      "W",             9,  "Main Process", "avg_forp"),
    ("Avg Ref.p",      "W",             9,  "Main Process", "avg_refp"),
    ("Avg Load",       "a.u.",          8,  "Main Process", "avg_load"),
    ("Avg Tune",       "a.u.",          8,  "Main Process", "avg_tune"),
    ("Avg Voltage",    "V",             9,  "Main Process", "avg_voltage"),
    ("Avg Current",    "A",             9,  "Main Process", "avg_current"),
    ("Duty Cycle",     "%",             8,  "Main Process", "duty_cycle"),
    ("Frequency",      "kHz",           8,  "Main Process", "frequency"),
    ("Off Time",       "μs",            8,  "Main Process", "off_time"),
    ("Soft Arc",       "count",         8,  "Main Process", "soft_arc"),
    ("Hard Arc",       "count",         8,  "Main Process", "hard_arc"),
]

# ── Plasma Cleaning 시트 컬럼 정의 ──────────────────────────────
_COLS_PC: List[Tuple] = [
    # 기본 정보 (5개 — Main Shutter/Power Select/Gun/Chuck 없음)
    ("날짜",           "YYYY-MM-DD HH:MM:SS", 19, "기본 정보",       "timestamp"),
    ("담당자",         "",              8,  "기본 정보",       "operator"),
    ("Process Name",   "",             16,  "기본 정보",       "process_name"),
    ("비고",           "",             18,  "기본 정보",       "note"),
    ("기판",           "",             12,  "기본 정보",       "substrate"),
    # Plasma Cleaning (11개)
    ("Time",           "min",           7,  "Plasma Cleaning", "pc_time"),
    ("Base Pressure",  "Torr",         12,  "Plasma Cleaning", "pc_base_pressure"),
    ("SP Ar",          "sccm",          8,  "Plasma Cleaning", "pc_sp_ar"),
    ("Avg Ar",         "sccm",          8,  "Plasma Cleaning", "pc_avg_ar"),
    ("SP Pressure",    "mTorr",         9,  "Plasma Cleaning", "pc_sp_pressure"),
    ("Avg Pressure",   "mTorr",         9,  "Plasma Cleaning", "pc_avg_pressure"),
    ("SP Power",       "W",             8,  "Plasma Cleaning", "pc_sp_power"),
    ("Avg For.p",      "W",             9,  "Plasma Cleaning", "pc_avg_forp"),
    ("Avg Ref.p",      "W",             9,  "Plasma Cleaning", "pc_avg_refp"),
    ("Avg Load",       "a.u.",          8,  "Plasma Cleaning", "pc_avg_load"),
    ("Avg Tune",       "a.u.",          8,  "Plasma Cleaning", "pc_avg_tune"),
]

# ── CH2 히터 열 (Main Process 시트 맨 뒤, CH2 파일에만) ───────────
#   값은 메인 공정 구간(폴링 중) 표본 — DataLogger.heater_stats(). 히터를 안 쓴 공정은 공란.
#   엑셀은 열 번호로 쓰므로 반드시 맨 뒤에 붙인다(중간에 넣으면 기존 행과 어긋남).
_COLS_HEATER: List[Tuple] = [
    ("SP Temp",        "°C",            8,  "Heater",       "heater_sv"),
    ("Avg Temp",       "°C",            8,  "Heater",       "heater_pv_avg"),
    ("Min Temp",       "°C",            8,  "Heater",       "heater_pv_min"),
    ("Max Temp",       "°C",            8,  "Heater",       "heater_pv_max"),
    ("Avg DAC",        "count",         8,  "Heater",       "heater_mv_avg"),
]
_COLS_MAIN_CH2: List[Tuple] = _COLS_MAIN + _COLS_HEATER

# ── 두께 모니터 열 (Main Process 시트 맨 뒤 — 히터 그룹 다음, SUPPORTS_THICKNESS=true 채널 파일에만) ──
#   값은 DataLogger.thickness_record(공정 종료 때 런타임이 넣음). 열 번호로 쓰므로 반드시 맨 뒤에 붙인다.
#   "Dep.rate"/"Thickness"(기본 정보)는 레시피 값, "Process Time" 은 실제 진행 시간 그대로.
_COLS_THICKNESS: List[Tuple] = [
    ("End Mode",       "",              9,  "Thickness",    "thk_end"),
    ("Target",         "nm",            8,  "Thickness",    "thk_target"),
    ("Final",          "nm",            8,  "Thickness",    "thk_final"),
    ("Avg Rate",       "nm/s",          9,  "Thickness",    "thk_avg_rate"),
    ("Tooling",        "%",             8,  "Thickness",    "thk_tooling"),
    ("Density",        "g/cm3",         8,  "Thickness",    "thk_density"),
    ("Z",              "",              7,  "Thickness",    "thk_z"),
    ("Fallback Reason", "",            18,  "Thickness",    "thk_fallback_reason"),
    ("Jump Count",     "회",            8,  "Thickness",    "thk_jump_count"),
    ("Material",       "",             16,  "Thickness",    "thk_material"),
    ("Factor Source",  "",             20,  "Thickness",    "thk_factor_source"),
]


def _group_index(cols: List[Tuple]) -> Tuple[Dict[str, List[int]], set]:
    """열 정의 → (그룹별 열 번호, 그룹 시작 열 번호 집합)."""
    grp_cols: Dict[str, List[int]] = {}
    for ci, (*_, grp, _key) in enumerate(cols, 1):
        grp_cols.setdefault(grp, []).append(ci)
    return grp_cols, {c[0] for c in grp_cols.values()}


# 시트별 그룹 인덱스 계산
_GRP_COLS_MAIN: Dict[str, List[int]] = {}
for _ci, (*_, _grp, _key) in enumerate(_COLS_MAIN, 1):
    _GRP_COLS_MAIN.setdefault(_grp, []).append(_ci)
_GRP_STARTS_MAIN = {cols[0] for cols in _GRP_COLS_MAIN.values()}
_GRP_COLS_MAIN_CH2, _GRP_STARTS_MAIN_CH2 = _group_index(_COLS_MAIN_CH2)


def _main_cols(ch: int, thickness: bool = False) -> Tuple[List[Tuple], Dict[str, List[int]], set]:
    """Main Process 시트 열 정의 — CH2 는 히터 열 포함, CH1 은 그대로. thickness(SUPPORTS_THICKNESS 채널)면 맨 뒤에 두께 열."""
    if int(ch) == 2:
        base = (_COLS_MAIN_CH2, _GRP_COLS_MAIN_CH2, _GRP_STARTS_MAIN_CH2)
    else:
        base = (_COLS_MAIN, _GRP_COLS_MAIN, _GRP_STARTS_MAIN)
    if not thickness:
        return base
    cols = base[0] + _COLS_THICKNESS
    grp_cols, grp_starts = _group_index(cols)
    return cols, grp_cols, grp_starts


_GRP_COLS_PC: Dict[str, List[int]] = {}
for _ci, (*_, _grp, _key) in enumerate(_COLS_PC, 1):
    _GRP_COLS_PC.setdefault(_grp, []).append(_ci)
_GRP_STARTS_PC = {cols[0] for cols in _GRP_COLS_PC.values()}

# ── 내부 락 ──────────────────────────────────────────────────────
_write_lock = threading.Lock()


# ════════════════════════════════════════════════════════════════
# 스타일 헬퍼
# ════════════════════════════════════════════════════════════════

def _dcp_off_time_value(pp: dict):
    """DC Pulse Off Time [µs] 값. 실측(dc_pulse_off_time_us) 우선, 없으면 레시피 값.
    "DC" 모드나 미지정은 None."""
    v = pp.get("dc_pulse_off_time_us", None)
    if v not in (None, ""):
        try:
            return float(v)
        except Exception:
            pass
    raw = pp.get("dc_pulse_off_time", None)
    if raw in (None, "") or isinstance(raw, str):
        return None
    try:
        return float(raw)
    except Exception:
        return None


def _cs(ws, r: int, c: int, val=None, bg: Optional[str] = None,
        fg: str = _FG, bold: bool = False, size: int = 9,
        ha: str = "center", wrap: bool = False,
        nf: Optional[str] = None, italic: bool = False):
    cell = ws.cell(row=r, column=c)
    cell.value = val
    cell.font = Font(name="Arial", bold=bold, size=size,
                     color=fg, italic=italic)
    cell.alignment = Alignment(horizontal=ha, vertical="center",
                                wrap_text=wrap)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    if nf:
        cell.number_format = nf
    return cell


def _build_headers(ws, cols: List[Tuple], grp_cols: Dict, grp_starts: set) -> None:
    """Row 1 그룹 / Row 2 컬럼명 / Row 3 단위 헤더 생성."""
    thin  = Side(style="thin",   color="E0E0E0")
    thick = Side(style="medium", color="FFFFFF")

    for grp, gcols in grp_cols.items():
        c1, c2 = gcols[0], gcols[-1]
        bg = _T[grp][0]
        if c1 < c2:
            ws.merge_cells(
                f"{get_column_letter(c1)}1:{get_column_letter(c2)}1")
        cell = ws.cell(row=1, column=c1)
        cell.value = grp
        cell.font = Font(name="Arial", bold=True, size=9, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in gcols[1:]:
            with contextlib.suppress(AttributeError):
                ws.cell(row=1, column=c).fill = \
                    PatternFill("solid", fgColor=bg)
    ws.row_dimensions[1].height = 16

    for ci, (name, unit, width, grp, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(ci)].width = width
        _cs(ws, 2, ci, name, bg=_T[grp][1], fg=_WHITE,
            bold=True, size=9, wrap=True)
        _cs(ws, 3, ci, unit or "", bg=_T[grp][2], fg="4A4A4A",
            size=8, italic=True)
    ws.row_dimensions[2].height = 32
    ws.row_dimensions[3].height = 13

    for r in range(1, 4):
        for ci in range(1, len(cols) + 1):
            left = thick if ci in grp_starts else thin
            ws.cell(row=r, column=ci).border = Border(
                left=left, right=thin, top=thin, bottom=thin)


def _extend_headers(ws, cols: List[Tuple], grp_cols: Dict, grp_starts: set) -> int:
    """기존 시트에 새로 생긴 뒤쪽 열의 머리글(1~3행)을 채운다. 채운 열 수를 돌려준다.
    2행 머리글이 cols 의 앞부분과 똑같고 그 뒤가 비어 있을 때만 — 앞부분이 다르면(옛 배치·손으로 고친 파일)
    새 열을 엉뚱한 자리에 만들지 않도록 건드리지 않는다. 데이터 행은 손대지 않는다."""
    names = [c[0] for c in cols]
    have: List[str] = []
    for ci in range(1, len(names) + 1):
        v = ws.cell(row=2, column=ci).value
        if v is None or str(v).strip() == "":
            break
        have.append(str(v).strip())
    n = len(have)
    if n == 0 or n >= len(names) or have != names[:n]:
        return 0
    if any(ws.cell(row=2, column=ci).value not in (None, "") for ci in range(n + 1, len(names) + 1)):
        return 0

    thin = Side(style="thin", color="E0E0E0")
    thick = Side(style="medium", color="FFFFFF")
    new_cols = range(n + 1, len(names) + 1)
    for grp, gcols in grp_cols.items():
        part = [c for c in gcols if c > n]
        if not part:
            continue
        bg = _T[grp][0]
        if part == gcols and len(gcols) > 1:          # 새 열로만 된 그룹 → 1행 병합
            ws.merge_cells(f"{get_column_letter(gcols[0])}1:{get_column_letter(gcols[-1])}1")
        if part[0] == gcols[0]:
            cell = ws.cell(row=1, column=gcols[0])
            cell.value = grp
            cell.font = Font(name="Arial", bold=True, size=9, color=_WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in part:
            with contextlib.suppress(AttributeError):
                ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor=bg)
    for ci in new_cols:
        name, unit, width, grp, _ = cols[ci - 1]
        ws.column_dimensions[get_column_letter(ci)].width = width
        _cs(ws, 2, ci, name, bg=_T[grp][1], fg=_WHITE, bold=True, size=9, wrap=True)
        _cs(ws, 3, ci, unit or "", bg=_T[grp][2], fg="4A4A4A", size=8, italic=True)
        for r in range(1, 4):
            left = thick if ci in grp_starts else thin
            with contextlib.suppress(AttributeError):
                ws.cell(row=r, column=ci).border = Border(left=left, right=thin, top=thin, bottom=thin)
    return len(new_cols)


def _get_or_create_wb(path: Path, sheet_name: str,
                       cols: List[Tuple], grp_cols: Dict, grp_starts: set):
    """파일 열기(기존) 또는 신규 생성. 시트명 기준으로 시트 선택. (wb, ws) 반환."""
    if path.exists():
        try:
            wb = load_workbook(str(path))
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                _extend_headers(ws, cols, grp_cols, grp_starts)   # 새 열(예: CH2 히터)이 생겼으면 머리글만 채움
            else:
                # 파일은 있는데 해당 시트가 없으면 새로 추가
                ws = wb.create_sheet(sheet_name)
                ws.sheet_view.showGridLines = False
                ws.freeze_panes = "A4"
                _build_headers(ws, cols, grp_cols, grp_starts)
            return wb, ws
        except Exception:
            broken = path.with_suffix(".broken.xlsx")
            with contextlib.suppress(Exception):
                path.rename(broken)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    _build_headers(ws, cols, grp_cols, grp_starts)
    return wb, ws


def _next_data_row(ws) -> int:
    """다음 빈 데이터 행 번호 반환 (헤더 3행 다음부터).

    [수정] ws.max_row는 서식(폰트/배경색)만 있고 값이 없는 phantom row도
    포함하므로 그대로 사용하면 데이터가 수백 행 아래로 밀린다.
    바닥에서 위로 스캔해 실제 값이 있는 마지막 행을 찾아 +1을 반환한다.
    """
    for r in range(ws.max_row, 3, -1):
        if any(cell.value is not None for cell in ws[r]):
            return r + 1
    return 4


# [수정 3] SP vs Avg 비교 쌍 정의
_SP_AVG_PAIRS = {
    # Main Process 시트
    "avg_ar"         : "sp_ar",
    "avg_n2"         : "sp_n2",
    "avg_o2"         : "sp_o2",
    "avg_pressure"   : "sp_pressure",
    "avg_forp"       : "sp_power",
    "avg_power"      : "sp_power",
    # Plasma Cleaning 시트 — PC 컬럼도 편차 하이라이트
    "pc_avg_ar"      : "pc_sp_ar",
    "pc_avg_pressure": "pc_sp_pressure",
    "pc_avg_forp"    : "pc_sp_power",
}


def _cell_bg(key: str, value: Any, row_bg: str,
             arc_thresh: int, refp_warn: float,
             data: Dict[str, Any]) -> str:
    """
    이상치 여부로 배경색 결정 (빨강으로 통일).
    - Arc 누적 >= arc_thresh → 빨강
    - SP vs Avg 10% 초과 편차 → 빨강
    """
    # Arc 누적 임계값 이상
    if key in ("soft_arc", "hard_arc"):
        if isinstance(value, (int, float)) and value >= arc_thresh:
            return _WARN_ARC

    # SP vs Avg 편차 (10% 초과 시 강조)
    if key in _SP_AVG_PAIRS:
        sp_key = _SP_AVG_PAIRS[key]
        sp_val = data.get(sp_key)
        if (isinstance(value, (int, float))
                and isinstance(sp_val, (int, float))
                and sp_val > 0):
            if abs(value - sp_val) / sp_val > _SP_DIFF_RATIO:
                return _WARN_ARC

    return row_bg


def _write_data_row(ws, row_num: int, data: Dict[str, Any],
                    arc_thresh: int, refp_warn: float,
                    cols: List[Tuple]) -> None:
    """데이터 딕셔너리를 ws의 row_num 행에 기록.

    [수정] 모든 값이 None인 빈 행에는 서식도 적용하지 않는다.
    서식만 적용된 행(phantom row)은 ws.max_row에 카운트되어
    _next_data_row()가 잘못된 행 번호를 반환하는 원인이 된다.
    """
    # 실제 쓸 값이 하나도 없으면 phantom row 방지를 위해 아무것도 하지 않음
    values = {key: data.get(key) for _, __, ___, ____, key in cols}
    if not any(v is not None and v != "" for v in values.values()):
        return

    thin = Side(style="thin", color="CCCCCC")
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)
    row_bg = _ROW_ODD if row_num % 2 == 0 else _ROW_EVEN
    # 컬럼 인덱스(ci)가 아닌 key 기반으로 왼쪽 정렬 결정.
    # _write_data_row()는 Main Process / Plasma Cleaning 시트 양쪽에서 호출되므로
    # ci 기준은 시트마다 컬럼 구조가 달라 부작용 발생.
    _LEFT_KEYS = {
        "process_name",  # Process Name (양 시트)
        "note",          # 비고 (양 시트)
        "substrate",     # 기판 (양 시트)
        "G1 Target",     # G1 Target (Main Process만)
        "G2 Target",     # G2 Target (Main Process만)
        "G3 Target",     # G3 Target (Main Process만)
    }

    for ci, (_, __, ___, ____, key) in enumerate(cols, 1):
        value = values[key]
        if value == "":
            value = None

        # [수정 3] data 딕셔너리 전달
        bg = _cell_bg(key, value, row_bg, arc_thresh, refp_warn, data)

        cell = ws.cell(row=row_num, column=ci)
        cell.value = value
        cell.font = Font(name="Arial", size=9, color=_FG)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.border = brd
        cell.alignment = Alignment(
            horizontal="left" if key in _LEFT_KEYS else "center",
            vertical="center")

        if key == "timestamp" and isinstance(value, datetime):
            cell.number_format = "YYYY-MM-DD HH:MM:SS"
        if key in ("base_pressure", "pc_base_pressure") \
                and isinstance(value, (int, float)):
            cell.number_format = "0.00E+00"

    ws.row_dimensions[row_num].height = 18
